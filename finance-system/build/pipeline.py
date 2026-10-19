# -*- coding: utf-8 -*-
"""خط الإنتاج: بناء ← إعادة حساب (LibreOffice) ← أخذ لقطة إقفال سبتمبر ← بناء نهائي ← تحقق ← سجل الاختبار."""
import argparse, datetime as dt, json, os, sys, time
import openpyxl
import build_system as bs
import verify as V
import b5_samples as b5

d = dt.date


def snapshot_from(wb):
    ws = wb["Period_Close"]
    hdr = {c.value: c.column for c in ws[4] if c.value}
    # find header cells by Arabic text prefix
    cols = {}
    for c in ws[4]:
        t = (c.value or "")
        if t.startswith("حي: صافي الإيرادات"):
            cols["rev"] = c.column
        elif t.startswith("حي: صافي الربح"):
            cols["np"] = c.column
        elif t.startswith("حي: إجمالي المدين"):
            cols["dr"] = c.column
        elif t.startswith("حي: إجمالي الدائن"):
            cols["cr"] = c.column
    for r in range(5, 40):
        if ws.cell(r, 1).value == 202609:
            return {k: round(float(ws.cell(r, v).value), 2) for k, v in cols.items()}
    raise RuntimeError("202609 not found")


def named(wb, nm):
    return V.named(wb, nm)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--log", default=None)
    ap.add_argument("--work", default=None)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    work = a.work or os.path.join(a.out, "_work")
    os.makedirs(work, exist_ok=True)
    log = []
    t0 = time.time()

    def L(s=""):
        print(s)
        log.append(s)

    # ---- تمرير 1: سبتمبر (الفترة الجارية = سبتمبر) لأخذ لقطة الإقفال والتحقق من جاهزية الإقفال
    p1, E = bs.build(work, demo=True, period=(d(2026, 9, 1), d(2026, 9, 30)), scale=a.scale, tag="_sep_pass")
    wb1, errs1 = V.report(p1, os.path.join(work, "calc1"))
    snap = snapshot_from(wb1)
    L(f"## تمرير 1 - Settings = سبتمبر 2026")
    L(f"- لقطة سبتمبر الحية: {json.dumps(snap, ensure_ascii=False)}")
    L(f"- البنود الحرجة المفتوحة (CK_Critical) = {named(wb1, 'CK_Critical')}، بنود المراجعة = {named(wb1, 'CK_Review')}")
    L(f"- ميزان المراجعة: TB_Diff={named(wb1,'TB_Diff')}؛ GL={named(wb1,'GL_TotalDiff')}؛ الميزانية={named(wb1,'BS_Diff')}؛ الموردون={named(wb1,'SB_Diff')}؛ التدفقات={named(wb1,'CF_Diff')}")
    L(f"- صافي المبيعات (قائمة الدخل) = {named(wb1,'IS_NetSales')} (المتوقع من المستندات مستقلًا: {E['sales_sep_net']})؛ صافي الربح = {named(wb1,'IS_NetProfit')}")
    ck = wb1["Close_Checklist"]
    open_items = []
    for r in range(10, 70):
        e = ck.cell(r, 5).value
        if e and str(e).startswith(("✖", "⚠")):
            open_items.append((ck.cell(r, 1).value, ck.cell(r, 3).value, ck.cell(r, 4).value, e))
    for it in open_items:
        L(f"  - {it}")
    assert abs(named(wb1, 'IS_NetSales') - E['sales_sep_net']) < 0.01, "صافي مبيعات سبتمبر لا يطابق الحساب المستقل"
    if errs1:
        L(f"- أخطاء صيغ في تمرير 1: {errs1}")

    # ---- التمرير النهائي: أكتوبر (الافتراضي) مع لقطة سبتمبر المقفلة
    ov = dict(sep_snapshot=snap, sep_closed=True)
    p2, E = bs.build(a.out, demo=True, period=(d(2026, 10, 1), d(2026, 10, 31)), ov=ov, scale=a.scale)
    wb2, errs2 = V.report(p2, os.path.join(work, "calc2"))
    L("")
    L("## التمرير النهائي - Settings = أكتوبر 2026 (الملف المسلَّم)")
    L(f"- أخطاء صيغ: {errs2 if errs2 else 'لا يوجد'}")
    for nm in ["TB_Diff", "GL_TotalDiff", "BS_Diff", "SB_Diff", "CF_Diff", "CF_InternalNet", "CK_Critical", "CK_Review", "IS_GrossSales", "IS_NetSales", "IS_COGS", "IS_GrossProfit", "IS_OpEx", "IS_NetProfit", "TEST_Pass", "TEST_Fail"]:
        L(f"- {nm} = {named(wb2, nm)}")
    ws = wb2["Test_Cases"]
    L("")
    L("| الرقم | الاختبار | الحالة | المتوقع | الفعلي |")
    L("|---|---|---|---|---|")
    fails = []
    for r in ws.iter_rows(min_row=8, values_only=True):
        if r[0]:
            esc = lambda v: str(v).replace("|", "\\|")
            L(f"| {r[0]} | {r[1]} | {r[5]} | {esc(r[3])} | {esc(r[4])} |")
            if r[5] == "فشل":
                fails.append(r[0])
    ck = wb2["Close_Checklist"]
    L("")
    L("### قائمة الإقفال لأكتوبر (متوقع وجود أخطاء مزروعة للاختبار)")
    for r in range(10, 70):
        e = ck.cell(r, 5).value
        if e and str(e).startswith(("✖", "⚠")):
            L(f"- {ck.cell(r,1).value}: {ck.cell(r,3).value} = {ck.cell(r,4).value} → {e}")
    assert not errs2, errs2
    assert not fails, fails

    # ---- النسخة النظيفة
    p3, _ = bs.build(a.out, demo=False, with_br04=False, scale=a.scale, tag="_CLEAN_TEMPLATE")
    wb3, errs3 = V.report(p3, os.path.join(work, "calc3"))
    L("")
    L("## النسخة النظيفة")
    L(f"- أخطاء صيغ: {errs3 if errs3 else 'لا يوجد'}")
    L(f"- TB_Diff={named(wb3,'TB_Diff')}؛ CK_Critical={named(wb3,'CK_Critical')} (الباقي بنود يدوية بانتظار التأكيد)")
    # ---- اختبارات سلبية/تحوّلية (مستندات خاطئة عمدًا) + اختبار تفعيل الضريبة
    import mutation_tests as MT
    res, _ = MT.run(os.path.join(work, "mut"), scale=min(a.scale, 0.3))
    L("")
    L("## اختبارات سلبية: إدخال مستندات خاطئة عمدًا في نسخة من الملف")
    L("| الحالة | الورقة | المتوقع | الفعلي | النتيجة |")
    L("|---|---|---|---|---|")
    for name, sh, row, exp, act, ok in res:
        a2 = "(لا أخطاء)" if (act is None or act == "" or (hasattr(act, "__len__") and len(act) == 0 and not isinstance(act, str))) else act
        L(f"| {name} | {sh} | {exp} | {a2} | {'نجح' if ok else 'فشل'} |")
    vat = MT.run_vat(os.path.join(work, "vat"), scale=min(a.scale, 0.3))
    L("")
    L(f"## اختبار تفعيل ضريبة القيمة المضافة (14%): ضريبة المخرجات المتوقعة {vat['expected_vat']} مقابل الأستاذ {vat['got']}؛ فرق الميزان {vat['tb']}؛ فرق الميزانية {vat['bs']} → {'نجح' if vat['ok'] else 'فشل'}")
    assert all(x[-1] for x in res), [x for x in res if not x[-1]]
    assert vat["ok"], vat
    L("")
    L(f"زمن التشغيل: {time.time() - t0:.0f} ثانية")
    if a.log:
        open(a.log, "w", encoding="utf-8").write("\n".join(log))
    return 0


if __name__ == "__main__":
    sys.exit(main())
