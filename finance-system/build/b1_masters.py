# -*- coding: utf-8 -*-
"""المجموعة A: الإعدادات والبيانات الأساسية."""
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from fs_common import *
import fs_data as D

CAP_COA = 250
CAP_LOC = 30
CAP_TRE = 30
CAP_PM = 15
CAP_CC = 30
CAP_SUP = 150
CAP_CUS = 100
CAP_EMP = 100
CAP_ITEM = 300
CAP_USER = 40

# مواقع أعمدة ورقة Lists
LIST_COLS = dict(ST="A", VT_AR="E", VT_CODE="F", VT_PFX="G", IT_AR="J", IT_CODE="K",
                 EM_AR="N", EM_CODE="O", YN="R", CAT_AR="T", CAT_ACCT="U")


def lcol(k):
    return f"Lists!${LIST_COLS[k]}:${LIST_COLS[k]}"


def lookup(sheet, ret, match, val):
    return f"INDEX({REG[sheet].ws_col(ret)},MATCH({val},{REG[sheet].ws_col(match)},0))"


def _wscol(self, key):
    L = self.L(key)
    return f"{self.name}!${L}:${L}"


Sheet.ws_col = _wscol


def band(ws, row, text, ncols=8, fill=C_HDR_BAND):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, bold=True, size=11, color=C_TITLE)
    for i in range(1, ncols + 1):
        ws.cell(row, i).fill = PatternFill("solid", fgColor=fill)


def mk_text_sheet(wb, name, title, subtitle, headers, rows, widths, tab=None, wrap=True, protect=False):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.zoomScale = 90
    if tab:
        ws.sheet_properties.tabColor = tab
    ws.cell(1, 1, title).font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, max(len(headers), 6) + 1):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.row_dimensions[1].height = 26
    ws.cell(2, 1, subtitle).font = Font(name=FONT, size=9, italic=True, color="404040")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(4, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=C_IN_H)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[GL(i)].width = widths[i - 1]
    for r, row in enumerate(rows, start=5):
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=wrap, vertical="top")
    ws.freeze_panes = "A5"
    if rows:
        ws.auto_filter.ref = f"A4:{GL(len(headers))}{4 + len(rows)}"
    if protect:
        protect_ws(ws)
    return ws


# --------------------------------------------------------------------- Lists
def build_lists(wb):
    ws = wb.create_sheet("Lists")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.cell(1, 1, "القوائم المنسدلة وحالات المستندات (Lists & Document Statuses)").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 44):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "هذه القوائم تغذي القوائم المنسدلة في كل الأوراق. لا تغيّر نصوص الحالات (مسودة/معتمد/مرحل/معكوس/ملغى) لأن المعادلات تعتمد عليها.").font = Font(name=FONT, size=9, italic=True)

    def block(col, title, rows, widths=None):
        c0 = ws[f"{col}4"]
        for j, h in enumerate(title):
            cc = ws.cell(4, c0.column + j, h)
            cc.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
            cc.fill = PatternFill("solid", fgColor=C_IN_H)
            cc.border = BORDER
            cc.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[GL(c0.column + j)].width = (widths[j] if widths else 16)
        for i, row in enumerate(rows, start=5):
            for j, v in enumerate(row):
                cc = ws.cell(i, c0.column + j, v)
                cc.font = Font(name=FONT, size=10)
                cc.border = BORDER
                cc.fill = PatternFill("solid", fgColor=C_IN)
                cc.protection = Protection(locked=False)

    block("A", ["حالة المستند", "Code", "المعنى"], D.STATUSES, [14, 12, 48])
    block("E", ["نوع السند", "Code", "بادئة الرقم", "الوصف"], D.VOUCHER_TYPES, [18, 8, 10, 44])
    block("J", ["نوع فاتورة المورد", "Code", "الوصف"], D.INVOICE_TYPES, [22, 8, 30])
    block("N", ["نمط سداد المصروف", "Code", "الوصف"], D.EXPENSE_MODES, [16, 8, 36])
    block("R", ["نعم/لا"], [("نعم",), ("لا",)], [10])
    block("T", ["فئة الإيراد", "حساب الإيراد"], D.REV_CATS, [16, 12])
    block("W", ["نوع القيد اليدوي"], [(t,) for t in D.JE_TYPES], [24])
    block("Y", ["سبب الفرق"], [(t,) for t in D.RECON_REASONS], [22])
    block("AA", ["حالة المعالجة"], [(t,) for t in D.RESOLUTIONS], [22])
    block("AC", ["تصنيف التدفق", "المعنى"], D.CF_CLASSES, [14, 28])
    block("AF", ["نوع الخزينة"], [(t,) for t in D.TREASURY_TYPES], [14])
    block("AH", ["نوع الموقع"], [(t,) for t in D.LOCATION_TYPES], [14])
    block("AJ", ["رمز النوع", "نوع الحساب"], D.ACCT_TYPES, [8, 16])
    block("AM", ["نوع وسيلة الدفع"], [(t,) for t in D.PM_TYPES], [14])
    block("AO", ["الأدوار"], [(t,) for t in D.ROLES], [18])
    names = dict(ST_DRAFT="Lists!$A$5", ST_APPR="Lists!$A$6", ST_POST="Lists!$A$7", ST_REV="Lists!$A$8", ST_CANC="Lists!$A$9",
                 LST_STATUS="Lists!$A$5:$A$9", LST_VTYPE="Lists!$E$5:$E$11", LST_ITYPE="Lists!$J$5:$J$7",
                 LST_EMODE="Lists!$N$5:$N$6", LST_YN="Lists!$R$5:$R$6", LST_CATS="Lists!$T$5:$T$8",
                 LST_JETYPE="Lists!$W$5:$W$14", LST_RSN="Lists!$Y$5:$Y$13", LST_RES="Lists!$AA$5:$AA$7",
                 LST_CF="Lists!$AC$5:$AC$9", LST_TTYPE="Lists!$AF$5:$AF$7", LST_LTYPE="Lists!$AH$5:$AH$8",
                 LST_ATYPE="Lists!$AJ$5:$AJ$10", LST_PMTYPE="Lists!$AM$5:$AM$8", LST_ROLE="Lists!$AO$5:$AO$14")
    for k, v in names.items():
        name_cell(wb, k, v)
        if k.startswith("LST_"):
            LIST_REFS[k] = v
    ws.protection.sheet = False  # قوائم يعدلها المسؤول فقط عبر صلاحيات جوجل
    return ws


# ------------------------------------------------------------------ Settings
def build_settings(wb, demo):
    ws = wb.create_sheet("Settings")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.cell(1, 1, "إعدادات النظام (System Settings)").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 6):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "الخلايا الصفراء قابلة للتعديل. فترة التقارير تتحكم في لوحة التحكم وميزان المراجعة وقائمة الدخل وقائمة الإقفال.").font = Font(name=FONT, size=9, italic=True)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 70
    rows = [
        ("اسم المنشأة", "مطاعم ستاموني (Satamoni)", None, "للعرض في التقارير", None),
        ("العملة الأساسية", "EGP", None, "جنيه مصري", None),
        ("المنطقة الزمنية", "Africa/Cairo", None, "اضبطها أيضًا من إعدادات ملف Google Sheets: ملف ← الإعدادات", None),
        ("بداية السنة المالية", dt.date(2026, 1, 1), "S_FYStart", "تُستخدم في الإقفال السنوي", "date"),
        ("بداية فترة التقارير", dt.date(2026, 10, 1), "S_From", "غيّرها لعرض شهر آخر (يجب أن تكون أول الشهر عند الإقفال)", "date"),
        ("نهاية فترة التقارير", dt.date(2026, 10, 31), "S_To", "آخر يوم في الفترة؛ هي أيضًا تاريخ الأرصدة في الميزانية", "date"),
        ("الفرع المختار للتقارير", "ALL", "S_Branch", "ALL = كل المواقع، أو كود موقع مثل BR01", None),
        ("مركز التكلفة المختار", "ALL", "S_CC", "ALL = كل المراكز، أو كود مثل CC10", None),
        ("رقم الفترة (آلي)", "=YEAR(C10)*100+MONTH(C10)", "S_PeriodID", "صيغة: سنة*100+شهر، تُستخدم في الإقفال", "f"),
        ("بداية فترة المقارنة (آلي)", "=C9-(C10-C9+1)", "S_PFrom", "فترة سابقة بنفس عدد الأيام", "f"),
        ("نهاية فترة المقارنة (آلي)", "=C9-1", "S_PTo", "", "f"),
        ("طريقة تسجيل تكلفة المبيعات", "PERIODIC", "S_CogsMethod", "PERIODIC = دوري مُعدّل (النسخة الأولى: قيد شهري بالجرد) | PERPETUAL = مستمر من حركات المخزون (المرحلة 2). لا تجمع بينهما في نفس الفترة.", None),
        ("هامش السماح بالفروق (جنيه)", 0.01, "S_Tol", "للمقارنات الحسابية (تقريب)", None),
        ("أيام تنبيه تأخر تسوية المنصات", 14, "S_PlatDays", "تنبيه إن مضى هذا العدد دون تسوية ومستحقات المنصة موجودة", None),
        ("حد فرق الخزينة المسموح (جنيه)", 0, "S_CashTol", "أي فرق أكبر منه يجب أن يُسجل بسند عجز/زيادة", None),
        ("وضع البيانات", "تجريبي" if demo else "تشغيلي", "S_Mode", "تجريبي = بيانات الاختبار؛ غيّرها إلى (تشغيلي) بعد تفريغ البيانات التجريبية", None),
    ]
    ws.cell(4, 2, "الإعداد").font = Font(name=FONT, bold=True)
    ws.cell(4, 3, "القيمة").font = Font(name=FONT, bold=True)
    ws.cell(4, 4, "ملاحظات").font = Font(name=FONT, bold=True)
    for i, (lab, val, nm, note, kind) in enumerate(rows, start=5):
        ws.cell(i, 2, lab).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(i, 3, val)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        if kind == "date" or kind == "f" and "C9" in str(val):
            c.number_format = "yyyy-mm-dd"
        if kind == "f":
            c.fill = PatternFill("solid", fgColor=C_F)
        else:
            c.fill = PatternFill("solid", fgColor=C_IN)
            c.protection = Protection(locked=False)
        ws.cell(i, 4, note).font = Font(name=FONT, size=9, color="404040")
        if nm:
            name_cell(wb, nm, f"Settings!$C${i}")
    ws["C14"].number_format = "yyyy-mm-dd"
    ws["C15"].number_format = "yyyy-mm-dd"
    dv = DataValidation(type="list", formula1='"PERIODIC,PERPETUAL"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C16")
    # أرقام الصفوف أعلاه ثابتة بالتصميم: C9=From, C10=To
    protect_ws(ws)
    return ws


# ---------------------------------------------------------------- Tax sheet
def build_tax(wb):
    ws = wb.create_sheet("Tax_Settings")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.cell(1, 1, "إعدادات الضرائب (Tax Settings) - إن انطبقت").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 5):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 80
    ws.cell(2, 1, "تنبيه: لم تتم مراجعة المعالجة الضريبية مع محاسب قانوني مصري. الافتراضي في النسخة الأولى = ضريبة القيمة المضافة غير مفعّلة.").font = Font(name=FONT, size=9, bold=True, color="C00000")
    data = [
        ("تفعيل ضريبة القيمة المضافة؟", "لا", "S_VATOn", "إذا (نعم): تعتبر مبالغ المبيعات والمشتريات شاملة للضريبة وتفصل المعادلات الضريبة تلقائيًا في قيد المبيعات."),
        ("نسبة ضريبة القيمة المضافة", 0.14, "S_VATRate", "النسبة القياسية المفترضة 14% - تحقق منها مع المحاسب القانوني"),
        ("معامل الفصل (آلي)", '=IF(S_VATOn="نعم",1+S_VATRate,1)', "S_VATF", "قيمة الإيراد بدون ضريبة = المبلغ الشامل ÷ هذا المعامل"),
        ("حساب ضريبة المخرجات", "=ACC_VAT_OUT", None, "من خريطة الحسابات (ورقة Posting_Rules)"),
        ("حساب ضريبة المدخلات", "=ACC_VAT_IN", None, "من خريطة الحسابات"),
    ]
    ws.cell(4, 2, "البند").font = Font(name=FONT, bold=True)
    ws.cell(4, 3, "القيمة").font = Font(name=FONT, bold=True)
    ws.cell(4, 4, "ملاحظات").font = Font(name=FONT, bold=True)
    for i, (lab, val, nm, note) in enumerate(data, start=5):
        ws.cell(i, 2, lab).font = Font(name=FONT, bold=True, size=10)
        c = ws.cell(i, 3, val)
        c.border = BORDER
        ws.cell(i, 4, note).font = Font(name=FONT, size=9)
        if isinstance(val, str) and val.startswith("="):
            c.fill = PatternFill("solid", fgColor=C_F)
        else:
            c.fill = PatternFill("solid", fgColor=C_IN)
            c.protection = Protection(locked=False)
        if nm:
            name_cell(wb, nm, f"Tax_Settings!$C${i}")
    ws["C6"].number_format = "0.0%"
    dv = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C5")
    notes = ["ملاحظات ضريبية (للمراجعة مع المحاسب القانوني):",
             "1) تسجيل ضريبة المدخلات في المشتريات والمصروفات يتم فقط عند وجود فاتورة ضريبية صحيحة.",
             "2) النظام لا يُعد إقرارات ضريبية ولا يُغني عن مراجعة قانونية؛ هو يفصل حسابات الضريبة فقط لتسهيل الإعداد.",
             "3) الضرائب الأخرى (دمغة، جدول، دخل) غير مفعلة في النسخة الأولى؛ استخدم حساب 2302 يدويًا عبر قيد يدوي."]
    for i, t in enumerate(notes, start=11):
        ws.cell(i, 2, t).font = Font(name=FONT, size=10, bold=(i == 11))
    protect_ws(ws)
    return ws


# ----------------------------------------------------------- Master sheets
def define_masters(demo, with_br04):
    # ---------------- COA
    coa = Sheet("COA", "دليل الحسابات الموحد (Chart of Accounts)",
                "ترقيم 4 خانات: الفئة (1-6) ثم المجموعة (100) ثم الحساب. أضف حسابًا جديدًا في أي صف فارغ بالنطاق المناسب ويظهر فورًا في ميزان المراجعة. الحسابات الرئيسية (لا) لا يُرحَّل عليها.",
                cap=CAP_COA, tab="2E75B6")
    coa.add("code", "كود الحساب", w=10, dv=("int",), en="Code", align="center")
    coa.add("ar", "اسم الحساب", w=44, en="Name (AR)")
    coa.add("en", "Account Name (EN)", w=34)
    coa.add("t", "النوع", w=8, dv=("list", "=LST_ATYPE"), align="center")
    coa.add("lvl", "المستوى", w=8, dv=("list", "1,2,3"), align="center")
    coa.add("post", "قابل للترحيل", w=10, dv=("list", "=LST_YN"), align="center")
    coa.add("normal", "الطبيعة", w=9, dv=("list", "مدين,دائن"), align="center")
    coa.add("cash", "نقدية/بنك (1/0)", w=10, dv=("list", "0,1"), align="center")
    coa.add("cfclass", "تصنيف التدفق", w=12, dv=("list", "=LST_CF"))
    coa.add("ae", "مسموح في المصروفات", w=10, dv=("list", "=LST_YN"), align="center")
    coa.add("ap", "مسموح في المشتريات", w=10, dv=("list", "=LST_YN"), align="center")
    coa.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    coa.add("rf", "من نطاق (للرئيسي)", w=10, dv=("int",))
    coa.add("rt", "إلى نطاق (للرئيسي)", w=10, dv=("int",))
    coa.add("note", "ملاحظات", w=40)
    coa.add("rng_from", "بداية النطاق", "f", 10, f=lambda x: f'=IF({x("code")}="","",IF({x("post")}="نعم",{x("code")},{x("rf")}))')
    coa.add("rng_to", "نهاية النطاق", "f", 10, f=lambda x: f'=IF({x("code")}="","",IF({x("post")}="نعم",{x("code")},{x("rt")}))')
    coa.add("valid", "صالح للترحيل (1/0)", "f", 9, f=lambda x: f'=IF({x("code")}="","",IF(AND({x("post")}="نعم",{x("active")}<>"لا"),1,0))')
    coa.add("tname", "نوع الحساب", "f", 12, f=lambda x: f'=IF({x("t")}="","",IFERROR(INDEX(Lists!$AK$5:$AK$10,MATCH({x("t")},Lists!$AJ$5:$AJ$10,0)),"نوع غير معرف"))')
    coa.add("chk", "فحص", "f", 40, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود حساب مكرر"),
        (f'{x("ar")}=""', "اسم الحساب ناقص"),
        (f'AND({x("post")}="نعم",{x("normal")}="")', "الطبيعة ناقصة"),
        (f'AND({x("post")}<>"نعم",OR({x("rf")}="",{x("rt")}=""))', "الحساب الرئيسي يحتاج نطاقًا (من/إلى)"),
        (f'AND({x("cash")}=1,{x("post")}<>"نعم")', "حساب نقدية يجب أن يكون قابلًا للترحيل"),
        (f'AND({x("post")}="نعم",{x("t")}="")', "النوع ناقص"),
    ]) + ')')
    rows = []
    for a in D.COA:
        if a["code"] == 1107 and not with_br04:
            continue
        r = dict(code=a["code"], ar=a["ar"], en=a["en"], t=a["t"], lvl=a["lvl"], post=a["post"], normal=a["normal"],
                 cash=a["cash"], cfclass=a["cf"], ae=a["ae"], ap=a["ap"], active="نعم", note=a["note"])
        if a["post"] == "لا":
            r["rf"], r["rt"] = D.HEADER_SPAN[a["code"]]
        rows.append(r)
    coa.samples = rows

    # ---------------- Locations
    loc = Sheet("Locations", "الفروع والمواقع (Branches & Locations)",
                "كل فرع/مخزن/مطبخ له كود ثابت. لإضافة فرع جديد: أضف صفًا هنا، ثم خزينة له في Treasuries وحسابًا في COA - وتظهر بياناته آليًا في التقارير.",
                cap=CAP_LOC, tab="2E75B6")
    loc.add("code", "كود الموقع", w=10, en="Code", align="center")
    loc.add("name", "اسم الموقع", w=34)
    loc.add("type", "النوع", w=14, dv=("list", "=LST_LTYPE"))
    loc.add("en", "Name (EN)", w=22)
    loc.add("mgr", "المسؤول", w=18)
    loc.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    loc.add("cash_tre", "خزينة النقدية الافتراضية", w=14, dv=("list", "=LST_TRE"))
    loc.add("opened", "تاريخ الافتتاح", w=12, fmt="yyyy-mm-dd", dv=("date",))
    loc.add("note", "ملاحظات", w=30)
    loc.add("seq", "ترتيب فرع البيع", "f", 9, f=lambda x: f'=IF(AND({x("type")}="فرع",{x("active")}<>"لا",{x("code")}<>""),MAX(${loc.L("seq")}${loc.hdr}:${loc.L("seq")}{x.r - 1})+1,"")')
    loc.add("chk", "فحص", "f", 36, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود موقع مكرر"),
        (f'{x("name")}=""', "اسم الموقع ناقص"),
        (f'{x("type")}=""', "نوع الموقع ناقص"),
        (f'AND({x("cash_tre")}<>"",COUNTIF({x.ext("Treasuries","code")},{x("cash_tre")})=0)', "الخزينة الافتراضية غير معرفة"),
        (f'AND({x("type")}="فرع",{x("cash_tre")}="")', "الفرع يحتاج خزينة افتراضية"),
    ]) + ')')
    rows = [dict(code=a, name=b, type=c, en=d, mgr=e, active=f, cash_tre=g or None, opened=dt.date(2026, 1, 1)) for a, b, c, d, e, f, g in D.LOCATIONS]
    if with_br04:
        a, b, c, d, e, f, g = D.LOCATION_TEST_BR04
        rows.append(dict(code=a, name=b, type=c, en=d, mgr=e, active=f, cash_tre=g, opened=dt.date(2026, 10, 1), note="موقع تجريبي لاختبار إضافة فرع جديد"))
    loc.samples = rows

    # ---------------- Treasuries
    tre = Sheet("Treasuries", "الخزائن والبنوك (Treasuries & Banks)",
                "كل خزينة/بنك مرتبط بحساب مستقل في دليل الحسابات (حساب واحد لكل خزينة). الأرصدة الافتتاحية تُسجَّل بقيد افتتاحي في Manual_JE وليس هنا.",
                cap=CAP_TRE, tab="2E75B6")
    tre.add("code", "كود الخزينة/البنك", w=14, align="center")
    tre.add("name", "الاسم", w=34)
    tre.add("type", "النوع", w=12, dv=("list", "=LST_TTYPE"))
    tre.add("loc", "الموقع", w=10, dv=("list", "=LST_LOC"), align="center")
    tre.add("gl", "حساب الأستاذ", w=11, dv=("list", "=LST_ACCT"), align="center")
    tre.add("custodian", "أمين الخزينة/المسؤول", w=22)
    tre.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    tre.add("note", "ملاحظات", w=26)
    tre.add("gl_name", "اسم الحساب", "f", 34, f=lambda x: f'=IF({x("gl")}="","",IFERROR(INDEX({REG["COA"].ws_col("ar")},MATCH({x("gl")},{REG["COA"].ws_col("code")},0)),"حساب غير موجود"))')
    tre.add("chk", "فحص", "f", 40, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر"),
        (f'COUNTIF({x.col("gl")},{x("gl")})>1', "حساب الأستاذ مستخدم لخزينة أخرى"),
        (f'COUNTIFS({x.ext("COA","code")},{x("gl")},{x.ext("COA","valid")},1,{x.ext("COA","cash")},1)<>1', "الحساب غير موجود أو غير نقدي أو غير قابل للترحيل"),
        (f'COUNTIF({x.ext("Locations","code")},{x("loc")})=0', "الموقع غير معرف"),
    ]) + ')')
    rows = [dict(code=a, name=b, type=c, loc=d, gl=e, custodian=f, active="نعم") for a, b, c, d, e, f in D.TREASURIES]
    if with_br04:
        a, b, c, d, e, f = D.TREASURY_TEST_BR04
        rows.append(dict(code=a, name=b, type=c, loc=d, gl=e, custodian=f, active="نعم", note="تجريبي"))
    tre.samples = rows

    # ---------------- Payment methods
    pm = Sheet("Payment_Methods", "وسائل الدفع (Payment Methods)",
               "وسيلة (نقدي) تُرحَّل على خزينة الفرع آليًا. الوسائل الأخرى تُرحَّل على حساب المستحقات المحدد هنا حتى يتم التحصيل الفعلي.",
               cap=CAP_PM, tab="2E75B6")
    pm.add("code", "كود الوسيلة", w=10, align="center")
    pm.add("name", "الاسم", w=30)
    pm.add("type", "النوع", w=12, dv=("list", "=LST_PMTYPE"))
    pm.add("gl", "حساب المستحقات/التسوية", w=12, dv=("list", "=LST_ACCT"), align="center")
    pm.add("comm_acct", "حساب العمولة/الرسوم", w=12, dv=("list", "=LST_ACCT"), align="center")
    pm.add("comm_pct", "نسبة العمولة المقدّرة", w=10, fmt="0.0%")
    pm.add("lag", "أيام التسوية المتوقعة", w=10, dv=("int",))
    pm.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    pm.add("note", "ملاحظات", w=30)
    pm.add("gl_name", "اسم حساب التسوية", "f", 34, f=lambda x: f'=IF({x("gl")}="","",IFERROR(INDEX({REG["COA"].ws_col("ar")},MATCH({x("gl")},{REG["COA"].ws_col("code")},0)),"حساب غير موجود"))')
    pm.add("chk", "فحص", "f", 40, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر"),
        (f'AND({x("type")}<>"نقدي",COUNTIFS({x.ext("COA","code")},{x("gl")},{x.ext("COA","valid")},1)<>1)', "حساب التسوية غير صالح"),
        (f'AND({x("type")}="منصة",{x("comm_acct")}="")', "المنصة تحتاج حساب عمولة"),
    ]) + ')')
    pm.samples = [dict(code=a, name=b, type=c, gl=d, comm_acct=e, comm_pct=f, lag=g, active="نعم") for a, b, c, d, e, f, g in D.PAYMENT_METHODS]

    # ---------------- Cost centers
    cc = Sheet("Cost_Centers", "مراكز التكلفة (Cost Centers)", "مراكز التكلفة تتيح تحليل المصروفات حسب الوظيفة بعيدًا عن الفرع.", cap=CAP_CC, tab="2E75B6")
    cc.add("code", "كود المركز", w=10, align="center")
    cc.add("name", "الاسم", w=34)
    cc.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    cc.add("note", "ملاحظات", w=30)
    cc.add("chk", "فحص", "f", 30, f=lambda x: '=IF(' + x("code") + '="","",' + chain([(f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر")]) + ')')
    cc.samples = [dict(code=a, name=b, active="نعم") for a, b in D.COST_CENTERS]

    # ---------------- Users
    us = Sheet("Users", "المستخدمون والصلاحيات (Users)", "تُستخدم أسماؤهم في خانات المُعِد والمعتمد. الصلاحيات الفعلية تُضبط في مشاركة Google Sheets وحماية الأوراق.", cap=CAP_USER, tab="2E75B6")
    us.add("code", "كود المستخدم", w=10, align="center")
    us.add("name", "الاسم (يُستخدم في الخانات)", w=30)
    us.add("email", "بريد جوجل", w=28)
    us.add("role", "الدور", w=18, dv=("list", "=LST_ROLE"))
    us.add("loc", "نطاق الموقع", w=12, dv=("list", "=LST_LOC"))
    us.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    us.add("note", "ملاحظات", w=26)
    us.add("chk", "فحص", "f", 30, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("name")},{x("name")})>1', "اسم مكرر"), (f'{x("role")}=""', "الدور ناقص")]) + ')')
    users = [("USR01", "المحاسب (تجريبي)", "المحاسب", "HQ01"), ("USR02", "المدير المالي (تجريبي)", "المدير المالي", "HQ01"),
             ("USR03", "مدير الفرع 1 (تجريبي)", "مدير فرع", "BR01"), ("USR04", "مدير الفرع 2 (تجريبي)", "مدير فرع", "BR02"),
             ("USR05", "مدير الفرع 3 (تجريبي)", "مدير فرع", "BR03"), ("USR06", "أمين صندوق فرع 1 (تجريبي)", "أمين صندوق", "BR01"),
             ("USR07", "أمين المخزن (تجريبي)", "أمين مخزن", "WH01"), ("USR08", "مشرف المشتريات (تجريبي)", "مشرف مشتريات", "HQ01"),
             ("USR09", "مراجع الحسابات (تجريبي)", "مراجع الحسابات", "HQ01"), ("USR10", "الإدارة العليا (تجريبي)", "الإدارة العليا", "HQ01"),
             ("USR11", "أمين صندوق فرع 2 (تجريبي)", "أمين صندوق", "BR02"), ("USR12", "مدير الإنتاج (تجريبي)", "مدير إنتاج", "CK01"),
             ("USR13", "أمين صندوق الإدارة (تجريبي)", "أمين صندوق", "HQ01")]
    us.samples = [dict(code=a, name=b, role=c, loc=d, active="نعم", email="user@example.com") for a, b, c, d in users]

    # ---------------- Suppliers
    sup = Sheet("Suppliers", "الموردون (Suppliers)", "كود المورد ثابت ولا يتكرر. شروط السداد (أيام) تحدد تاريخ الاستحقاق في فواتير الموردين.", cap=CAP_SUP, tab="2E75B6")
    sup.add("code", "كود المورد", w=10, align="center")
    sup.add("name", "اسم المورد", w=30)
    sup.add("cat", "التصنيف", w=16)
    sup.add("terms", "شروط السداد (أيام)", w=10, dv=("int",))
    sup.add("phone", "الهاتف", w=14)
    sup.add("taxid", "الرقم الضريبي", w=14)
    sup.add("dacct", "حساب المدين الافتراضي", w=11, dv=("list", "=LST_ACCT"), align="center")
    sup.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    sup.add("note", "ملاحظات", w=26)
    sup.add("chk", "فحص", "f", 34, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود مورد مكرر"), (f'{x("name")}=""', "اسم المورد ناقص"),
        (f'{x("terms")}=""', "شروط السداد ناقصة")]) + ')')
    if demo:
        sup.samples = [dict(code="SUP001", name="مورد أجبان تجريبي", cat="ألبان وأجبان", terms=30, dacct=1401, active="نعم"),
                       dict(code="SUP002", name="مورد دقيق ومخبوزات تجريبي", cat="دقيق", terms=15, dacct=1401, active="نعم"),
                       dict(code="SUP003", name="مورد عبوات وتغليف تجريبي", cat="تغليف", terms=0, dacct=1405, active="نعم"),
                       dict(code="SUP004", name="مورد معدات مطابخ تجريبي", cat="أصول", terms=30, dacct=1701, active="نعم"),
                       dict(code="SUP005", name="مورد خضروات تجريبي", cat="خضروات", terms=7, dacct=1401, active="نعم")]

    # ---------------- Customers
    cus = Sheet("Customers", "العملاء والمدينون (Customers)", "للعملاء الآجلين فقط (شركات/تموين). مبيعات الفروع والمنصات لا تحتاج عميلًا.", cap=CAP_CUS, tab="2E75B6")
    cus.add("code", "كود العميل", w=10, align="center")
    cus.add("name", "اسم العميل", w=30)
    cus.add("terms", "شروط السداد (أيام)", w=10, dv=("int",))
    cus.add("limit", "حد الائتمان", w=12, fmt="#,##0.00", dv=("dec",))
    cus.add("phone", "الهاتف", w=14)
    cus.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    cus.add("note", "ملاحظات", w=26)
    cus.add("chk", "فحص", "f", 30, f=lambda x: '=IF(' + x("code") + '="","",' + chain([(f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر"), (f'{x("name")}=""', "الاسم ناقص")]) + ')')
    if demo:
        cus.samples = [dict(code="CUS001", name="شركة تموين تجريبية", terms=30, limit=50000, active="نعم")]

    # ---------------- Employees
    emp = Sheet("Employees", "الموظفون والسلف والعهد (Employees)", "أرصدة السلف والعهد تُحسب من الأستاذ العام (حسابا 1501 و1502) وفق كود الموظف.", cap=CAP_EMP, tab="2E75B6")
    emp.add("code", "كود الموظف", w=10, align="center")
    emp.add("name", "الاسم", w=26)
    emp.add("loc", "الموقع", w=10, dv=("list", "=LST_LOC"), align="center")
    emp.add("job", "الوظيفة", w=16)
    emp.add("hired", "تاريخ التعيين", w=12, fmt="yyyy-mm-dd", dv=("date",))
    emp.add("salary", "الراتب الأساسي", w=12, fmt="#,##0.00", dv=("dec",))
    emp.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    emp.add("note", "ملاحظات", w=22)
    emp.add("adv", "رصيد السلف", "f", 12, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",SUMIFS({G("effdr")},{G("acct")},1501,{G("party")},{x("code")})-SUMIFS({G("effcr")},{G("acct")},1501,{G("party")},{x("code")}))')
    emp.add("cust", "رصيد العهدة", "f", 12, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",SUMIFS({G("effdr")},{G("acct")},1502,{G("party")},{x("code")})-SUMIFS({G("effcr")},{G("acct")},1502,{G("party")},{x("code")}))')
    emp.add("chk", "فحص", "f", 30, f=lambda x: '=IF(' + x("code") + '="","",' + chain([(f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر"), (f'{x("name")}=""', "الاسم ناقص")]) + ')')
    if demo:
        emp.samples = [dict(code="EMP001", name="موظف تجريبي 1", loc="BR01", job="كاشير", salary=4500, active="نعم"),
                       dict(code="EMP002", name="موظف تجريبي 2", loc="BR02", job="طباخ", salary=5200, active="نعم")]

    # ---------------- Units
    un = Sheet("Units", "وحدات القياس (Units of Measure)", "الوحدة الأساسية لكل صنف تُحدد في بطاقة الصنف. التحويل بين الوحدات يتم فقط عبر Unit_Conversions المعتمد.", cap=30, tab="2E75B6")
    un.add("code", "كود الوحدة", w=10, align="center")
    un.add("name", "الاسم", w=18)
    un.add("type", "النوع", w=12, dv=("list", "وزن,حجم,عدد,عبوة"))
    un.add("note", "ملاحظات", w=30)
    un.add("chk", "فحص", "f", 26, f=lambda x: '=IF(' + x("code") + '="","",' + chain([(f'COUNTIF({x.col("code")},{x("code")})>1', "كود مكرر")]) + ')')
    un.samples = [dict(code="KG", name="كيلوجرام", type="وزن"), dict(code="G", name="جرام", type="وزن"), dict(code="L", name="لتر", type="حجم"),
                  dict(code="ML", name="مللي لتر", type="حجم"), dict(code="PC", name="قطعة", type="عدد"), dict(code="PACK", name="عبوة", type="عبوة"),
                  dict(code="CTN", name="كرتونة", type="عبوة"), dict(code="BAG", name="شكارة", type="عبوة")]

    # ---------------- Unit conversions
    uc = Sheet("Unit_Conversions", "جدول تحويل الوحدات (Unit Conversions)",
               "التحويل لا يعمل آليًا إلا إذا كان المعامل معرّفًا ومعتمدًا (نعم). اكتب GENERAL في خانة الصنف للمعامل العام، أو كود الصنف للمعامل الخاص به (مثل: كرتونة البيض = 30 قطعة).",
               cap=60, tab="2E75B6")
    uc.add("frm", "من وحدة", w=10, dv=("list", "=LST_UNIT"), align="center")
    uc.add("to", "إلى وحدة", w=10, dv=("list", "=LST_UNIT"), align="center")
    uc.add("factor", "المعامل (1 من = ؟ إلى)", w=14, dv=("dec",))
    uc.add("item", "الصنف (GENERAL أو كود)", w=16)
    uc.add("appr", "معتمد؟", w=9, dv=("list", "=LST_YN"), align="center")
    uc.add("by", "المعتمد", w=18)
    uc.add("note", "ملاحظات", w=26)
    uc.add("chk", "فحص", "f", 40, f=lambda x: '=IF(' + x("frm") + '&' + x("to") + '="","",' + chain([
        (f'OR({x("frm")}="",{x("to")}="")', "الوحدة ناقصة"), (f'{x("frm")}={x("to")}', "الوحدتان متطابقتان"),
        (f'NOT(ISNUMBER({x("factor")}))', "المعامل غير معرف"), (f'{x("factor")}<=0', "المعامل يجب أن يكون موجبًا"),
        (f'{x("item")}=""', "الصنف ناقص (GENERAL أو كود)"),
        (f'COUNTIFS({x.col("frm")},{x("frm")},{x.col("to")},{x("to")},{x.col("item")},{x("item")})>1', "معامل مكرر لنفس الوحدتين والصنف"),
        (f'AND({x("appr")}="نعم",{x("by")}="")', "اسم المعتمد ناقص"),
    ]) + ')')
    uc.samples = [dict(frm="KG", to="G", factor=1000, item="GENERAL", appr="نعم", by="المدير المالي (تجريبي)"),
                  dict(frm="L", to="ML", factor=1000, item="GENERAL", appr="نعم", by="المدير المالي (تجريبي)"),
                  dict(frm="CTN", to="PC", factor=30, item="ITM010", appr="نعم", by="المدير المالي (تجريبي)"),
                  dict(frm="BAG", to="KG", factor=25, item="ITM002", appr="نعم", by="المدير المالي (تجريبي)"),
                  dict(frm="PACK", to="KG", factor=2.5, item="ITM001", appr="لا", note="غير معتمد: لن يعمل التحويل")]

    # ---------------- Unit converter (يرفض التحويل غير المعتمد)
    ucv = Sheet("Unit_Converter", "محوّل الوحدات المضبوط (Unit Converter)",
                "يحوّل كمية من أي وحدة إلى الوحدة الأساسية للصنف فقط عبر معامل معرّف ومعتمد في Unit_Conversions؛ وإلا يظهر (التحويل غير مسموح) ولا يُنتج رقمًا. لا يُعكس معامل تلقائيًا.",
                cap=40, tab="2E75B6")
    ucv.add("item", "كود الصنف", w=12, dv=("list", "=LST_ITEM"), align="center")
    ucv.add("qty", "الكمية", w=12, dv=("dec",))
    ucv.add("frm", "الوحدة المدخلة", w=12, dv=("list", "=LST_UNIT"), align="center")
    UCc = lambda k: REG["Unit_Conversions"].ws_col(k)
    ucv.add("base", "الوحدة الأساسية للصنف", "f", 14, align="center",
            f=lambda x: f'=IF({x("item")}="","",IF(COUNTIF({REG["Items"].ws_col("code")},{x("item")})=0,"صنف غير معرف",INDEX({REG["Items"].ws_col("unit")},MATCH({x("item")},{REG["Items"].ws_col("code")},0))))')
    ucv.add("factor", "المعامل المعتمد", "f", 12,
            f=lambda x: (f'=IF(OR({x("item")}="",{x("frm")}=""),"",IF({x("base")}="صنف غير معرف","",IF({x("frm")}={x("base")},1,'
                         f'IF(COUNTIFS({UCc("frm")},{x("frm")},{UCc("to")},{x("base")},{UCc("item")},{x("item")},{UCc("appr")},"نعم",{UCc("factor")},">0")=1,'
                         f'SUMIFS({UCc("factor")},{UCc("frm")},{x("frm")},{UCc("to")},{x("base")},{UCc("item")},{x("item")},{UCc("appr")},"نعم"),'
                         f'IF(COUNTIFS({UCc("frm")},{x("frm")},{UCc("to")},{x("base")},{UCc("item")},"GENERAL",{UCc("appr")},"نعم",{UCc("factor")},">0")=1,'
                         f'SUMIFS({UCc("factor")},{UCc("frm")},{x("frm")},{UCc("to")},{x("base")},{UCc("item")},"GENERAL",{UCc("appr")},"نعم"),"")))))'))
    ucv.add("conv", "الكمية بالوحدة الأساسية", "f", 16, fmt="#,##0.####",
            f=lambda x: f'=IF({x("factor")}="","",IF(ISNUMBER({x("qty")}),{x("qty")}*{x("factor")},""))')
    ucv.add("flag", "الحالة", "f", 44,
            f=lambda x: f'=IF(OR({x("item")}="",{x("frm")}=""),"",IF({x("base")}="صنف غير معرف","صنف غير معرف",IF({x("factor")}="","التحويل غير مسموح: معامل غير معرّف أو غير معتمد","")))')
    if demo:
        ucv.samples = [dict(item="ITM010", qty=2, frm="CTN"), dict(item="ITM002", qty=3, frm="BAG"), dict(item="ITM001", qty=4, frm="PACK"), dict(item="ITM001", qty=500, frm="G"), dict(item="ITM006", qty=2500, frm="ML")]

    # ---------------- Items
    it = Sheet("Items", "بطاقات الأصناف (Item Master)", "لكل صنف وحدة أساسية واحدة لا تتغير. الأرصدة تُحسب من حركات المخزون (المرحلة 2) وليس من هنا.", cap=CAP_ITEM, tab="2E75B6")
    it.add("code", "كود الصنف", w=10, align="center")
    it.add("name", "اسم الصنف", w=30)
    it.add("cat", "التصنيف", w=16)
    it.add("unit", "الوحدة الأساسية", w=10, dv=("list", "=LST_UNIT"), align="center")
    it.add("kind", "النوع", w=16, dv=("list", "مادة خام,نصف مصنع,منتج تام,عبوات,مستهلكات"))
    it.add("reorder", "حد إعادة الطلب", w=10, dv=("dec",))
    it.add("shelf", "مدة الصلاحية (يوم)", w=10, dv=("int",))
    it.add("sup", "المورد الافتراضي", w=10, dv=("list", "=LST_SUP"), align="center")
    it.add("stdcost", "تكلفة قياسية (اختياري)", w=12, fmt="#,##0.00", dv=("dec",))
    it.add("inv_acct", "حساب المخزون", w=10, dv=("list", "=LST_ACCT"), align="center")
    it.add("active", "نشط", w=7, dv=("list", "=LST_YN"), align="center")
    it.add("chk", "فحص", "f", 36, f=lambda x: '=IF(' + x("code") + '="","",' + chain([
        (f'COUNTIF({x.col("code")},{x("code")})>1', "كود صنف مكرر"), (f'{x("unit")}=""', "الوحدة الأساسية ناقصة"),
        (f'COUNTIF({x.ext("Units","code")},{x("unit")})=0', "وحدة غير معرفة")]) + ')')
    if demo:
        items = [("ITM001", "موزاريلا", "ألبان", "KG", "مادة خام", 50, 45, "SUP001", 1401), ("ITM002", "دقيق", "دقيق", "KG", "مادة خام", 200, 180, "SUP002", 1401),
                 ("ITM003", "صلصة طماطم", "صلصات", "KG", "نصف مصنع", 40, 10, None, 1403), ("ITM004", "عجينة بيتزا", "عجين", "KG", "نصف مصنع", 60, 3, None, 1403),
                 ("ITM005", "جبنة رومي", "ألبان", "KG", "مادة خام", 20, 60, "SUP001", 1401), ("ITM006", "زيت", "زيوت", "L", "مادة خام", 40, 365, "SUP005", 1401),
                 ("ITM007", "مكرونة", "جافة", "KG", "مادة خام", 80, 365, "SUP002", 1401), ("ITM008", "علب بيتزا", "تغليف", "PC", "عبوات", 500, 730, "SUP003", 1405),
                 ("ITM009", "أكياس تغليف", "تغليف", "PC", "عبوات", 1000, 730, "SUP003", 1405), ("ITM010", "بيض", "بيض", "PC", "مادة خام", 300, 21, "SUP005", 1401)]
        it.samples = [dict(code=a, name=b, cat=c, unit=d, kind=e, reorder=f, shelf=g, sup=h, inv_acct=i, active="نعم") for a, b, c, d, e, f, g, h, i in items]

    # ---------------- Doc numbering
    dn = Sheet("Doc_Numbering", "ترقيم المستندات (Document Numbering)",
               "صيغة الرقم: البادئة-كود الموقع-تسلسل 4 خانات (مثال SAL-BR01-0001). التسلسل مطبوع مسبقًا على النماذج الورقية. السجل المركزي Doc_Register يمنع التكرار.",
               cap=40, tab="2E75B6")
    dn.add("prefix", "البادئة", w=10, align="center")
    dn.add("name", "نوع المستند", w=30)
    dn.add("form", "النموذج الورقي", w=26)
    dn.add("sheet", "ورقة التسجيل", w=22)
    dn.add("phase", "المرحلة", w=8, dv=("list", "1,2,3,4"), align="center")
    dn.add("example", "مثال", w=20)
    dn.add("used", "عدد المستخدم في النظام", "f", 12, f=lambda x: f'=IF({x("prefix")}="","",COUNTIF({REG["Doc_Register"].ws_col("docno")},{x("prefix")}&"-*"))')
    dn.add("next", "الرقم التالي المقترح (للموقع في الخلية)", "f", 24, f=lambda x: f'=IF({x("prefix")}="","",{x("prefix")}&"-[الموقع]-"&TEXT({x("used")}+1,"0000"))')
    dn.samples = [dict(prefix=p, name=n, form=f, sheet=s, phase=ph, example=f"{p}-BR01-0001") for p, n, f, s, ph in [
        ("SAL", "تقرير مبيعات يومي", "نموذج 15", "Sales_Daily", 1), ("PLT", "كشف تسوية منصة/بطاقات", "نموذج 16", "Platform_Settlements", 1),
        ("RCV", "سند قبض", "نموذج 11", "Cash_Vouchers", 1), ("PAY", "سند صرف", "نموذج 12", "Cash_Vouchers", 1),
        ("TRF", "إذن تحويل بين الخزائن/إيداع/سحب", "نموذج 13", "Cash_Vouchers", 1), ("ADJ", "تسوية عجز/زيادة خزينة", "نموذج 30", "Cash_Vouchers", 1),
        ("EXP", "نموذج مصروفات", "نموذج 17", "Expenses", 1), ("SINV", "تسجيل فاتورة مورد/إشعار", "نموذج 20 (مرفق)", "Supplier_Invoices", 1),
        ("SPAY", "دفعة مورد", "نموذج 12", "Supplier_Payments", 1), ("MJE", "قيد يومية/تسوية", "نموذج 24-25", "Manual_JE", 1),
        ("PR", "طلب شراء", "نموذج 1", "PR", 2), ("PO", "أمر شراء", "نموذج 2", "PO", 2), ("GRN", "إذن استلام مخزني", "نموذج 3", "GRN", 2),
        ("ISS", "إذن صرف مخزني", "نموذج 4", "Inventory_Moves", 2), ("STR", "إذن تحويل بين المخازن", "نموذج 5", "Inventory_Moves", 2),
        ("KRQ", "طلب صرف إنتاج", "نموذج 6", "Production", 2), ("KPR", "تقرير إنتاج يومي", "نموذج 7", "Production", 2),
        ("KTR", "إذن تحويل إنتاج لفرع", "نموذج 8", "Production", 2), ("WST", "تقرير هالك وتالف", "نموذج 9", "Inventory_Moves", 2),
        ("CNT", "محضر جرد", "نموذج 10", "Inventory_Moves", 2), ("PRT", "مرتجع مشتريات", "نموذج 28", "Supplier_Invoices", 2),
        ("ADV", "طلب سلفة/عهدة", "نموذج 18", "Advances_Custody", 3), ("CST", "تسوية عهدة", "نموذج 19", "Advances_Custody", 3),
        ("PAYR", "كشف رواتب", "نموذج 22", "Payroll", 3), ("CLS", "محضر إقفال شهر", "نموذج 26", "Period_Close", 1)]]
    return dict(coa=coa, loc=loc, tre=tre, pm=pm, cc=cc, us=us, sup=sup, cus=cus, emp=emp, un=un, uc=uc, ucv=ucv, it=it, dn=dn)


def G(key):
    """مرجع عمود كامل في GL_Journal"""
    return REG["GL_Journal"].ws_col(key)


def define_named_lists(wb):
    n = {
        "LST_LOC": f"Locations!$A${5}:$A${4 + CAP_LOC}",
        "LST_TRE": f"Treasuries!$A$5:$A${4 + CAP_TRE}",
        "LST_PM": f"Payment_Methods!$A$5:$A${4 + CAP_PM}",
        "LST_CC": f"Cost_Centers!$A$5:$A${4 + CAP_CC}",
        "LST_SUP": f"Suppliers!$A$5:$A${4 + CAP_SUP}",
        "LST_CUS": f"Customers!$A$5:$A${4 + CAP_CUS}",
        "LST_EMP": f"Employees!$A$5:$A${4 + CAP_EMP}",
        "LST_ACCT": f"COA!$A$5:$A${4 + CAP_COA}",
        "LST_USER": f"Users!$B$5:$B${4 + CAP_USER}",
        "LST_UNIT": "Units!$A$5:$A$34",
        "LST_ITEM": f"Items!$A$5:$A${4 + CAP_ITEM}",
    }
    for k, v in n.items():
        name_cell(wb, k, v)
        LIST_REFS[k] = v
