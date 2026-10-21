# -*- coding: utf-8 -*-
"""يولّد النماذج بصيغة Word (.docx) قابلة للرفع إلى Google Docs والتعديل."""
import os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import fs_content as C

LOGO = os.path.join(os.path.dirname(__file__), "..", "forms", "assets", "satamoni_logo.jpg")
FONT = "Arial"


def rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi")
    pPr.append(b)
    p.alignment = align


def run(p, text, bold=False, size=10, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    rf = r._r.get_or_add_rPr().find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); r._r.get_or_add_rPr().append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FONT)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)


def table_rtl(t):
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:bidiVisual")
    tblPr.append(b)


def cell_text(cell, text, bold=False, size=9, center=False):
    p = cell.paragraphs[0]
    rtl(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, text, bold, size)


def add_form(doc, f, first):
    wide = len(f["cols"]) > 7
    if not first:
        sec = doc.add_section()
    else:
        sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
    w, h = (29.7, 21.0) if wide else (21.0, 29.7)
    sec.page_width, sec.page_height = Cm(w), Cm(h)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.2))
    # رأس: شعار | عنوان | رقم مسلسل
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"; table_rtl(t)
    c0, c1, c2 = t.rows[0].cells
    c0.paragraphs[0].add_run().add_picture(LOGO, width=Cm(2.6))
    rtl(c0.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    p = c1.paragraphs[0]; rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    run(p, f["title"], True, 18, "0B4F5C")
    p2 = c1.add_paragraph(); rtl(p2, WD_ALIGN_PARAGRAPH.CENTER)
    run(p2, f"مطاعم ستاموني | نموذج رقم {f['no']} | الإصدار 1.0", False, 9)
    p3 = c2.paragraphs[0]; rtl(p3)
    run(p3, "الرقم المسلسل:", False, 9)
    p4 = c2.add_paragraph(); rtl(p4)
    run(p4, f"{f['prefix']} - ______ - ________", True, 11, "B00020")
    p5 = c2.add_paragraph(); rtl(p5)
    run(p5, "(كود الموقع - تسلسل 4 خانات)", False, 7)
    doc.add_paragraph()
    # بيانات ثابتة + خاصة بالنموذج
    fields = ["التاريخ: ____ / ____ / 20____", "الفرع / الموقع (الكود): ____________", "الوردية / الوقت: ____________"] + [f"{x}: ____________________" for x in f["fields"]]
    n = 2
    rows = (len(fields) + n - 1) // n
    ft = doc.add_table(rows=rows, cols=n); table_rtl(ft)
    for i, x in enumerate(fields):
        cell_text(ft.rows[i // n].cells[i % n], x, size=10)
    doc.add_paragraph()
    # جدول
    if f.get("table_custom"):
        t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; table_rtl(t)
        for c, txt in zip(t.rows[0].cells, ["البند", "المبلغ (جنيه)"]):
            cell_text(c, txt, True, 10, True); shade(c, "E6EFF1")
        for a, _ in f["table_custom"]:
            r = t.add_row().cells
            cell_text(r[0], a, size=10); cell_text(r[1], "")
    elif f["cols"]:
        t = doc.add_table(rows=1, cols=len(f["cols"])); t.style = "Table Grid"; table_rtl(t)
        for c, txt in zip(t.rows[0].cells, f["cols"]):
            cell_text(c, txt, True, 9, True); shade(c, "E6EFF1")
        for i in range(f["rows"]):
            r = t.add_row()
            r.height = Cm(0.8)
            if f["cols"][0] == "م":
                cell_text(r.cells[0], str(i + 1), size=9, center=True)
    if f.get("note"):
        p = doc.add_paragraph(); rtl(p); run(p, f["note"], False, 10)
    p = doc.add_paragraph(); rtl(p)
    run(p, "المرفقات: ☐ فاتورة  ☐ إيصال  ☐ عقد  ☐ إذن  ☐ أخرى: ________        ملاحظات: ________________________________", False, 9)
    sigs = f["sigs"]
    st = doc.add_table(rows=1, cols=len(sigs)); st.style = "Table Grid"; table_rtl(st)
    st.rows[0].height = Cm(2.2)
    for c, s in zip(st.rows[0].cells, sigs):
        cell_text(c, s, True, 9)
        q = c.add_paragraph(); rtl(q); run(q, "\nالاسم / التوقيع / التاريخ", False, 7, "777777")
    cp = "، ".join(f"{n_} ({col}) ← {h}" for n_, col, h, w in f["copies"])
    when = "، ".join(f"{n_}: {w}" for n_, col, h, w in f["copies"] if w) or "—"
    p = doc.add_paragraph(); rtl(p)
    run(p, f"عدد النسخ: {len(f['copies'])} — {cp}\nالتسليم للمحاسب: {when}. لا يُعدَّل بعد الترحيل؛ التصحيح بمستند عكسي.", False, 8, "333333")


def build(outdir):
    os.makedirs(os.path.join(outdir, "individual"), exist_ok=True)
    for f in C.FORMS:
        d = Document(); add_form(d, f, True)
        d.save(os.path.join(outdir, "individual", f"{f['no']:02d}_{f['prefix']}_{f['title']}.docx"))
    d = Document()
    for i, f in enumerate(C.FORMS):
        add_form(d, f, i == 0)
    d.save(os.path.join(outdir, "Satamoni_Forms_All.docx"))


if __name__ == "__main__":
    build(sys.argv[1])
