# -*- coding: utf-8 -*-
"""يعيد حساب الملف بـ LibreOffice ويفحص: أخطاء الصيغ، نتائج Test_Cases، المؤشرات الرئيسية."""
import subprocess, sys, os, shutil, collections
import openpyxl


def recalc(path, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path], check=True, capture_output=True, timeout=900)
    return os.path.join(outdir, os.path.basename(path))


def scan_errors(wb, limit=8):
    errs = collections.OrderedDict()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and (v.startswith("#") or v.startswith("Err:")):
                    errs.setdefault(ws.title, []).append((c.coordinate, v))
    return errs


def named(wb, name):
    dn = wb.defined_names[name]
    sheet, ref = dn.attr_text.split("!")
    return wb[sheet.strip("'")][ref.replace("$", "")].value


def report(path, outdir):
    rp = recalc(path, outdir)
    wb = openpyxl.load_workbook(rp, data_only=True)
    errs = scan_errors(wb)
    print("== أخطاء الصيغ ==")
    if not errs:
        print("لا أخطاء")
    for k, v in errs.items():
        print(k, len(v), v[:6])
    return wb, errs


if __name__ == "__main__":
    wb, errs = report(sys.argv[1], sys.argv[2])
    for nm in ["TB_Diff", "GL_TotalDiff", "BS_Diff", "SB_Diff", "CF_Diff", "CF_InternalNet", "CK_Critical", "CK_Review", "IS_NetSales", "IS_GrossProfit", "IS_NetProfit", "TEST_Pass", "TEST_Fail"]:
        try:
            print(nm, named(wb, nm))
        except Exception as e:
            print(nm, "ERR", e)
    ws = wb["Test_Cases"]
    for r in ws.iter_rows(min_row=8, values_only=True):
        if r[0]:
            print(r[0], "|", r[5], "|", r[3], "|", r[4])
