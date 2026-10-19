# -*- coding: utf-8 -*-
"""اختبارات سلبية/تحوّلية: تُدخل مستندات خاطئة عمدًا في نسخة من الملف وتتأكد أن النظام يرفضها وينبّه."""
import datetime as dt, os, sys, tempfile
import openpyxl
import build_system as bs
import verify as V
import fs_common as fc
import b5_samples as b5

d = dt.date


def run(workdir, scale=0.3):
    os.makedirs(workdir, exist_ok=True)
    path, E = bs.build(workdir, demo=True, scale=scale, tag="_mut", ov=dict(sep_snapshot=dict(rev=139670.0, np=-35820.0, dr=1482570.0, cr=1482570.0), sep_closed=True))
    REG = fc.REG
    wb = openpyxl.load_workbook(path)
    nxt = {}
    log = []

    def add(sheet, rec):
        sh = REG[sheet]
        r = nxt.get(sheet, sh.first + len(sh.samples))
        nxt[sheet] = r + 1
        for k, v in rec.items():
            wb[sheet].cell(r, sh.idx(k), v)
        return r

    base = dict(prepared=b5.HQCASH, approved=b5.FM, poster=b5.ACC, status="مرحل", ref="مرفق")
    sc = {}
    # M1 المُعِد = المعتمد
    sc["M1"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0101", date=d(2026, 10, 20), loc="HQ01", cc="CC50", acct=6901, amount=700, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا", approved=b5.HQCASH)), "المُعِد والمعتمد نفس الشخص")
    # M2 أصل ثابت كمصروف
    sc["M2"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0102", date=d(2026, 10, 20), loc="HQ01", cc="CC50", acct=1701, amount=9000, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا")), "غير مسموح للمصروفات")
    # M3 سند صرف بحساب مقابل نقدي
    sc["M3"] = ("Cash_Vouchers", add("Cash_Vouchers", dict(base, doc="PAY-HQ01-0103", date=d(2026, 10, 20), vtype="صرف", tre="CSH-HQ01", cacct=1101, amount=100, payee="س", method="نقدًا", cc="CC00")), "الحساب المقابل نقدي")
    # M4 مجموع وسائل الدفع لا يساوي الصافي
    sc["M4"] = ("Sales_Daily", add("Sales_Daily", dict(base, doc="SAL-BR01-0104", date=d(2026, 10, 20), loc="BR01", orders=10, g1=1000, g2=0, g3=0, g4=0, disc=0, ret=0, m1=900, prepared=b5.CSH["BR01"], approved=b5.MGR["BR01"])), "مجموع وسائل الدفع لا يساوي صافي المبيعات")
    # M5 تقرير مبيعات مكرر لنفس الفرع واليوم
    sc["M5"] = ("Sales_Daily", add("Sales_Daily", dict(base, doc="SAL-BR01-0105", date=d(2026, 10, 5), loc="BR01", orders=10, g1=1000, disc=0, ret=0, m1=1000, prepared=b5.CSH["BR01"], approved=b5.MGR["BR01"])), "تقرير مبيعات مكرر")
    # M6 رقم مستند مكرر عبر الأوراق
    sc["M6"] = ("Expenses", add("Expenses", dict(base, doc="SINV-WH01-0001", date=d(2026, 10, 20), loc="HQ01", cc="CC50", acct=6901, amount=50, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا")), "رقم مستند مكرر")
    # M7 مستند جديد (مسودة) في فترة مقفلة
    sc["M7"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0107", date=d(2026, 9, 15), loc="HQ01", cc="CC50", acct=6901, amount=50, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا", status="مسودة")), "الفترة مقفلة")
    # M8 مستند مرحل بأثر رجعي في فترة مقفلة -> تتغير لقطة الإقفال
    sc["M8"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0108", date=d(2026, 9, 15), loc="HQ01", cc="CC50", acct=6901, amount=1234, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا")), "")
    # M9 مورد غير معرف
    sc["M9"] = ("Supplier_Invoices", add("Supplier_Invoices", dict(base, doc="SINV-WH01-0109", sinv="X-1", sup="SUP999", inv_date=d(2026, 10, 20), date=d(2026, 10, 20), itype="فاتورة", dacct=1401, loc="WH01", cc="CC70", net=100, prepared=b5.PURCH)), "المورد غير معرف")
    # M10 المرحِّل = المُعِد
    sc["M10"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0110", date=d(2026, 10, 20), loc="HQ01", cc="CC50", acct=6901, amount=50, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا", poster=b5.HQCASH)), "المرحِّل يجب أن يختلف")
    # M11 فرع جديد بدون خزينة
    ws = wb["Locations"]
    sh = REG["Locations"]
    r = sh.first + len(sh.samples)
    for k, v in dict(code="BR05", name="فرع تجريبي بلا خزينة", type="فرع", active="نعم").items():
        ws.cell(r, sh.idx(k), v)
    sc["M11"] = ("Locations", r, "الفرع يحتاج خزينة افتراضية")
    # M12 سند تحويل لنفس الخزينة
    sc["M12"] = ("Cash_Vouchers", add("Cash_Vouchers", dict(base, doc="TRF-HQ01-0112", date=d(2026, 10, 20), vtype="تحويل بين خزائن", tre="CSH-HQ01", tre2="CSH-HQ01", amount=100, method="نقدًا")), "التحويل إلى نفس الخزينة")
    # M13 مبلغ سالب (يتجاوز التحقق عبر لصق) في المصروفات
    sc["M13"] = ("Expenses", add("Expenses", dict(base, doc="EXP-HQ01-0113", date=d(2026, 10, 20), loc="HQ01", cc="CC50", acct=6901, amount=-50, mode="فوري", tre="CSH-HQ01", payee="مورد", method="نقدًا")), "المبلغ يجب أن يكون أكبر من صفر")
    # M14 قيد يدوي متوازن صحيح (يجب أن يُقبل)
    for acct, dr, cr in [(6990, 200, None), (2201, None, 200)]:
        add("Manual_JE", dict(doc="MJE-HQ01-0114", date=d(2026, 10, 21), jtype="تسوية", acct=acct, loc="HQ01", cc="CC50", desc="قيد سليم", debit=dr, credit=cr, status="مرحل",
                              prepared=b5.ASSIST, approved=b5.FM, poster=b5.ACC, ref="مرفق"))
    sc["M14"] = ("Manual_JE", nxt["Manual_JE"] - 1, "")
    for acct, dr, cr in [(4101, 1000, None), (3501, None, 1000)]:
        add("Manual_JE", dict(doc="MJE-HQ01-0115", date=d(2026, 10, 31), jtype="إقفال", acct=acct, loc="HQ01", cc="CC50", desc="قيد إقفال تجريبي", debit=dr, credit=cr, status="مرحل",
                              prepared=b5.ASSIST, approved=b5.FM, poster=b5.ACC, ref="مرفق"))
    sc["M15"] = ("Manual_JE", nxt["Manual_JE"] - 1, "")
    out = os.path.join(workdir, "mut_input.xlsx")
    wb.save(out)
    rp = V.recalc(out, os.path.join(workdir, "mut_calc"))
    wv = openpyxl.load_workbook(rp, data_only=True)

    def gl_eff(doc):
        sh = REG["GL_Journal"]
        ws = wv["GL_Journal"]
        tot = 0
        n = 0
        for r in range(sh.first, sh.last + 1):
            if ws.cell(r, sh.idx("doc")).value == doc:
                n += 1
                tot += float(ws.cell(r, sh.idx("effdr")).value or 0)
        return n, tot

    res = []
    for name, (sheet, row, expect) in sc.items():
        sh = REG[sheet]
        key = "chk" if sheet != "Cash_Close" else "flag"
        chk = wv[sheet].cell(row, sh.idx(key)).value
        ok = (expect in (chk or "")) if expect else (not chk)
        res.append((name, sheet, row, expect or "(مقبول بلا أخطاء)", chk, ok))
    # تأثير على الدفاتر
    n, eff = gl_eff("EXP-HQ01-0101")
    res.append(("M1-GL", "GL_Journal", "-", "أثر فعلي 0 للمستند المرفوض", eff, eff == 0))
    n, eff = gl_eff("EXP-HQ01-0102")
    res.append(("M2-GL", "GL_Journal", "-", "أثر فعلي 0", eff, eff == 0))
    n, eff = gl_eff("MJE-HQ01-0114")
    res.append(("M14-GL", "GL_Journal", "-", "قيد سليم مقبول: مدين 200", eff, abs(eff - 200) < 0.001))
    # مكرر عبر الأوراق: الأصل أيضًا موسوم
    sh = REG["Supplier_Invoices"]
    chk0 = wv["Supplier_Invoices"].cell(sh.first, sh.idx("chk")).value
    res.append(("M6-orig", "Supplier_Invoices", sh.first, "المستند الأصلي SINV-WH01-0001 موسوم أيضًا (التكرار لا يمرّ بصمت)", chk0, bool(chk0) and "مكرر" in chk0))
    # M8: لقطة الإقفال تكشف
    pc = REG["Period_Close"]
    wsp = wv["Period_Close"]
    chg = None
    for r in range(pc.first, pc.last + 1):
        if wsp.cell(r, pc.idx("period")).value == 202609:
            chg = wsp.cell(r, pc.idx("chg")).value
    res.append(("M8-close", "Period_Close", "-", "تغيّر بعد الإقفال", chg, bool(chg)))
    crit = V.named(wv, "CK_Critical")
    tb = V.named(wv, "TB_Diff")
    res.append(("TB-after", "Trial_Balance", "-", "الميزان متوازن رغم المستندات المرفوضة", tb, abs(tb) < 0.01))
    gs = V.named(wv, "IS_GrossSales")
    res.append(("M15-closing", "Income_Statement", "-", f"قيد الإقفال لا يُنقص إجمالي مبيعات قائمة الدخل ({E['sales_oct_gross']})", gs, abs(gs - E["sales_oct_gross"]) < 0.01))
    res.append(("M15-BS", "Balance_Sheet", "-", "الميزانية متوازنة بعد قيد الإقفال", V.named(wv, "BS_Diff"), abs(V.named(wv, "BS_Diff")) < 0.01))
    errs = V.scan_errors(wv)
    res.append(("no-formula-errors", "كل الأوراق", "-", "لا أخطاء #", errs, not errs))
    return res, E


def run_vat(workdir, scale=0.3):
    os.makedirs(workdir, exist_ok=True)
    path, E = bs.build(workdir, demo=True, scale=scale, tag="_vat", ov=dict(sep_snapshot=dict(rev=139670.0, np=-35820.0, dr=1482570.0, cr=1482570.0), sep_closed=True))
    wb = openpyxl.load_workbook(path)
    wb["Tax_Settings"]["C5"].value = "نعم"
    out = os.path.join(workdir, "vat_input.xlsx")
    wb.save(out)
    rp = V.recalc(out, os.path.join(workdir, "vat_calc"))
    wv = openpyxl.load_workbook(rp, data_only=True)
    f = 1.14
    exp_vat = 0.0
    for r in E["sales_rows"].values():
        g = [r["g1"], r["g2"], r["g3"], r["g4"]]
        rev = [round(x / f + 1e-9, 2) for x in g]
        dex = round(r["disc"] / f + 1e-9, 2)
        rex = round(r["ret"] / f + 1e-9, 2)
        net = r["_net"]
        exp_vat += round(net + dex + rex - sum(rev), 2)
    sh = fc.REG["GL_Journal"]
    ws = wv["GL_Journal"]
    got = 0.0
    for rr in range(sh.first, sh.last + 1):
        if ws.cell(rr, sh.idx("acct")).value == 2301:
            got += float(ws.cell(rr, sh.idx("effcr")).value or 0)
    tb = V.named(wv, "TB_Diff")
    bsd = V.named(wv, "BS_Diff")
    return dict(expected_vat=round(exp_vat, 2), got=round(got, 2), tb=tb, bs=bsd, ok=abs(exp_vat - got) < 0.05 and abs(tb) < 0.01 and abs(bsd) < 0.01)


if __name__ == "__main__":
    wd = sys.argv[1]
    os.makedirs(wd, exist_ok=True)
    res, E = run(wd)
    for r in res:
        print(r)
    print("FAIL" if not all(x[-1] for x in res) else "ALL OK")
    print(run_vat(wd))
