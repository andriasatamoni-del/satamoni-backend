# -*- coding: utf-8 -*-
"""قائمة الإقفال، إقفال الفترات، لوحة التحكم، الاختبارات، وبقية الأوراق."""
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from fs_common import *
from b1_masters import G, band, mk_text_sheet
from b3_ledger import put, selector, sumifs, netdr
import fs_data as D
import fs_content as C

d = dt.date

CK_ITEMS_MANUAL = [
    ("M01", "مطابقة الخزائن", "عدّ فعلي لكل الخزائن وتوقيع كشوف الإقفال اليومية للشهر كله (Cash_Close)"),
    ("M02", "مطابقة البنوك", "مطابقة أرصدة حسابات البنوك مع كشوف الحساب البنكية"),
    ("M03", "مطابقة مستحقات منصات التوصيل", "مطابقة الرصيد 1303/1304 مع كشوف المنصات وتسوية الفروق في 1309"),
    ("M04", "مطابقة الموردين", "مطابقة أرصدة الموردين الرئيسيين بكشوف حساباتهم"),
    ("M05", "جرد المخزون واعتماد فروقه", "جرد فعلي في المخزن والمطبخ والفروع واعتماد محضر الفروق (المرحلة 2: آلي)"),
    ("M06", "مراجعة الهالك", "مراجعة تقارير الهالك الطبيعي/غير الطبيعي واعتمادها"),
    ("M07", "مراجعة التحويلات المفتوحة", "تحويلات نقدية أو مخزنية لم يوقع المستلم عليها"),
    ("M08", "مراجعة الرواتب والسلف والعهد", "اعتماد كشف الرواتب وتسوية العهد (المرحلة 3: آلي)"),
    ("M09", "إثبات المصروفات المستحقة", "كهرباء/غاز/إنترنت/خدمات لم تصل فواتيرها بعد"),
    ("M10", "مراجعة المصروفات المدفوعة مقدمًا", "إطفاء حصة الشهر من الإيجار والتأمين المدفوع مقدمًا"),
    ("M11", "احتساب الإهلاك", "إثبات إهلاك الشهر للأصول الثابتة"),
    ("M12", "عمولة المنصات المقدّرة غير المسواة", "قيد استحقاق بالنسبة التقديرية عن مبيعات المنصات غير المسواة"),
    ("M13", "مراجعة ميزان المراجعة والحسابات المعلقة", "مراجعة بشرية للأرصدة غير المعتادة"),
    ("M14", "اعتماد قائمة الدخل والميزانية", "توقيع المدير المالي على نتائج الفترة"),
]


def define_close_confirm():
    cc = Sheet("Close_Confirm", "تأكيدات الإقفال اليدوية (Close Confirmations)",
               "لكل فترة وبند يدوي من Close_Checklist: سجّل التأكيد واسم المؤكِّد وتاريخه. (لا ينطبق) تتطلب كتابة السبب في الملاحظات.",
               cap=300, tab="C00000")
    cc.add("period", "رقم الفترة (yyyymm)", w=12, dv=("int",), align="center")
    cc.add("item", "رقم البند", w=9, dv=("list", ",".join(x[0] for x in CK_ITEMS_MANUAL)), align="center")
    cc.add("value", "التأكيد", w=14, dv=("list", "نعم,لا ينطبق"))
    cc.add("by", "المؤكِّد", w=24, dv=("list", "=LST_USER"))
    cc.add("on", "تاريخ التأكيد", w=12, fmt="yyyy-mm-dd", dv=("date",))
    cc.add("notes", "ملاحظات / سبب عدم الانطباق", w=40)
    cc.add("chk", "فحص", "f", 36, f=lambda x: f'=IF(COUNTA(${cc.input_span()[0]}{x.r}:${cc.input_span()[1]}{x.r})=0,"",' + chain([
        (f'OR({x("period")}="",{x("item")}="")', "الفترة أو البند ناقص"),
        (f'COUNTIFS({x.col("period")},{x("period")},{x.col("item")},{x("item")})>1', "تأكيد مكرر لنفس الفترة والبند"),
        (f'{x("by")}=""', "اسم المؤكِّد ناقص"),
        (f'NOT(ISNUMBER({x("on")}))', "تاريخ التأكيد ناقص"),
        (f'AND({x("value")}="لا ينطبق",{x("notes")}="")', "اكتب سبب عدم الانطباق"),
    ]) + ')')
    return cc


def define_period_close():
    pc = Sheet("Period_Close", "إقفال الفترات (Period Close Log)",
               "سجل إقفال الشهر: اسم المعتمد وتاريخ الإقفال ولقطة النتائج وقت الإقفال (انسخ القيم من الأعمدة الحية والصقها كقيم). يُمنع إدخال (مقفلة) ما لم تكن الفترة جاهزة (قائمة الإقفال بلا أخطاء حرجة).",
               cap=60, tab="C00000")
    pc.add("period", "رقم الفترة (yyyymm)", w=12, dv=("int",), align="center")
    pc.add("status", "الحالة", w=11, dv=("custom", '=OR({c}<>"مقفلة",$A5=$H$3)', "لا يمكن الإقفال: الفترة غير جاهزة (ضع Settings على هذا الشهر الكامل وعالج الأخطاء الحرجة أولًا)."), align="center")
    pc.add("by", "المُقفِل", w=24, dv=("list", "=LST_USER"))
    pc.add("on", "تاريخ الإقفال", w=12, fmt="yyyy-mm-dd", dv=("date",))
    pc.add("by2", "المعتمد الثاني", w=24, dv=("list", "=LST_USER"))
    pc.add("s_rev", "لقطة: صافي الإيرادات", w=14, fmt="#,##0.00", dv=("dec",))
    pc.add("s_np", "لقطة: صافي الربح", w=14, fmt="#,##0.00")
    pc.add("s_dr", "لقطة: إجمالي المدين", w=14, fmt="#,##0.00", dv=("dec",))
    pc.add("s_cr", "لقطة: إجمالي الدائن", w=14, fmt="#,##0.00", dv=("dec",))
    pc.add("notes", "ملاحظات", w=26)
    frm = lambda x: f'DATE(INT({x("period")}/100),MOD({x("period")},100),1)'
    to = lambda x: f'EOMONTH({frm(x)},0)'
    def rev(x, a, b, side):
        cr = [("acct", f'">={a}"'), ("acct", f'"<={b}"'), ("date", f'">="&{frm(x)}'), ("date", f'"<="&{to(x)}')]
        if a >= 4000:  # حسابات النتيجة تستثني قيود الإقفال السنوي
            cr.append(("clos", "0"))
        return sumifs(side, cr)
    pc.add("l_rev", "حي: صافي الإيرادات", "f", 14, fmt="#,##0.00", f=lambda x: f'=IF({x("period")}="","",{rev(x,4000,4999,"effcr")}-{rev(x,4000,4999,"effdr")})')
    pc.add("l_np", "حي: صافي الربح", "f", 14, fmt="#,##0.00", f=lambda x: f'=IF({x("period")}="","",{x("l_rev")}-({rev(x,5000,6999,"effdr")}-{rev(x,5000,6999,"effcr")}))')
    pc.add("l_dr", "حي: إجمالي المدين", "f", 14, fmt="#,##0.00", f=lambda x: f'=IF({x("period")}="","",{rev(x,1000,6999,"effdr")})')
    pc.add("l_cr", "حي: إجمالي الدائن", "f", 14, fmt="#,##0.00", f=lambda x: f'=IF({x("period")}="","",{rev(x,1000,6999,"effcr")})')
    pc.add("chg", "تغيّر بعد الإقفال؟", "f", 36, f=lambda x: (
        f'=IF(AND({x("status")}="مقفلة",{x("period")}<>""),IF(OR(ABS({x("l_rev")}-N({x("s_rev")}))>S_Tol,ABS({x("l_np")}-N({x("s_np")}))>S_Tol,ABS({x("l_dr")}-N({x("s_dr")}))>S_Tol,ABS({x("l_cr")}-N({x("s_cr")}))>S_Tol),"تغيّر بعد الإقفال: قيد مرحّل بأثر رجعي","") ,"")'))
    pc.add("flag", "فحص", "f", 36, f=lambda x: f'=IF({x("period")}="","",' + chain([
        (f'COUNTIF({x.col("period")},{x("period")})>1', "فترة مكررة"),
        (f'AND({x("status")}="مقفلة",{x("by")}=""),', "") if False else (f'AND({x("status")}="مقفلة",{x("by")}="")', "اسم المُقفِل ناقص"),
        (f'AND({x("status")}="مقفلة",NOT(ISNUMBER({x("on")})))', "تاريخ الإقفال ناقص"),
        (f'AND({x("status")}="مقفلة",{x("by2")}="")', "المعتمد الثاني ناقص"),
        (f'AND({x("status")}="مقفلة",{x("by2")}<>"",{x("by2")}={x("by")})', "المُقفِل والمعتمد الثاني نفس الشخص"),
    ]) + ')')

    def top(ws):
        put(ws, "G3", "الفترة الجاهزة للإقفال الآن:", bold=True, fill=C_HDR_BAND)
        put(ws, "H3", '=IF(AND(CK_Critical=0,DAY(S_From)=1,S_To=EOMONTH(S_From,0)),S_PeriodID,0)', fill=C_F, name=None)
        ws["H3"].number_format = "0"
    pc.extra.append(top)
    return pc


def snapshot_rows(ov):
    """عينات إقفال سبتمبر (القيم من التشغيل الآلي الأول) وفتح أكتوبر."""
    s = ov.get("sep_snapshot", dict(rev=0, np=0, dr=0, cr=0))
    return [dict(period=202609, status="مقفلة" if ov.get("sep_closed") else "مفتوحة", by="المحاسب (تجريبي)", on=d(2026, 10, 2), by2="المدير المالي (تجريبي)",
                 s_rev=s["rev"], s_np=s["np"], s_dr=s["dr"], s_cr=s["cr"], notes="إقفال تجريبي لشهر سبتمبر"),
            dict(period=202610, status="مفتوحة", notes="الشهر الجاري")]


# ============================================================ Close checklist
def build_close_checklist(wb, pm_sheet):
    ws = wb.create_sheet("Close_Checklist")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    ws.cell(1, 1, "قائمة الإقفال اليومي والشهري (Close Checklist) - أخطاء تمنع الإقفال وملاحظات للمراجعة").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 9):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "تعمل على فترة Settings. البنود (آلية) تحسب نفسها؛ البنود (يدوية) تؤكدها من ورقة Close_Confirm. الإقفال غير مسموح إلا بلا أي بند حرج مفتوح.").font = Font(name=FONT, size=9, italic=True)
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 9
    ws.column_dimensions["G"].width = 9
    ws.column_dimensions["H"].width = 60
    heads = ["رقم", "المجال", "الفحص", "القيمة", "الحالة", "النوع", "الأهمية", "كيف تعالجه"]
    # summary rows 4-7
    put(ws, "A4", "الفترة", bold=True, fill=C_HDR_BAND)
    put(ws, "B4", '=YEAR(S_From)&"-"&TEXT(MONTH(S_From),"00")&"-"&TEXT(DAY(S_From),"00")&" → "&YEAR(S_To)&"-"&TEXT(MONTH(S_To),"00")&"-"&TEXT(DAY(S_To),"00")', fill=C_F)
    put(ws, "A5", "بنود حرجة مفتوحة (تمنع الإقفال)", bold=True, fill=C_HDR_BAND)
    put(ws, "A6", "بنود للمراجعة", bold=True, fill=C_HDR_BAND)
    put(ws, "A7", "جاهزية الإقفال", bold=True, fill=C_HDR_BAND)
    H = 9
    for i, h in enumerate(heads, start=1):
        c = ws.cell(H, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_F_H)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    R = lambda sh, k: REG[sh].ws_col(k)
    inp = lambda sh: (R(sh, "date"), R(sh, "status"))
    D_FROM, D_TO = '">="&S_From', '"<="&S_To'

    def cnt_sources(key_extra, status_cond):
        parts = []
        for sh, dk in [("Sales_Daily", "date"), ("Platform_Settlements", "date"), ("Cash_Vouchers", "date"), ("Expenses", "date"),
                       ("Supplier_Invoices", "date"), ("Supplier_Payments", "date"), ("Manual_JE", "date")]:
            parts.append(f'COUNTIFS({R(sh,"used")},1,{R(sh,dk)},{D_FROM},{R(sh,dk)},{D_TO},{status_cond(sh)}{key_extra(sh)})')
        return "+".join(parts)

    nodraft = lambda sh: f'{R(sh,"status")},"<>"&ST_DRAFT,{R(sh,"status")},"<>"&ST_CANC,{R(sh,"status")},"<>"&ST_REV'
    errs = cnt_sources(lambda sh: f',{R(sh,"bad")},1', nodraft)
    drafts = cnt_sources(lambda sh: '', lambda sh: f'{R(sh,"status")},ST_DRAFT')
    apprs = cnt_sources(lambda sh: '', lambda sh: f'{R(sh,"status")},ST_APPR')
    blank_status = cnt_sources(lambda sh: '', lambda sh: f'{R(sh,"status")},""')
    noref = "+".join([f'COUNTIFS({R(sh,"used")},1,{R(sh,"date")},{D_FROM},{R(sh,"date")},{D_TO},{R(sh,"status")},ST_POST,{R(sh,"ref")},"")' for sh in
                      ["Sales_Daily", "Platform_Settlements", "Cash_Vouchers", "Expenses", "Supplier_Invoices", "Supplier_Payments"]])
    dupsi = f'COUNTIFS({R("Supplier_Invoices","dupf")},1,{R("Supplier_Invoices","date")},{D_FROM},{R("Supplier_Invoices","date")},{D_TO})'
    gl_blocked = f'COUNTIFS({G("act")},1,{G("eff")},0,{G("status")},ST_POST,{G("date")},{D_FROM},{G("date")},{D_TO})'
    sd = REG["Sales_Daily"]
    master_flags = "+".join([f'COUNTIF({REG[s].rng("chk")},"?*")' for s in ["COA", "Locations", "Treasuries", "Payment_Methods", "Cost_Centers", "Suppliers", "Customers", "Employees", "Users"]])
    nonsales_rev = "+".join([f'COUNTIFS({G("eff")},1,{G("acct")},">=4101",{G("acct")},"<=4104",{G("stype")},"{t}")' for t in ["CV", "EXP", "SINV", "SPAY", "PLT"]])
    bal_of = lambda acct: netdr([("acct", acct), ("date", '"<="&S_To')])
    items = [
        # key, group, text, value formula, severity, how, kind
        ("errs", "المستندات", "مستندات مرحلة/معتمدة بها أخطاء فحص في الفترة", errs, "حرج", "افتح الورقة المصدرية وراجع عمود (فحص المستند) وصحّح أو اعكس المستند"),
        ("drafts", "المستندات", "مستندات (مسودة) في الفترة لم تُعتمد", drafts, "حرج", "اعتمدها وسجلها أو اجعلها (ملغى) بتوقيع المسؤول"),
        ("apprs", "المستندات", "مستندات (معتمد) لم يُرحّلها المحاسب بعد", apprs, "حرج", "ترحيل المستندات بتغيير الحالة إلى (مرحل) بواسطة المحاسب"),
        ("blank", "المستندات", "مستندات بلا حالة", blank_status, "حرج", "حدد حالة المستند"),
        ("docdup", "المستندات", "أرقام مستندات مكررة في السجل المركزي", f'COUNTIF({R("Doc_Register","dup")},"مكرر")', "حرج", "راجع Doc_Register وصحّح الرقم المكرر بموافقة المحاسب"),
        ("invdup", "الموردون", "فواتير موردين مكررة (نفس المورد ونفس الرقم)", dupsi, "حرج", "ألغِ المستند المكرر (ملغى) أو تحقق إن كان فاتورة مختلفة"),
        ("unbal", "اليومية", "قيود يدوية غير متوازنة", f'SUMIFS({R("Manual_JE","jebal")},{R("Manual_JE","date")},{D_FROM},{R("Manual_JE","date")},{D_TO})', "حرج", "صحّح الأسطر حتى يتساوى المدين والدائن، ولا يُرحَّل غير المتوازن"),
        ("glblocked", "اليومية", "أسطر مرحّلة لم تؤثر على الدفاتر (محجوبة بخطأ)", gl_blocked, "حرج", "سببها في GL_Journal عمود (سبب عدم الترحيل)"),
        ("tb", "الأرصدة", "فرق ميزان المراجعة (مدين - دائن)", "ABS(TB_Diff)", "حرج", "راجع GL_Journal؛ أي فرق يعني حساب خارج الدليل أو قيدًا محجوبًا"),
        ("gl", "الأرصدة", "فرق إجمالي الأستاذ العام (مدين - دائن)", "ABS(GL_TotalDiff)", "حرج", "نفس البند السابق"),
        ("bs", "الأرصدة", "فرق الميزانية (الأصول - الالتزامات - الملكية)", "ABS(BS_Diff)", "حرج", "الميزانية تقرأ من الأستاذ؛ الفرق يعني خطأ في المصدر"),
        ("sbdiff", "الأرصدة", "فرق أرصدة الموردين (دفتر فرعي - حساب 2101)", "ABS(SB_Diff)", "حرج", "حركات على 2101 بدون كود مورد صحيح في Manual_JE"),
        ("cbdiff", "الأرصدة", "فرق أرصدة العملاء (دفتر فرعي - حساب 1301)", "ABS(CB_Diff)", "حرج", "حركات على 1301 بدون كود عميل"),
        ("cfdiff", "الأرصدة", "فرق التدفقات النقدية مع أرصدة النقدية", "ABS(CF_Diff)", "حرج", "تحقق من تصنيف التدفق للأسطر النقدية"),
        ("cfint", "الخزينة", "تحويلات نقدية داخلية صافيها لا يساوي صفرًا", "ABS(CF_InternalNet)", "حرج", "سند تحويل غير مكتمل أو بسطر واحد"),
        ("reconflag", "المبيعات", "صفوف مطابقة مبيعات بها فرق غير موضح أو ناقصة", f'COUNTIFS({R("Daily_Recon","flag")},"?*",{R("Daily_Recon","date")},{D_FROM},{R("Daily_Recon","date")},{D_TO})', "حرج", "اكتب سبب الفرق والمسؤول وحالة المعالجة في Daily_Recon"),
        ("reconopen", "المبيعات", "فروق مطابقة مبيعات ما زالت مفتوحة", f'COUNTIFS({R("Daily_Recon","open")},1,{R("Daily_Recon","date")},{D_FROM},{R("Daily_Recon","date")},{D_TO})', "مراجعة", "سوّها أو اعتمدها بموافقة الإدارة"),
        ("norecon", "المبيعات", "أيام مبيعات مرحلة بدون مطابقة يومية", f'COUNTIFS({R("Sales_Daily","recon")},"لا",{R("Sales_Daily","date")},{D_FROM},{R("Sales_Daily","date")},{D_TO},{R("Sales_Daily","status")},ST_POST)', "مراجعة", "سجل صف مطابقة لليوم والفرع في Daily_Recon"),
        ("cashflag", "الخزينة", "صفوف إقفال خزينة بها أخطاء أو فروق بلا تسوية", f'COUNTIFS({R("Cash_Close","flag")},"?*",{R("Cash_Close","date")},{D_FROM},{R("Cash_Close","date")},{D_TO})', "حرج", "سجل سند عجز/زيادة (نموذج 30) واربطه برقمه في Cash_Close"),
        ("cashopen", "الخزينة", "فروق خزينة مفتوحة بعد التسوية", f'COUNTIFS({R("Cash_Close","open")},1,{R("Cash_Close","date")},{D_FROM},{R("Cash_Close","date")},{D_TO})', "حرج", "الفرق المتبقي لا يساوي صفرًا"),
        ("negcash", "الخزينة", "خزائن/بنوك برصيد سالب", "TR_Negative", "مراجعة", "رصيد سالب غير منطقي؛ غالبًا مستند ناقص أو تاريخ خاطئ"),
        ("susp", "الحسابات المعلقة", "رصيد الحساب المعلق 2901 غير صفر", f'ABS(-{bal_of("ACC_SUSPENSE")})', "حرج", "صفّر الحساب المعلق بقيد تسوية"),
        ("pltdiff", "الحسابات المعلقة", "فروق تسوية المنصات 1309 غير صفرية", f'ABS({bal_of("ACC_PLT_DIFF")})', "مراجعة", "تسوية الفرق مع المنصة أو قيد شطب معتمد"),
        ("stale", "المنصات", "منصات/بطاقات بمستحقات ومضى على آخر تسوية أكثر من الحد", f'COUNTIFS({REG["Payment_Methods"].ws_col("stale")},1)', "مراجعة", "اطلب كشف التسوية من المنصة"),
        ("cogs", "التكلفة", "قيد تسوية تكلفة المبيعات الدوري غير موجود لهذه الفترة", f'IF(S_CogsMethod="PERIODIC",IF(COUNTIFS({R("Manual_JE","jtype")},"تسوية تكلفة المبيعات",{R("Manual_JE","ok")},1,{R("Manual_JE","date")},{D_FROM},{R("Manual_JE","date")},{D_TO})=0,1,0),0)', "حرج", "سجل قيد المثال 17/26 بعد الجرد، وإلا يظهر الربح مبالغًا فيه"),
        ("revsrc", "الإيراد", "إيراد من مصدر غير المبيعات (مثلًا سند/مصروف)", nonsales_rev, "حرج", "الإيراد لا يسجل إلا من Sales_Daily؛ التحويلات الداخلية ليست إيرادًا"),
        ("master", "البيانات الأساسية", "أخطاء في البيانات الأساسية (حسابات/مواقع/خزائن/موردون...)", master_flags, "حرج", "افتح الورقة المعنية وراجع عمود (فحص)"),
        ("fullmonth", "الإقفال", "الفترة المختارة ليست شهرًا كاملًا", 'IF(AND(DAY(S_From)=1,S_To=EOMONTH(S_From,0)),0,1)', "حرج", "اضبط Settings على أول وآخر يوم في الشهر قبل الإقفال"),
        ("closedchg", "الإقفال", "فترات مقفلة تغيّرت أرقامها بعد الإقفال", f'COUNTIF({REG["Period_Close"].rng("chg")},"?*")', "حرج", "راجع القيود المرحلة بأثر رجعي في الفترة المقفلة"),
        ("closeflag", "الإقفال", "أخطاء في سجل الإقفال/التأكيدات", f'COUNTIF({REG["Period_Close"].rng("flag")},"?*")+COUNTIF({REG["Close_Confirm"].rng("chk")},"?*")', "حرج", "راجع Period_Close وClose_Confirm"),
        ("noref", "المستندات", "مستندات مرحلة بدون مرجع/مرفق ورقي", noref, "مراجعة", "اكتب رقم المرفق أو مرجع الدوسيه الورقي"),
        ("overdue", "الموردون", "مبالغ موردين متأخرة السداد", "ROUND(SB_Overdue,2)", "مراجعة", "راجع Supplier_Balances وجدول السداد"),
        ("sbneg", "الموردون", "موردون برصيد مدين (دفع زيادة)", "SB_Negative", "مراجعة", "تحقق من الدفعات أو المرتجعات"),
    ]
    r = H + 1
    keyrow = {}
    for n, (key, grp, text, fx, sev, how) in enumerate(items, start=1):
        ws.cell(r, 1, f"A{n:02d}").border = BORDER
        ws.cell(r, 2, grp).border = BORDER
        ws.cell(r, 3, text).border = BORDER
        c = ws.cell(r, 4, "=" + fx)
        c.number_format = "#,##0.##"
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor=C_F)
        e = ws.cell(r, 5, f'=IF(ROUND(D{r},2)=0,"✔ سليم",IF(G{r}="حرج","✖ خطأ يمنع الإقفال","⚠ مراجعة"))')
        e.border = BORDER
        ws.cell(r, 6, "آلي").border = BORDER
        ws.cell(r, 7, sev).border = BORDER
        h = ws.cell(r, 8, how)
        h.border = BORDER
        h.alignment = Alignment(wrap_text=True)
        keyrow[key] = r
        r += 1
    for n, (mid, grp, text) in enumerate(CK_ITEMS_MANUAL, start=1):
        ws.cell(r, 1, mid).border = BORDER
        ws.cell(r, 2, "مراجعة بشرية").border = BORDER
        ws.cell(r, 3, f"{grp}: {text}").border = BORDER
        cc = REG["Close_Confirm"]
        yes = f'COUNTIFS({R("Close_Confirm","period")},S_PeriodID,{R("Close_Confirm","item")},A{r},{R("Close_Confirm","value")},"نعم")'
        na = f'COUNTIFS({R("Close_Confirm","period")},S_PeriodID,{R("Close_Confirm","item")},A{r},{R("Close_Confirm","value")},"لا ينطبق",{R("Close_Confirm","notes")},"?*")'
        c = ws.cell(r, 4, f'=IF({yes}>0,0,IF({na}>0,0,1))')
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor=C_F)
        e = ws.cell(r, 5, f'=IF(D{r}=0,"✔ تم","✖ بانتظار التأكيد")')
        e.border = BORDER
        ws.cell(r, 6, "يدوي").border = BORDER
        ws.cell(r, 7, "حرج").border = BORDER
        h = ws.cell(r, 8, "سجل التأكيد في Close_Confirm (الفترة + رقم البند + اسمك + التاريخ)")
        h.border = BORDER
        keyrow[mid] = r
        r += 1
    last = r - 1
    put(ws, "B5", f'=COUNTIF(E{H + 1}:E{last},"✖*")', bold=True, name="CK_Critical", wb=wb)
    put(ws, "B6", f'=COUNTIF(E{H + 1}:E{last},"⚠*")', bold=True, name="CK_Review", wb=wb)
    put(ws, "B7", '=IF(CK_Critical=0,"✔ جاهزة للإقفال","✖ غير جاهزة: عالج البنود الحرجة")', bold=True)
    for k, nm in [("drafts", "CK_Drafts"), ("apprs", "CK_Approved"), ("docdup", "CK_DocDup"), ("invdup", "CK_InvDup"), ("reconopen", "CK_ReconOpen"), ("norecon", "CK_NoRecon"),
                  ("cashopen", "CK_CashOpen"), ("stale", "CK_Stale"), ("susp", "CK_Susp"), ("errs", "CK_Errs"), ("noref", "CK_NoRef"), ("cogs", "CK_CogsMissing"), ("unbal", "CK_Unbal"), ("negcash", "CK_NegCash")]:
        name_cell(wb, nm, f"Close_Checklist!$D${keyrow[k]}")
    name_cell(wb, "CK_ManualOpen", f"Close_Checklist!$D${keyrow['M01']}")  # placeholder (not used)
    ws.conditional_formatting.add(f"E{H + 1}:E{last}", FormulaRule(formula=[f'LEFT(E{H + 1},1)="✖"'], fill=PatternFill("solid", bgColor=C_RED, fgColor=C_RED)))
    ws.conditional_formatting.add(f"E{H + 1}:E{last}", FormulaRule(formula=[f'LEFT(E{H + 1},1)="⚠"'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    ws.conditional_formatting.add(f"E{H + 1}:E{last}", FormulaRule(formula=[f'LEFT(E{H + 1},1)="✔"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    ws.conditional_formatting.add("B7", FormulaRule(formula=['LEFT(B7,1)="✖"'], fill=PatternFill("solid", bgColor=C_RED, fgColor=C_RED)))
    ws.conditional_formatting.add("B7", FormulaRule(formula=['LEFT(B7,1)="✔"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    ws.freeze_panes = ws.cell(H + 1, 1)
    ws.auto_filter.ref = f"A{H}:H{last}"
    protect_ws(ws)
    return ws, keyrow


# ============================================================ Dashboard
def build_dashboard(wb, keymap, H_is):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0B4F5C"
    for c in range(1, 15):
        ws.column_dimensions[GL(c)].width = 15
    ws.column_dimensions["A"].width = 2
    ws.cell(1, 2, "لوحة التحكم الرئيسية - مطاعم ستاموني").font = Font(name=FONT, size=18, bold=True, color="FFFFFF")
    for c in range(1, 15):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.row_dimensions[1].height = 34
    put(ws, "B2", '="الفترة: "&YEAR(S_From)&"-"&TEXT(MONTH(S_From),"00")&"-"&TEXT(DAY(S_From),"00")&" → "&YEAR(S_To)&"-"&TEXT(MONTH(S_To),"00")&"-"&TEXT(DAY(S_To),"00")&"   |   الفرع: "&S_Branch&"   |   وضع البيانات: "&S_Mode&"   |   العملة: EGP"', size=10, bold=True)
    ws.merge_cells("B2:N2")
    put(ws, "B3", "=IS_Banner", size=11, bold=True)
    ws.merge_cells("B3:N3")
    ws.conditional_formatting.add("B3", FormulaRule(formula=['LEFT(B3,1)="⚠"'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    ws.conditional_formatting.add("B3", FormulaRule(formula=['LEFT(B3,1)="✔"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    SD = REG["Sales_Daily"]
    Sd = lambda k: SD.ws_col(k)
    prelim = '&IF(CK_Critical>0," (أولي)","")'

    def tile(row, col, label, formula, fmt="#,##0.00", sub=None, red_if_neg=False):
        a = ws.cell(row, col, label)
        a.font = Font(name=FONT, size=9, bold=True, color="404040")
        a.fill = PatternFill("solid", fgColor=C_HDR_BAND)
        a.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        v = ws.cell(row + 1, col, formula)
        v.font = Font(name=FONT, size=16, bold=True, color=C_TITLE)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="center", vertical="center")
        v.fill = PatternFill("solid", fgColor="FFFFFF")
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        s = ws.cell(row + 2, col, sub if sub else "")
        s.font = Font(name=FONT, size=8, color="595959")
        s.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
        for rr in range(row, row + 3):
            for cc in (col, col + 1):
                ws.cell(rr, cc).border = BORDER
        ws.row_dimensions[row + 1].height = 28
        if red_if_neg:
            ws.conditional_formatting.add(f"{GL(col)}{row + 1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

    chg = lambda cur, prev: f'=IF({prev}=0,"المقارنة: لا توجد فترة سابقة","السابقة: "&TEXT({prev},"#,##0")&"  ("&TEXT(({cur}-{prev})/ABS({prev}),"+0%;-0%")&")")'
    cols = [2, 4, 6, 8, 10, 12]
    r1, r2, r3, r4 = 5, 9, 13, 17
    last_day = f'_xlfn.MAXIFS({Sd("date")},{Sd("date")},"<="&S_To,{Sd("ok")},1)'
    tile(r1, cols[0], "مبيعات آخر يوم مسجل", f'=SUMIFS({Sd("gross")},{Sd("date")},{last_day},{Sd("ok")},1)',
         sub=f'=IF({last_day}=0,"لا مبيعات مرحلة","بتاريخ "&YEAR({last_day})&"-"&TEXT(MONTH({last_day}),"00")&"-"&TEXT(DAY({last_day}),"00"))')
    tile(r1, cols[1], "إجمالي مبيعات الفترة", "=IS_GrossSales", sub=chg("IS_GrossSales", "IS_GrossSales_P") if False else "قبل الخصومات والمرتجعات")
    tile(r1, cols[2], "صافي المبيعات", "=IS_NetSales", sub=chg("IS_NetSales", "IS_NetSales_P")[1:] and None)
    ws.cell(r1 + 2, cols[2], chg("IS_NetSales", "IS_NetSales_P"))
    tile(r1, cols[3], "عدد الطلبات", f'=SUMIFS({Sd("orders")},{Sd("date")},{chr(34)}>={chr(34)}&S_From,{Sd("date")},{chr(34)}<={chr(34)}&S_To,{Sd("ok")},1)', fmt="#,##0", sub="من تقارير المبيعات المرحلة")
    tile(r1, cols[4], "متوسط قيمة الطلب", f'=IF({ws.cell(r1 + 1, cols[3]).coordinate}=0,"غير متوفر",SUMIFS({Sd("net")},{Sd("date")},{chr(34)}>={chr(34)}&S_From,{Sd("date")},{chr(34)}<={chr(34)}&S_To,{Sd("ok")},1)/{ws.cell(r1 + 1, cols[3]).coordinate})', sub="صافي المبيعات ÷ عدد الطلبات")
    tile(r1, cols[5], '="مجمل الربح"' + prelim.replace('&IF', '&IF'), "=IS_GrossProfit")
    ws.cell(r1, cols[5]).value = '="مجمل الربح"' + prelim
    ws.cell(r1 + 2, cols[5], chg("IS_GrossProfit", "IS_GrossProfit_P"))
    tile(r2, cols[0], "x", "=IS_NetProfit", red_if_neg=True)
    ws.cell(r2, cols[0]).value = '="صافي الربح"&IF(CK_Critical>0," (أولي - غير مؤكد)"," (مؤكد)")'
    ws.cell(r2 + 2, cols[0], chg("IS_NetProfit", "IS_NetProfit_P"))
    tile(r2, cols[1], "نسبة تكلفة المبيعات", '=IF(IS_NetSales=0,"-",IS_COGS/IS_NetSales)', fmt="0.0%", sub="تكلفة المبيعات ÷ صافي المبيعات")
    tile(r2, cols[2], "المصروفات التشغيلية", "=IS_OpEx")
    ws.cell(r2 + 2, cols[2], chg("IS_OpEx", "IS_OpEx_P"))
    tile(r2, cols[3], "قيمة المخزون", "=" + netdr([("acct", '">=1400"'), ("acct", '"<=1499"'), ("date", '"<="&S_To')]), sub="حسابات 1400-1499 (دفتريًا)")
    tile(r2, cols[4], "مستحقات منصات التوصيل والبطاقات", "=" + netdr([("acct", '">=1302"'), ("acct", '"<=1304"'), ("date", '"<="&S_To')]), sub="1302-1304 غير محصلة")
    tile(r2, cols[5], "أرصدة الموردين", "=SB_Total", sub="إجمالي المستحق للموردين")
    tile(r3, cols[0], "الهالك (طبيعي وغير طبيعي)", "=" + netdr([("acct", '">=5200"'), ("acct", '"<=5299"'), ("date", '">="&S_From'), ("date", '"<="&S_To')]), sub="المرحلة 2 للتفصيل")
    tile(r3, cols[1], "فروق الجرد (مخزون)", "=" + netdr([("acct", '">=5400"'), ("acct", '"<=5499"'), ("date", '">="&S_From'), ("date", '"<="&S_To')]), sub="موجب = عجز", red_if_neg=False)
    tile(r3, cols[2], "فروق الخزائن (عجز - زيادة)", "=" + netdr([("acct", "ACC_CASH_SHORT"), ("date", '">="&S_From'), ("date", '"<="&S_To')]) + "-" + f'(-{netdr([("acct", "ACC_CASH_OVER"), ("date", chr(34)+">="+chr(34)+"&S_From"), ("date", chr(34)+"<="+chr(34)+"&S_To")])})', sub="موجب = عجز صافٍ")
    tile(r3, cols[3], "أرصدة البنوك", "=" + netdr([("acct", '">=1200"'), ("acct", '"<=1299"'), ("date", '"<="&S_To')]), red_if_neg=True)
    tile(r3, cols[4], "النقدية بالخزائن", "=" + netdr([("acct", '">=1100"'), ("acct", '"<=1199"'), ("date", '"<="&S_To')]), red_if_neg=True)
    tile(r3, cols[5], "مديونيات العملاء المستحقة", "=" + netdr([("acct", "ACC_CUSTOMERS"), ("date", '"<="&S_To')]), sub="حساب 1301")
    # treasuries
    TR = REG["Treasury_Balances"]
    band(ws, 21, "رصيد كل خزينة وبنك (Treasury_Balances)", 14)
    for i, h in enumerate(["الخزينة/البنك", "الاسم", "", "الرصيد", "تنبيه"]):
        pass
    heads = [(2, "الكود"), (3, "الاسم"), (6, "الرصيد الختامي"), (8, "تنبيه")]
    for col, t in heads:
        c = ws.cell(22, col, t)
        c.font = Font(name=FONT, bold=True, size=9)
        c.fill = PatternFill("solid", fgColor=C_HDR_BAND)
    for k in range(10):
        rr = 23 + k
        src = TR.first + k
        ws.cell(rr, 2, f'=Treasury_Balances!${TR.L("code")}${src}')
        ws.cell(rr, 3, f'=Treasury_Balances!${TR.L("name")}${src}')
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=5)
        c = ws.cell(rr, 6, f'=Treasury_Balances!${TR.L("close")}${src}')
        c.number_format = "#,##0.00;[Red]-#,##0.00"
        ws.cell(rr, 8, f'=Treasury_Balances!${TR.L("flag")}${src}')
        for col in (2, 3, 6, 8):
            ws.cell(rr, col).border = BORDER
            ws.cell(rr, col).font = Font(name=FONT, size=9)
    # branches
    band(ws, 35, "مقارنة أداء الفروع (من قائمة الدخل - بدون توزيع تكاليف الإدارة)", 14)
    bh = ["الفرع", "إجمالي المبيعات", "صافي المبيعات", "تكلفة المبيعات", "مجمل الربح", "% مجمل الربح", "مصروفات مباشرة", "الربح التشغيلي"]
    for i, t in enumerate(bh):
        c = ws.cell(36, 2 + i, t)
        c.font = Font(name=FONT, bold=True, size=9)
        c.fill = PatternFill("solid", fgColor=C_HDR_BAND)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for k in range(6):
        rr = 37 + k
        col = GL(9 + k)
        ws.cell(rr, 2, f'=Income_Statement!{col}${H_is}')
        ws.cell(rr, 3, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["GS"]})')
        ws.cell(rr, 4, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["NS"]})')
        ws.cell(rr, 5, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["COGS"]})')
        ws.cell(rr, 6, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["GP"]})')
        ws.cell(rr, 7, f'=IF(B{rr}="","",IF(D{rr}=0,0,F{rr}/D{rr}))')
        ws.cell(rr, 8, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["OPEX"]})')
        ws.cell(rr, 9, f'=IF(B{rr}="","",Income_Statement!{col}{keymap["OP"]})')
        for cc in range(2, 10):
            c = ws.cell(rr, cc)
            c.border = BORDER
            c.font = Font(name=FONT, size=9)
            c.number_format = "0.0%" if cc == 7 else "#,##0;[Red]-#,##0"
    ws.conditional_formatting.add("G37:G42", FormulaRule(formula=['AND(ISNUMBER(G37),G37<0.5)'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    # expense breakdown
    band(ws, 44, "المصروفات التشغيلية حسب النوع", 14)
    for k, key in enumerate(["o1", "o2", "o3", "o4", "o5", "o6", "o7", "o8", "o9"]):
        rr = 45 + k
        ws.cell(rr, 2, f'=Income_Statement!$A${keymap[key]}')
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=4)
        c = ws.cell(rr, 5, f'=Income_Statement!$E${keymap[key]}')
        c.number_format = "#,##0;[Red]-#,##0"
        c2 = ws.cell(rr, 6, f'=Income_Statement!$G${keymap[key]}')
        c2.number_format = "#,##0;[Red]-#,##0"
        for cc in (2, 5, 6):
            ws.cell(rr, cc).border = BORDER
            ws.cell(rr, cc).font = Font(name=FONT, size=9)
    ws.cell(44, 5, "الحالية").font = Font(name=FONT, size=9, bold=True)
    ws.cell(44, 6, "السابقة").font = Font(name=FONT, size=9, bold=True)
    # alerts
    band(ws, 58, "التنبيهات (Alerts)", 14)
    alerts = [
        '=IF(CK_Critical>0,"✖ "&CK_Critical&" بند حرج مفتوح في قائمة الإقفال - النتائج أولية","✔ لا توجد بنود إقفال حرجة مفتوحة")',
        '=IF(CK_Drafts+CK_Approved>0,"⚠ "&CK_Drafts&" مستند مسودة و"&CK_Approved&" معتمد غير مرحّل في الفترة","✔ كل مستندات الفترة مرحّلة أو ملغاة")',
        '=IF(CK_DocDup+CK_InvDup>0,"✖ مستندات مكررة: أرقام "&CK_DocDup&" وفواتير موردين "&CK_InvDup,"✔ لا توجد مستندات مكررة")',
        '=IF(CK_Unbal>0,"✖ "&CK_Unbal&" قيد يدوي غير متوازن لم يُرحَّل","✔ كل القيود اليدوية متوازنة")',
        '=IF(CK_NoRecon>0,"⚠ "&CK_NoRecon&" يوم مبيعات بدون مطابقة يومية","✔ المطابقة اليومية للمبيعات مكتملة")',
        '=IF(CK_ReconOpen+CK_CashOpen>0,"⚠ فروق مفتوحة: مبيعات "&CK_ReconOpen&" وخزائن "&CK_CashOpen,"✔ لا فروق مفتوحة في المبيعات والخزائن")',
        '=IF(CK_Stale>0,"⚠ "&CK_Stale&" منصة/بطاقات بمستحقات متأخرة التسوية","✔ تسويات المنصات منتظمة")',
        '=IF(SB_Overdue>0,"⚠ مستحقات موردين متأخرة: "&TEXT(SB_Overdue,"#,##0"),"✔ لا مستحقات موردين متأخرة")',
        '=IF(CK_NoRef>0,"⚠ "&CK_NoRef&" مستند مرحل بدون مرجع/مرفق ورقي","✔ كل المستندات المرحلة لها مرجع ورقي")',
        '=IF(CK_CogsMissing>0,"✖ قيد تسوية تكلفة المبيعات غير مسجل: الربح مبالغ فيه","✔ تكلفة المبيعات مسجلة")',
        '=IF(ABS(TB_Diff)>S_Tol,"✖ ميزان المراجعة غير متوازن","✔ ميزان المراجعة متوازن")',
    ]
    for i, a in enumerate(alerts):
        rr = 59 + i
        ws.cell(rr, 2, a).font = Font(name=FONT, size=10)
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=12)
    ws.conditional_formatting.add("B59:B69", FormulaRule(formula=['LEFT(B59,1)="✖"'], fill=PatternFill("solid", bgColor=C_RED, fgColor=C_RED)))
    ws.conditional_formatting.add("B59:B69", FormulaRule(formula=['LEFT(B59,1)="⚠"'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    ws.conditional_formatting.add("B59:B69", FormulaRule(formula=['LEFT(B59,1)="✔"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    # charts
    ch = BarChart()
    ch.type = "col"
    ch.title = "صافي المبيعات حسب الفرع"
    ch.height, ch.width = 7.5, 14
    data = Reference(ws, min_col=4, min_row=36, max_row=42)
    cats = Reference(ws, min_col=2, min_row=37, max_row=42)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ws.add_chart(ch, "J22")
    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "المصروفات التشغيلية (الفترة الحالية)"
    ch2.height, ch2.width = 5.6, 14
    d2 = Reference(ws, min_col=5, min_row=44, max_row=53)
    c2 = Reference(ws, min_col=2, min_row=45, max_row=53)
    ch2.add_data(d2, titles_from_data=True)
    ch2.set_categories(c2)
    ch2.legend = None
    ws.add_chart(ch2, "J44")
    ws.freeze_panes = "A4"
    protect_ws(ws)
    return ws
