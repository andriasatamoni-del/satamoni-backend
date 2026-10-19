# -*- coding: utf-8 -*-
"""مولّد Satamoni_Financial_Management_System.xlsx (يعمل في Excel وGoogle Sheets).

الاستخدام:
    python3 build_system.py --out ../output --scale 1.0
ينتج الملفين: النسخة التجريبية (ببيانات اختبار) والنسخة النظيفة (قوالب بلا بيانات تشغيل).
"""
import argparse
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import fs_common as fc
from fs_common import *
import fs_data as D
import fs_content as C
import fs_manual as M
import b1_masters as b1
import b2_sources as b2
import b3_ledger as b3
import b4_control as b4
import b5_samples as b5
import b6_tests as b6

d = dt.date

SYSTEM_MAP = [
    ("مبيعات الفرع اليومية", "تقرير مبيعات يومي (15)", "الكاشير / أمين الصندوق", "مدير الفرع ← المحاسب يرحّل", "مدين: تحصيل كل وسيلة دفع (خزينة / بطاقات / منصات) + خصومات + مرتجعات | دائن: إيراد كل فئة (+ضريبة)", "نقدية↑ أو مستحقات↑ | ربح↑ | مخزون: لا (التكلفة في التسوية الدورية)", "Income_Statement, Dashboard, Daily_Recon", "1"),
    ("مطابقة المبيعات اليومية", "كشف إقفال (14) + تقرير الجهاز/المنصة", "أمين الصندوق", "مدير الفرع", "لا قيد؛ الفرق يُسجل بسند عجز/زيادة", "يكشف الفروق قبل أن تدخل الأرباح", "Daily_Recon, Close_Checklist", "1"),
    ("تسوية المنصات والبطاقات", "كشف تسوية (16)", "المحاسب", "المدير المالي", "مدين: بنك + عمولة + خصومات (+1309) | دائن: مستحقات المنصة", "نقدية↑ | مستحقات↓ | ربح↓ بقيمة العمولة فقط", "Platform_Settlements, Balance_Sheet", "1"),
    ("الخزائن وسندات القبض والصرف", "سند قبض (11) / صرف (12)", "أمين الصندوق", "المعتمد ← المحاسب", "قبض: مدين خزينة / دائن الحساب المقابل | صرف: العكس", "حسب الحساب المقابل", "Treasury_Balances, GL_Ledger", "1"),
    ("التحويل بين الخزائن والبنوك", "إذن تحويل (13)", "المسلِّم والمستلم", "المعتمد", "مدين: الخزينة المستلمة / دائن: المسلِّمة", "لا أثر على الإيراد أو المصروف", "Treasury_Balances, Cash_Flow (تحويلات داخلية = 0)", "1"),
    ("عجز وزيادة الخزينة", "كشف إقفال (14) + تسوية (30)", "أمين الصندوق", "مدير الفرع + المدير المالي", "عجز: مدين 6910 / دائن خزينة | زيادة: مدين خزينة / دائن 4210", "نقدية↓/↑ | ربح↓/↑", "Cash_Close, Dashboard", "1"),
    ("المصروفات والمستحقات", "نموذج مصروفات (17)", "مقدم الطلب", "مدير الفرع/الإدارة ← المحاسب", "فوري: مدين مصروف / دائن خزينة | آجل: دائن 2201", "ربح↓ بتاريخ الاستحقاق | نقدية بتاريخ الدفع", "Expenses, Income_Statement", "1"),
    ("المشتريات (طلب ← أمر ← استلام ← فاتورة ← سداد)", "نماذج 1 و2 و3 + فاتورة المورد", "مشرف المشتريات / أمين المخزن", "معتمد الشراء ← المدير المالي", "فاتورة: مدين مخزون/أصل/مصروف / دائن 2101 | الطلب والأمر لا قيد لهما", "مخزون↑ والتزام↑ | لا أثر على الربح", "Supplier_Invoices, Supplier_Balances", "1 (الفاتورة) / 2 (الطلب والأمر والاستلام)"),
    ("الموردون ودفعاتهم", "سند صرف (12)", "مشرف المشتريات", "المدير المالي", "مدين 2101 / دائن خزينة أو بنك", "التزام↓ ونقدية↓", "Supplier_Payments, Supplier_Balances", "1"),
    ("المخزون والتحويلات", "إذن استلام (3) / صرف (4) / تحويل (5) / جرد (10)", "أمين المخزن", "معتمد الصرف", "التحويل: مدين مخزون المستلم / دائن المسلِّم (بالتكلفة الدفترية)", "لا إيراد ولا تكلفة مبيعات", "Inventory_Moves (المرحلة 2)", "2"),
    ("السنتر كيتشن والإنتاج", "نماذج 6 و7 و8 و9", "مدير الإنتاج", "المدير المالي (للهالك غير الطبيعي)", "خام ← تحت التشغيل ← تام ← مخزون الفرع، والهالك بحساب منفصل", "لا إيراد؛ التكلفة تنتقل مع المنتج", "Production, Recipes (المرحلة 2)", "2"),
    ("تكلفة المبيعات", "محضر جرد (10) + تسوية (25)", "المحاسب + أمين المخزن", "المدير المالي", "النسخة 1: مدين 5101 / دائن 1401.. بالمستهلك شهريًا (دوري مُعدّل) | المرحلة 2: مستمر", "ربح↓ | مخزون↓", "Income_Statement, Close_Checklist", "1 / 2"),
    ("الرواتب والسلف والعهد", "كشف رواتب (22) / سلفة (18) / عهدة (19)", "الموارد البشرية / المحاسب", "المدير المالي", "استحقاق: مدين 6101 / دائن 2202 | سلفة: مدين 1501", "ربح↓ بالاستحقاق | نقدية بالدفع", "Payroll, Employees (المرحلة 3)", "3 (مؤقتًا بقيد يدوي)"),
    ("الأصول الثابتة والإهلاك", "فاتورة أصل + بطاقة أصل", "المحاسب", "المدير المالي", "شراء: مدين 17xx / دائن 2101 | إهلاك: مدين 6801 / دائن 179x", "ربح↓ بالإهلاك فقط", "Fixed_Assets (المرحلة 3)", "3 (مؤقتًا بقيد يدوي)"),
    ("القيود والتسويات", "قيد يومية (24) / تسوية (25)", "المحاسب", "المدير المالي", "حسب القيد: يجب أن يتساوى المدين والدائن قبل الترحيل", "حسب القيد", "GL_Journal, Manual_JE", "1"),
    ("الإقفال الشهري", "محضر إقفال (26)", "المحاسب", "المدير المالي + الإدارة العليا", "لا قيد شهري؛ لقطة نتائج وقت الإقفال", "يثبت أرقام الفترة", "Close_Checklist, Period_Close", "1"),
    ("التقارير", "—", "المحاسب", "—", "تقرأ من دفتر اليومية فقط", "—", "Trial_Balance, Income_Statement, Balance_Sheet, Cash_Flow, Dashboard", "1"),
]

SHEET_GUIDE = [
    ("Dashboard", "A/K", "لوحة التحكم الرئيسية", "لا إدخال", "مؤشرات ورسوم وتنبيهات", "—"),
    ("Guide", "A", "هذا الدليل وخريطة الأوراق", "—", "—", "—"),
    ("System_Map", "A", "خريطة النظام: الوحدات والمستندات والقيود", "—", "—", "—"),
    ("Settings", "A", "فترة التقارير وطريقة التكلفة وحدود السماح", "الخلايا الصفراء", "تتحكم في كل التقارير", "—"),
    ("Lists", "A", "القوائم المنسدلة وحالات المستندات", "المسؤول فقط", "—", "—"),
    ("COA", "A/J", "دليل الحسابات الموحد", "حساب جديد في صف فارغ", "يغذي القوائم والتقارير", "—"),
    ("Locations", "A", "الفروع والمخازن والمطابخ والإدارة", "صف لكل موقع", "—", "—"),
    ("Treasuries", "A/C", "الخزائن والبنوك وحساب كل منها", "صف لكل خزينة", "—", "—"),
    ("Users / Permissions", "A", "المستخدمون ومصفوفة الصلاحيات", "—", "—", "—"),
    ("Suppliers / Customers / Employees", "A", "الموردون والعملاء والموظفون", "—", "—", "—"),
    ("Items / Units / Unit_Conversions / Unit_Converter", "A/E", "الأصناف والوحدات والتحويل المعتمد ومحوّل مضبوط", "—", "—", "—"),
    ("Payment_Methods / Cost_Centers / Tax_Settings", "A", "وسائل الدفع ومراكز التكلفة والضرائب", "—", "—", "—"),
    ("Doc_Numbering / Doc_Register", "A", "ترقيم المستندات وسجل الأرقام المركزي", "—", "منع التكرار", "—"),
    ("Sales_Daily", "B", "المبيعات اليومية (صف لكل فرع ويوم)", "نموذج 15", "GL_Journal", "اليومية"),
    ("Platform_Settlements", "B", "تسويات المنصات والبطاقات", "نموذج 16", "GL_Journal", "—"),
    ("Daily_Recon", "B", "المطابقة اليومية للمبيعات", "الفعلي والسبب", "Close_Checklist", "—"),
    ("Cash_Vouchers", "C", "قبض/صرف/تحويل/إيداع/عجز/زيادة", "نماذج 11-13 و30", "GL_Journal", "—"),
    ("Cash_Close", "C", "إقفال الخزينة اليومي وفروق الجرد", "نموذج 14", "Close_Checklist", "—"),
    ("Treasury_Balances / Cash_Flow", "C/K", "أرصدة الخزائن وحركة النقدية والتدفقات", "لا إدخال", "تقارير", "—"),
    ("Supplier_Invoices / Supplier_Payments", "D", "فواتير الموردين ودفعاتهم", "فاتورة المورد / نموذج 12", "GL_Journal", "—"),
    ("Supplier_Balances / Customer_Balances", "D/K", "الأرصدة وأعمار الديون", "لا إدخال", "تقارير", "—"),
    ("Expenses", "H", "المصروفات والمستحقات", "نموذج 17", "GL_Journal", "—"),
    ("Manual_JE", "J", "القيود اليدوية والتسويات", "نماذج 24-25", "GL_Journal", "—"),
    ("Posting_Rules", "J", "خريطة الحسابات وقواعد القيود", "المسؤول فقط", "—", "—"),
    ("GL_Journal", "J", "دفتر اليومية العام المولّد آليًا", "لا إدخال", "كل التقارير", "—"),
    ("GL_Ledger", "K", "الأستاذ العام وكشف الحساب", "اختيار حساب/جهة", "—", "—"),
    ("Trial_Balance / Income_Statement / Balance_Sheet", "K", "ميزان المراجعة وقائمة الدخل والميزانية", "لا إدخال", "—", "—"),
    ("Close_Checklist / Close_Confirm / Period_Close", "L", "الإقفال وتأكيداته وسجله", "تأكيدات", "—", "—"),
    ("Audit_Log", "L", "سجل المراجعة (يدوي + Apps Script)", "تسجيل الأحداث", "—", "—"),
    ("Test_Cases", "—", "اختبارات الجودة الحية", "—", "—", "—"),
    ("JE_Guide / Forms_Catalog / Glossary", "—", "دليل القيود والنماذج والقاموس", "—", "—", "—"),
    ("Known_Issues / Roadmap / Expansion_Guide", "—", "المشكلات المعروفة وخطة التطوير وإرشادات التوسعة", "—", "—", "—"),
    ("PR, Quotations, PO, GRN, Inventory_Moves, Stock_Count, Recipes, Production, Payroll, Fixed_Assets, Advances_Custody", "المراحل 2-3", "هياكل جاهزة للمراحل القادمة", "—", "—", "لا تؤثر على الدفاتر بعد"),
]


def skeleton(name, title, sub, cols, tab="A6A6A6"):
    sh = Sheet(name, title, sub, cap=40, tab=tab)
    for k, (ar, w) in enumerate(cols):
        sh.add(f"c{k}", ar, w=w)
    return sh


def D_(a, b, c, e):
    pass


def build(out, demo=True, with_br04=True, period=None, ov=None, scale=1.0, tag=""):
    fc.REG.clear()
    fc.ALL.clear()
    ov = ov or {}
    b2.SCALE = scale
    b3.ROWMAP.clear()
    b3.DR_MAP.clear()
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- تعريف الأوراق
    m = b1.define_masters(demo, with_br04)
    if not demo:
        m["us"].samples = []
    else:
        extra = [("USR14", "أمين صندوق فرع 3 (تجريبي)", "أمين صندوق", "BR03"), ("USR15", "مدير الفرع 4 (تجريبي)", "مدير فرع", "BR04"),
                 ("USR16", "أمين صندوق فرع 4 (تجريبي)", "أمين صندوق", "BR04"), ("USR17", "محاسب مساعد (تجريبي)", "المحاسب", "HQ01")]
        m["us"].samples += [dict(code=a, name=b, role=c, loc=l, active="نعم", email="user@example.com") for a, b, c, l in extra]
    pm = m["pm"]
    pm.add("bal", "رصيد المستحقات", "f", 14, fmt="#,##0.00;[Red]-#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",IF({x("type")}="نقدي",0,{b3.netdr([("acct", x("gl")), ("date", chr(34)+"<="+chr(34)+"&S_To")])}))')
    pm.add("lastset", "آخر تسوية", "f", 12, fmt="yyyy-mm-dd",
           f=lambda x: f'=IF({x("code")}="","",IF(COUNTIFS({REG["Platform_Settlements"].ws_col("pm")},{x("code")},{REG["Platform_Settlements"].ws_col("ok")},1)=0,"",_xlfn.MAXIFS({REG["Platform_Settlements"].ws_col("date")},{REG["Platform_Settlements"].ws_col("pm")},{x("code")},{REG["Platform_Settlements"].ws_col("ok")},1)))')
    pm.add("stale", "تسوية متأخرة؟", "f", 10, align="center",
           f=lambda x: (f'=IF({x("code")}="","",IF(AND(OR({x("type")}="منصة",{x("type")}="بطاقة"),{x("bal")}>S_Tol),IF(S_To-IF({x("lastset")}="",'
                        f'_xlfn.MINIFS({b1.G("date")},{b1.G("acct")},{x("gl")},{b1.G("eff")},1),{x("lastset")})>S_PlatDays,1,0),0))'))
    pm.add("nextdue", "موعد التسوية المتوقع", "f", 14, fmt="yyyy-mm-dd",
           f=lambda x: f'=IF({x("code")}="","",IF(OR({x("lastset")}="",NOT(ISNUMBER({x("lag")}))),"",{x("lastset")}+{x("lag")}))')
    S = b2.define_sources(demo, with_br04, ov)
    gl, dreg = b3.define_ledger(wb, None)
    lg = b3.define_gl_ledger(wb)
    tb = b3.define_tb(wb)
    tr = b3.define_treasury_balances(wb)
    sb = b3.define_supplier_balances(wb)
    cb = b3.define_customer_balances(wb)
    cc = b4.define_close_confirm()
    pc = b4.define_period_close()
    al = Sheet("Audit_Log", "سجل المراجعة (Audit Log)", "سجّل هنا أي تعديل جوهري أو عكس أو إقفال أو تغيير بيانات أساسية. Apps Script الاختياري يسجل تعديلات الأوراق المصدرية آليًا.", cap=400, tab="C00000")
    al.add("ts", "التاريخ", w=12, fmt="yyyy-mm-dd", dv=("date",))
    al.add("user", "المستخدم", w=24, dv=("list", "=LST_USER"))
    al.add("sheet", "الورقة", w=20)
    al.add("ref", "رقم المستند / الخلية", w=22)
    al.add("action", "الإجراء", w=14, dv=("list", "إضافة,تعديل,اعتماد,ترحيل,عكس,إلغاء,إقفال,تغيير بيانات أساسية,استعادة نسخة"))
    al.add("old", "القيمة السابقة", w=20)
    al.add("new", "القيمة الجديدة", w=20)
    al.add("reason", "السبب", w=30)
    al.add("appr", "المعتمد", w=24, dv=("list", "=LST_USER"))
    # هياكل المراحل القادمة
    skels = [
        skeleton("PR", "طلبات الشراء PR (المرحلة 2 - هيكل فقط)", "لا يؤثر على الدفاتر بعد. يُفعَّل في المرحلة 2 مع أوامر الشراء والاستلام.",
                 [("رقم الطلب", 14), ("التاريخ", 12), ("طالب الشراء", 20), ("الموقع", 9), ("الصنف", 12), ("الكمية", 10), ("الوحدة", 8), ("تاريخ الاحتياج", 12), ("الحالة", 10), ("المعتمد", 20)]),
        skeleton("Quotations", "عروض الأسعار (المرحلة 2 - هيكل فقط)", "مقارنة عروض الموردين قبل أمر الشراء عند الحاجة.",
                 [("رقم العرض", 14), ("التاريخ", 12), ("المورد", 10), ("الصنف", 12), ("الكمية", 10), ("السعر", 10), ("صلاحية العرض", 12), ("شروط السداد", 12), ("المختار؟", 9), ("ملاحظات", 24)]),
        skeleton("PO", "أوامر الشراء PO (المرحلة 2 - هيكل فقط)", "أمر الشراء لا يُعد مصروفًا ولا قيدًا نهائيًا.",
                 [("رقم الأمر", 14), ("التاريخ", 12), ("المورد", 10), ("رقم الطلب PR", 14), ("الصنف", 12), ("الكمية", 10), ("السعر", 10), ("الإجمالي", 12), ("شروط السداد", 12), ("الحالة", 10), ("المعتمد", 20)]),
        skeleton("GRN", "محاضر الاستلام GRN (المرحلة 2 - هيكل فقط)", "الاستلام قبل الفاتورة يُسجل في 2110 (بضاعة مستلمة غير مفوترة) عند التفعيل.",
                 [("رقم الاستلام", 14), ("التاريخ", 12), ("رقم الأمر", 14), ("المورد", 10), ("الصنف", 12), ("الكمية المستلمة", 12), ("الكمية المرفوضة", 12), ("الوحدة", 8), ("أمين المخزن", 20), ("ملاحظات", 24)]),
        skeleton("Inventory_Moves", "حركات المخزون (المرحلة 2 - هيكل فقط)", "الرصيد يُحسب من هذا السجل فقط، وتكلفة التحويل = التكلفة الدفترية المعتمدة.",
                 [("رقم الحركة", 14), ("التاريخ", 12), ("الصنف", 12), ("الكمية", 10), ("وحدة القياس", 10), ("تكلفة الوحدة", 12), ("القيمة", 12), ("المخزن المصدر", 11), ("المخزن المستلم", 11), ("نوع الحركة", 14), ("المستند المرجعي", 16), ("منفذ الحركة", 18), ("المعتمد", 18)]),
        skeleton("Stock_Count", "الجرد وفروق الجرد (المرحلة 2 - هيكل فقط)", "الرصيد الدفتري ثم الفعلي ثم الفرق ثم مستند التسوية.",
                 [("رقم محضر الجرد", 14), ("التاريخ", 12), ("الموقع", 9), ("الصنف", 12), ("الرصيد الدفتري", 12), ("الرصيد الفعلي", 12), ("الفرق", 10), ("قيمة الفرق", 12), ("السبب", 20), ("معتمد", 18)]),
        skeleton("Recipes", "وصفات الإنتاج والمنيو (المرحلة 2 - هيكل فقط)", "كود الوصفة، المكونات، الفاقد المعياري، الناتج الصالح، تكلفة الدفعة والوحدة، ومقارنة بالفعلي.",
                 [("كود الوصفة", 12), ("الاسم", 22), ("النوع (إنتاج/منيو)", 12), ("المكوّن", 12), ("الكمية", 10), ("الوحدة", 8), ("الفاقد المعياري %", 10), ("الناتج الصالح", 11), ("تكلفة الدفعة", 12), ("تكلفة الوحدة", 12), ("سعر البيع", 10)]),
        skeleton("Production", "الإنتاج في السنتر كيتشن (المرحلة 2 - هيكل فقط)", "طلبات الإنتاج، الصرف، الناتج، الهالك الطبيعي/غير الطبيعي، التحويل للفروع.",
                 [("رقم تقرير الإنتاج", 14), ("التاريخ", 12), ("الوصفة", 12), ("الكمية المخططة", 11), ("الكمية الفعلية", 11), ("هالك طبيعي", 10), ("هالك غير طبيعي", 10), ("محوّل إلى", 10), ("رقم التحويل", 14), ("المسؤول", 18)]),
        skeleton("Payroll", "الرواتب (المرحلة 3 - هيكل فقط)", "الاستحقاق منفصل عن الدفع، والجزاءات لا تُخصم دون اعتماد ومراجعة قانونية.",
                 [("الموظف", 10), ("الفرع", 9), ("الشهر", 9), ("الأساسي", 11), ("البدلات", 10), ("الإضافي", 10), ("الجزاءات المعتمدة", 11), ("السلف", 10), ("الخصومات", 10), ("صافي المستحق", 12), ("المدفوع", 10), ("المتبقي", 10)]),
        skeleton("Fixed_Assets", "سجل الأصول الثابتة (المرحلة 3 - هيكل فقط)", "كود الأصل، العمر الإنتاجي، طريقة الإهلاك، مجمع الإهلاك والقيمة الدفترية.",
                 [("كود الأصل", 12), ("الوصف", 24), ("الحساب", 9), ("تاريخ الشراء", 12), ("التكلفة", 12), ("العمر (سنة)", 9), ("قيمة الخردة", 10), ("طريقة الإهلاك", 12), ("الموقع", 9), ("مجمع الإهلاك", 12), ("القيمة الدفترية", 12)]),
        skeleton("Advances_Custody", "السلف والعهد (المرحلة 3 - هيكل فقط)", "السلفة والعهدة أصل وليست مصروفًا حتى تُسوى بالمستندات.",
                 [("الرقم", 14), ("التاريخ", 12), ("الموظف", 10), ("النوع (سلفة/عهدة)", 12), ("المبلغ", 12), ("المُسدَّد/المُسوّى", 12), ("الرصيد", 12), ("الحالة", 10), ("المعتمد", 20)]),
    ]

    # ---------------- عينات
    if demo:
        SAMPLE, E = b5.build(with_br04)
        for key, sh in [("sd", "Sales_Daily"), ("dr", "Daily_Recon"), ("ps", "Platform_Settlements"), ("cv", "Cash_Vouchers"), ("ex", "Expenses"),
                        ("si", "Supplier_Invoices"), ("sp", "Supplier_Payments"), ("mj", "Manual_JE")]:
            REG[sh].samples = SAMPLE[key]
        sr = E["sales_rows"]
        c = lambda doc: sr[doc]["m1"]
        cl = [
            dict(date=d(2026, 9, 20), tre="CSH-BR02", counted=12000 + c("SAL-BR02-0001") + 5000 + c("SAL-BR02-0002") - 150, reason="عجز نقدي", resp=b5.CSH["BR02"], res="تمت التسوية", adj="ADJ-BR02-0001"),
            dict(date=d(2026, 9, 30), tre="CSH-BR01", counted=15000 + c("SAL-BR01-0001") + c("SAL-BR01-0002") + c("SAL-BR01-0003") - 10000 - 600),
            dict(date=d(2026, 9, 30), tre="CSH-BR03", counted=10000 + c("SAL-BR03-0001") + c("SAL-BR03-0002") + c("SAL-BR03-0003")),
            dict(date=d(2026, 10, 12), tre="CSH-BR02", counted=12000 + c("SAL-BR02-0001") + c("SAL-BR02-0002") + c("SAL-BR02-0003") + 5000 - 150 + c("SAL-BR02-0004") - 150,
                 reason="عجز نقدي", resp=b5.CSH["BR02"], res="تمت التسوية", adj="ADJ-BR02-0002"),
        ]
        REG["Cash_Close"].samples = cl
        REG["Close_Confirm"].samples = [dict(period=202609, item=mid, value="نعم", by=(b5.ACC if i < 9 else b5.FM), on=d(2026, 10, 1 if i < 9 else 2),
                                             notes=("مطابقة الدورة المخزنية يدوية - المرحلة 2" if mid in ("M05", "M06", "M07") else "")) for i, (mid, _, _) in enumerate(b4.CK_ITEMS_MANUAL)]
        pc.samples = b4.snapshot_rows(ov)
    else:
        E = dict(sales_rows={})
        pc.samples = []

    # ---------------- إنشاء الأوراق (الترتيب = ترتيب التبويبات النهائي بعد إعادة الترتيب)
    b1.define_named_lists(wb)
    b1.build_lists(wb)
    b1.build_settings(wb, demo)
    b1.build_tax(wb)
    # فترة التقارير
    p = period or (d(2026, 10, 1), d(2026, 10, 31))
    ws_set = wb["Settings"]
    ws_set["C9"].value, ws_set["C10"].value = p
    for sh in fc.ALL:
        sh.render(wb, with_samples=True)
    # حسابات النظام (Posting_Rules)
    build_posting_rules(wb)
    ws_is, keymap, H_is = b3.build_income_statement(wb)
    b3.build_balance_sheet(wb)
    b3.build_cash_flow(wb)
    b4.build_close_checklist(wb, None)
    b4.build_dashboard(wb, keymap, H_is)
    b6.build_tests(wb, E, with_br04 and demo, keymap, H_is) if demo else build_static_tests(wb)
    # نصية
    t = b1.mk_text_sheet
    t(wb, "Glossary", "قاموس المصطلحات المحاسبية البسيط", "اشرح لكل موظف بمثال من مطعم ستاموني", ["المصطلح", "بالإنجليزية", "الشرح المبسط", "مثال من ستاموني"],
      [list(g) for g in C.GLOSSARY], [16, 18, 64, 56], tab="7F7F7F")
    t(wb, "System_Map", "خريطة النظام (System Map)", "الوحدات المالية والتشغيلية وحركة البيانات بينها", ["الوحدة", "المستند الذي يبدأ الحركة", "من يعدّه", "من يعتمده", "القيد المحاسبي الناتج", "الأثر (خزينة / مخزون / تكلفة / ربح)", "التقارير", "المرحلة"],
      [list(r) for r in SYSTEM_MAP], [26, 26, 20, 24, 56, 40, 34, 14], tab="7F7F7F")
    guide_rows = [list(r) for r in SHEET_GUIDE]
    t(wb, "Guide", "دليل الأوراق (Sheet Guide)", "ابدأ من هنا: ماذا تفعل كل ورقة وماذا تدخل فيها. التفاصيل الكاملة في ملفات docs", ["الورقة", "المجموعة", "الوظيفة", "المدخلات", "المخرجات", "ملاحظات"], guide_rows, [44, 10, 46, 26, 28, 24], tab="7F7F7F")
    manual_rows = []
    for p_ in M.MANUAL:
        manual_rows.append([p_["no"], p_["title"], p_["who"], p_["doc"], "\n".join(f"{i}) {s}" for i, s in enumerate(p_["steps"], 1)), p_["fields"], p_["result"], p_["verify"], "\n".join("• " + x for x in p_["mistakes"]), p_["approver"]])
    t(wb, "User_Manual", "دليل التشغيل المبسط لغير المحاسبين", "14 إجراء بالترتيب: من المسؤول، المستند، الخطوات، الخانات، النتيجة، التحقق، الأخطاء الشائعة، المعتمِد",
      ["م", "الإجراء", "من المسؤول؟", "المستند المطلوب", "الخطوات بالترتيب", "الخانات التي يملؤها", "النتيجة المتوقعة", "كيف يعرف أنها تمت صحيحة", "الأخطاء الشائعة", "من يعتمد"], manual_rows, [5, 24, 24, 24, 70, 36, 30, 30, 34, 22], tab="7F7F7F")
    je_rows = []
    for e in C.JE_EXAMPLES:
        for stp, lines in e["steps"]:
            dr = " + ".join(f'{c} {D.ACCT_BY_CODE[c]["ar"]}: {a:,.0f}' for s, c, a in lines if s == "D")
            cr = " + ".join(f'{c} {D.ACCT_BY_CODE[c]["ar"]}: {a:,.0f}' for s, c, a in lines if s == "K")
            tot = sum(a for s, c, a in lines if s == "D")
            je_rows.append([e["no"], e["title"] + (f" - {stp}" if stp else ""), e["story"], e["doc"], dr, cr, tot, e["why"], e["effect"]["bs"], e["effect"]["pl"], e["effect"]["cash"], e["effect"]["inv"], e["how"], e["review"]])
    t(wb, "JE_Guide", "دليل القيود المحاسبية - 28 حالة بأرقام تجريبية", "كل مثال يستخدم حسابات دليل الحسابات الفعلي. الأرقام تجريبية فقط. القيد دائمًا متوازن.",
      ["م", "الحالة", "القصة", "المستند المطلوب", "المدين", "الدائن", "المبلغ التجريبي", "سبب المدين والدائن", "الأثر على الميزانية", "الأثر على قائمة الدخل", "الأثر على الخزينة", "الأثر على المخزون", "طريقة التسجيل في Google Sheets", "طريقة المراجعة"],
      je_rows, [5, 30, 42, 30, 44, 44, 12, 44, 30, 30, 26, 28, 56, 40], tab="7F7F7F")
    form_rows = []
    for f in C.FORMS:
        cp = " | ".join(f"{n} ({col}) ← {h}" for n, col, h, w in f["copies"])
        when = " | ".join(f"{n}: {w}" for n, col, h, w in f["copies"] if w)
        form_rows.append([f["no"], f["title"], f["prefix"], len(f["copies"]), cp, when, ", ".join(f["sigs"]), "رقم مسلسل، تاريخ، الفرع/الموقع، البيانات، التوقيعات، الاعتماد، المرفقات والملاحظات"])
    t(wb, "Forms_Catalog", "كتالوج النماذج الورقية (30 نموذجًا)", "ملف forms/Satamoni_Forms_A4.html (وPDF) يحتوي النماذج جاهزة للطباعة A4", ["رقم النموذج", "الاسم", "بادئة الترقيم", "عدد النسخ", "النسخ ومن يحتفظ بها", "التسليم للمحاسب", "التوقيعات", "الخانات الثابتة"],
      form_rows, [8, 30, 10, 8, 56, 56, 40, 40], tab="7F7F7F")
    t(wb, "Doc_Cancel_Rules", "إلغاء المستندات وفقدها واكتشاف الأخطاء", "ماذا يحدث إذا أُلغي مستند أو فُقد أو اكتُشف فيه خطأ", ["الحالة", "الإجراء"], [list(x) for x in C.CANCEL_RULES], [30, 100], tab="7F7F7F")
    perm_rows = [[role] + vals for role, vals in C.PERMS.items()]
    t(wb, "Permissions", "مصفوفة الصلاحيات والفصل بين المسؤوليات", "تطبيق فعلي: Google Sheets sharing + حماية الأوراق + مراجعة مستقلة. (✔ مسموح، ✖ غير مسموح)", ["الدور"] + C.ACTIONS, perm_rows, [20] + [20] * len(C.ACTIONS), tab="2E75B6")
    ws = wb["Permissions"]
    r0 = 5 + len(perm_rows) + 2
    ws.cell(r0, 1, "قواعد الفصل بين المسؤوليات").font = Font(name=FONT, bold=True, size=12)
    for i, (a, b) in enumerate(C.SOD_RULES, start=1):
        ws.cell(r0 + i, 1, a).alignment = Alignment(wrap_text=True)
        ws.cell(r0 + i, 2, b)
    r1 = r0 + len(C.SOD_RULES) + 2
    ws.cell(r1, 1, "ضوابط تعويضية عند محدودية الموظفين").font = Font(name=FONT, bold=True, size=12)
    for i, a in enumerate(C.COMPENSATING, start=1):
        ws.cell(r1 + i, 1, a)
    t(wb, "Known_Issues", "المشكلات والقيود المعروفة (Known Issues)", "سجل صريح بما لا يفعله النظام الآن", ["الأهمية", "المشكلة", "الوصف", "التخفيف / الحل", "الحالة"], [list(x) for x in C.KNOWN_ISSUES], [10, 36, 70, 70, 22], tab="7F7F7F")
    t(wb, "Expansion_Guide", "إعدادات وتوسعة النظام (Expansion Guide)", "كيف تضيف فرعًا أو مخزنًا أو حسابًا أو وسيلة دفع أو تفعّل مرحلة جديدة دون إعادة بناء الملف", ["الإجراء", "الخطوات بالترتيب", "ما يتغير تلقائيًا", "ملاحظات/حدود"],
      [list(x) for x in C.EXPANSION], [30, 80, 44, 44], tab="7F7F7F")
    t(wb, "Roadmap", "خطة التطوير (Roadmap)", "مراحل التنفيذ", ["المرحلة", "المحتوى"], [list(x) for x in C.ROADMAP], [24, 120], tab="7F7F7F")
    t(wb, "Open_Questions", "الأسئلة المفتوحة للمالك", "أسئلة ضرورية فقط؛ النظام يعمل بافتراضات مؤقتة قابلة للتعديل حتى تُجاب", ["م", "السؤال"], [[i, q] for i, q in enumerate(C.QUESTIONS, 1)], [6, 130], tab="7F7F7F")

    order = ["Dashboard", "Guide", "System_Map", "User_Manual", "Settings", "Lists", "COA", "Locations", "Treasuries", "Users", "Permissions", "Suppliers", "Customers", "Employees",
             "Items", "Units", "Unit_Conversions", "Unit_Converter", "Payment_Methods", "Cost_Centers", "Tax_Settings", "Doc_Numbering", "Doc_Register", "Doc_Cancel_Rules",
             "Sales_Daily", "Platform_Settlements", "Daily_Recon", "Cash_Vouchers", "Cash_Close", "Treasury_Balances", "Cash_Flow",
             "Supplier_Invoices", "Supplier_Payments", "Supplier_Balances", "Customer_Balances", "Expenses", "Manual_JE", "Posting_Rules", "GL_Journal", "GL_Ledger",
             "Trial_Balance", "Income_Statement", "Balance_Sheet", "Close_Checklist", "Close_Confirm", "Period_Close", "Audit_Log", "Test_Cases", "JE_Guide", "Forms_Catalog", "Glossary",
             "Known_Issues", "Roadmap", "Open_Questions", "Expansion_Guide", "PR", "Quotations", "PO", "GRN", "Inventory_Moves", "Stock_Count", "Recipes", "Production", "Payroll", "Fixed_Assets", "Advances_Custody"]
    names = [s.title for s in wb._sheets]
    missing = [n for n in names if n not in order]
    assert not missing, missing
    wb._sheets = [wb[n] for n in order if n in names]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Satamoni_Financial_Management_System"
    wb.properties.subject = "نظام الحسابات والإدارة المالية - مطاعم ستاموني (المرحلة 1)"
    wb.properties.creator = "Satamoni Finance Builder"
    path = os.path.join(out, f"Satamoni_Financial_Management_System{tag}.xlsx")
    wb.save(path)
    return path, E


def build_posting_rules(wb):
    ws = wb.create_sheet("Posting_Rules")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    ws.cell(1, 1, "قواعد القيود المحاسبية وخريطة الحسابات (Posting Rules)").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 9):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "أ) خريطة الحسابات الثابتة تستخدمها المعادلات بالاسم - تعديل رقم الحساب هنا يغيّر القيود الجديدة والقديمة معًا (لا تغيّره بعد بدء التشغيل). ب) جدول القواعد يشرح مدين/دائن كل نوع حركة.").font = Font(name=FONT, size=9, italic=True)
    for col, w in zip("ABCDEFGH", [34, 26, 40, 40, 40, 34, 22, 28]):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(["اسم الحساب في الصيغ", "رقم الحساب", "اسم الحساب (آلي)", "الاستخدام"], start=1):
        c = ws.cell(4, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_IN_H)
        c.border = BORDER
    for k, (nm, code, use) in enumerate(D.ACCOUNT_MAP, start=5):
        ws.cell(k, 1, nm).border = BORDER
        c = ws.cell(k, 2, code)
        c.fill = PatternFill("solid", fgColor=C_IN)
        c.border = BORDER
        from openpyxl.styles import Protection
        c.protection = Protection(locked=False)
        ws.cell(k, 3, f'=IFERROR(INDEX({REG["COA"].ws_col("ar")},MATCH(B{k},{REG["COA"].ws_col("code")},0)),"حساب غير موجود")').border = BORDER
        ws.cell(k, 4, use).border = BORDER
        name_cell(wb, nm, f"Posting_Rules!$B${k}")
    r = 5 + len(D.ACCOUNT_MAP) + 2
    ws.cell(r, 1, "ب) قواعد القيود لكل نوع حركة").font = Font(name=FONT, bold=True, size=12)
    heads = ["نوع الحركة", "ورقة التسجيل", "الحساب المدين", "الحساب الدائن", "شروط الإثبات", "المستند المؤيد", "توقيت التسجيل", "الإلغاء / التصحيح"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(r + 1, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_IN_H)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    for k, row in enumerate(C.POSTING_RULES, start=r + 2):
        for i, v in enumerate(row, start=1):
            c = ws.cell(k, i, v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(r + 2, 2)
    ws.protection.sheet = False
    protect_ws(ws)


def build_static_tests(wb):
    ws = wb.create_sheet("Test_Cases")
    ws.sheet_view.rightToLeft = True
    ws.cell(1, 1, "اختبارات الجودة - النسخة النظيفة: استخدم النسخة التجريبية لاختبارات T01-T18 و V01-V14").font = Font(name=FONT, size=12, bold=True)
    ws.cell(3, 1, "هذه النسخة بلا بيانات؛ بعد إدخال بيانات حقيقية شغّل فحوص التكامل: TB_Diff و GL_TotalDiff و BS_Diff و SB_Diff و CF_Diff في Close_Checklist.")
    return ws


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="../output")
    ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--clean", action="store_true", help="النسخة النظيفة فقط")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    if not a.clean:
        p, E = build(a.out, demo=True, with_br04=True, scale=a.scale)
        print(p)
    p, _ = build(a.out, demo=False, with_br04=False, scale=a.scale, tag="_CLEAN_TEMPLATE")
    print(p)
