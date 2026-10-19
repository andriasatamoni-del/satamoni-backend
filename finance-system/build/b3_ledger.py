# -*- coding: utf-8 -*-
"""سجل المستندات + دفتر اليومية العام + الأستاذ + ميزان المراجعة + القوائم المالية + تقارير الخزينة والموردين."""
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from fs_common import *
from b1_masters import G, band, CAP_COA, CAP_TRE, CAP_SUP, CAP_CUS
import fs_data as D


def src(sheet, key, sr):
    return f"{sheet}!${REG[sheet].L(key)}{sr}"


def COAc(k):
    return REG["COA"].ws_col(k)


# =================================================================== blocks
def make_blocks():
    s = lambda sh: (lambda k: (lambda sr: src(sh, k, sr)))
    z = lambda sr: "0"
    cc0 = lambda sr: '"CC00"'
    dash = lambda sr: '"-"'
    blocks = []
    def _blk1():
        f = s("Sales_Daily")
        base = dict(sheet="Sales_Daily", stype="SAL", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        slots = []
        for k, code in enumerate(D.SALES_PM_CODES + ["OTH"], start=1):
            slots.append(dict(act=(lambda k: lambda sr: f'AND({f("used")(sr)}=1,N({f(f"m{k}")(sr)})>0)')(k),
                              acct=f(f"a{k}"), loc=f("loc"), cc=lambda sr: '"CC10"', party=lambda sr: '"-"',
                              desc=(lambda code: lambda sr: f'"مبيعات - تحصيل {code}"')(code),
                              dr=f(f"m{k}"), cr=lambda sr: "0", cfc=lambda sr: '"OPERATING"'))
        slots.append(dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("dex")(sr)})>0)', acct=lambda sr: "ACC_DISC", loc=f("loc"), cc=lambda sr: '"CC10"', party=lambda sr: '"-"',
                          desc=lambda sr: '"خصومات المبيعات والعروض"', dr=f("dex"), cr=lambda sr: "0", cfc=lambda sr: '""'))
        slots.append(dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("rex")(sr)})>0)', acct=lambda sr: "ACC_RET", loc=f("loc"), cc=lambda sr: '"CC10"', party=lambda sr: '"-"',
                          desc=lambda sr: '"مرتجعات وإلغاءات المبيعات"', dr=f("rex"), cr=lambda sr: "0", cfc=lambda sr: '""'))
        for k in range(1, 5):
            slots.append(dict(act=(lambda k: lambda sr: f'AND({f("used")(sr)}=1,N({f(f"r{k}")(sr)})>0)')(k),
                              acct=f(f"ra{k}"), loc=f("loc"), cc=lambda sr: '"CC10"', party=lambda sr: '"-"',
                              desc=(lambda k: lambda sr: f'"إيراد مبيعات - فئة {["فطير","بيتزا","مكرونات","أخرى"][k-1]}"')(k),
                              dr=lambda sr: "0", cr=f(f"r{k}"), cfc=lambda sr: '""'))
        slots.append(dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("vat")(sr)})>0)', acct=lambda sr: "ACC_VAT_OUT", loc=f("loc"), cc=lambda sr: '"CC10"', party=lambda sr: '"-"',
                          desc=lambda sr: '"ضريبة القيمة المضافة - مخرجات"', dr=lambda sr: "0", cr=f("vat"), cfc=lambda sr: '""'))
        return dict(base, slots=slots)

    blocks.append(_blk1())
    def _blk2():
        f = s("Platform_Settlements")
        base = dict(sheet="Platform_Settlements", stype="PLT", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        slots = [
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("actual")(sr)})>0)', acct=f("ba"), loc=f("loc"), cc=cc0, party=dash, desc=lambda sr: '"استلام تسوية منصة/بطاقات"', dr=f("actual"), cr=z, cfc=lambda sr: '"OPERATING"'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("comm")(sr)})>0)', acct=f("ca"), loc=f("loc"), cc=lambda sr: '"CC30"', party=dash, desc=lambda sr: '"عمولة المنصة/رسوم البطاقات"', dr=f("comm"), cr=z, cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("other")(sr)})>0)', acct=lambda sr: "ACC_PLT_OTHER", loc=f("loc"), cc=lambda sr: '"CC30"', party=dash, desc=lambda sr: '"خصومات/غرامات أخرى من المنصة"', dr=f("other"), cr=z, cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("diff")(sr)})>0)', acct=lambda sr: "ACC_PLT_DIFF", loc=f("loc"), cc=cc0, party=dash, desc=lambda sr: '"فرق تسوية منصة (نقص) تحت المراجعة"', dr=f("diff"), cr=z, cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("diff")(sr)})<0)', acct=lambda sr: "ACC_PLT_DIFF", loc=f("loc"), cc=cc0, party=dash, desc=lambda sr: '"فرق تسوية منصة (زيادة) تحت المراجعة"', dr=z, cr=lambda sr: f'-{f("diff")(sr)}', cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("settled")(sr)})>0)', acct=f("ra"), loc=f("loc"), cc=cc0, party=dash, desc=lambda sr: '"تخفيض مستحقات المنصة/البطاقات"', dr=z, cr=f("settled"), cfc=lambda sr: '""'),
        ]
        return dict(base, slots=slots)

    blocks.append(_blk2())
    def _blk3():
        f = s("Cash_Vouchers")
        base = dict(sheet="Cash_Vouchers", stype="CV", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        dsc = lambda sr: f'IF({f("desc")(sr)}="",{f("vtype")(sr)}&" - "&{f("payee")(sr)},{f("desc")(sr)})'
        slots = [
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("amount")(sr)})>0)', acct=f("dracct"), loc=f("drloc"), cc=f("drcc"), party=f("drparty"), desc=dsc, dr=f("amount"), cr=z, cfc=f("cfc")),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("amount")(sr)})>0)', acct=f("cracct"), loc=f("crloc"), cc=f("crcc"), party=f("crparty"), desc=dsc, dr=z, cr=f("amount"), cfc=f("cfc")),
        ]
        return dict(base, slots=slots)

    blocks.append(_blk3())
    def _blk4():
        f = s("Expenses")
        base = dict(sheet="Expenses", stype="EXP", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        dsc = lambda sr: f'IF({f("desc")(sr)}="","مصروف - "&{f("payee")(sr)},{f("desc")(sr)})'
        slots = [
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("amount")(sr)})>0)', acct=f("acct"), loc=f("loc"), cc=f("cc"), party=dash, desc=dsc, dr=f("amount"), cr=z, cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("vat")(sr)})>0)', acct=lambda sr: "ACC_VAT_IN", loc=f("loc"), cc=f("cc"), party=dash, desc=lambda sr: '"ضريبة قيمة مضافة - مدخلات"', dr=f("vat"), cr=z, cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("total")(sr)})>0)', acct=f("cr"), loc=f("loc"), cc=cc0, party=dash, desc=dsc, dr=z, cr=f("total"), cfc=lambda sr: '"OPERATING"'),
        ]
        return dict(base, slots=slots)

    blocks.append(_blk4())
    def _blk5():
        f = s("Supplier_Invoices")
        base = dict(sheet="Supplier_Invoices", stype="SINV", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        isCN = lambda sr: f'{f("tc")(sr)}="CN"'
        dsc = lambda sr: f'{f("itype")(sr)}&" "&{f("sinv")(sr)}&" - "&{f("sup")(sr)}'
        pay_act = lambda sr: f'AND({f("used")(sr)}=1,{f("paynow")(sr)}<>"",{f("tc")(sr)}<>"CN",N({f("total")(sr)})>0)'
        slots = [
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("net")(sr)})>0)', acct=f("dacct"), loc=f("loc"), cc=f("cc"), party=dash, desc=dsc,
                 dr=lambda sr: f'IF({isCN(sr)},0,{f("net")(sr)})', cr=lambda sr: f'IF({isCN(sr)},{f("net")(sr)},0)', cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("vat")(sr)})>0)', acct=lambda sr: "ACC_VAT_IN", loc=f("loc"), cc=f("cc"), party=dash, desc=lambda sr: '"ضريبة قيمة مضافة - مدخلات"',
                 dr=lambda sr: f'IF({isCN(sr)},0,{f("vat")(sr)})', cr=lambda sr: f'IF({isCN(sr)},{f("vat")(sr)},0)', cfc=lambda sr: '""'),
            dict(act=lambda sr: f'AND({f("used")(sr)}=1,N({f("total")(sr)})>0)', acct=lambda sr: "ACC_SUPPLIERS", loc=f("loc"), cc=cc0, party=f("sup"), desc=dsc,
                 dr=lambda sr: f'IF({isCN(sr)},{f("total")(sr)},0)', cr=lambda sr: f'IF({isCN(sr)},0,{f("total")(sr)})', cfc=lambda sr: '""'),
            dict(act=pay_act, acct=lambda sr: "ACC_SUPPLIERS", loc=f("loc"), cc=cc0, party=f("sup"), desc=lambda sr: f'"سداد فوري - "&{dsc(sr)}', dr=f("total"), cr=z, cfc=lambda sr: '""'),
            dict(act=pay_act, acct=f("tgl"), loc=f("loc"), cc=cc0, party=dash, desc=lambda sr: f'"سداد فوري - "&{dsc(sr)}', dr=z, cr=f("total"), cfc=f("cfp")),
        ]
        return dict(base, slots=slots)

    blocks.append(_blk5())
    def _blk6():
        f = s("Supplier_Payments")
        base = dict(sheet="Supplier_Payments", stype="SPAY", je=lambda sr: f'"J-"&{f("doc")(sr)}', date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"))
        dsc = lambda sr: f'IF({f("desc")(sr)}="","دفعة مورد "&{f("sup")(sr)},{f("desc")(sr)})'
        act = lambda sr: f'AND({f("used")(sr)}=1,N({f("amount")(sr)})>0)'
        slots = [
            dict(act=act, acct=lambda sr: "ACC_SUPPLIERS", loc=f("tloc"), cc=cc0, party=f("sup"), desc=dsc, dr=f("amount"), cr=z, cfc=lambda sr: '""'),
            dict(act=act, acct=f("tgl"), loc=f("tloc"), cc=cc0, party=dash, desc=dsc, dr=z, cr=f("amount"), cfc=f("cfp")),
        ]
        return dict(base, slots=slots)

    blocks.append(_blk6())
    def _blk7():
        f = s("Manual_JE")
        base = dict(sheet="Manual_JE", stype="MJE", je=f("doc"), date=f("date"), doc=f("doc"), status=f("status"),
                    srcchk=f("chk"), ok=f("ok"), prep=f("prepared"), appr=f("approved"), poster=f("poster"), ref=f("ref"), rev=f("rev_ref"), used=f("used"),
                    closing=lambda sr: f'IF({f("jtype")(sr)}="إقفال",1,0)')
        slots = [dict(act=lambda sr: f'{f("used")(sr)}=1', acct=f("acct"), loc=f("loc"), cc=lambda sr: f'IF({f("cc")(sr)}="","CC00",{f("cc")(sr)})',
                      party=lambda sr: f'IF({f("party")(sr)}="","-",{f("party")(sr)})', desc=f("desc"),
                      dr=lambda sr: f'N({f("debit")(sr)})', cr=lambda sr: f'N({f("credit")(sr)})', cfc=f("cf"))]
        return dict(base, slots=slots)



    blocks.append(_blk7())
    return blocks

ROWMAP = {}
BLOCKS = []
DR_MAP = {}


def build_rowmaps(gl, dreg):
    BLOCKS[:] = make_blocks()
    r = gl.first
    rr = dreg.first
    for b in BLOCKS:
        sh = REG[b["sheet"]]
        for i in range(sh.cap):
            sr = sh.first + i
            for j in range(len(b["slots"])):
                ROWMAP[r] = (b, j, sr)
                r += 1
            DR_MAP[rr] = (b, sr)
            rr += 1
    return r - gl.first, rr - dreg.first


def count_rows():
    tot = 0
    dtot = 0
    for b in make_blocks():
        n = REG[b["sheet"]].cap
        tot += n * len(b["slots"])
        dtot += n
    return tot, dtot


# =================================================================== GL / Doc register
def define_ledger(wb, scale_caps):
    tot, dtot = count_rows()
    dreg = Sheet("Doc_Register", "السجل المركزي لأرقام المستندات (Document Register)",
                 "سجل آلي بكل أرقام المستندات من كل الأوراق لمنع التكرار. لا تُدخل شيئًا هنا؛ أي رقم مكرر يظهر في عمود (تكرار) ويمنع ترحيل المستند.",
                 cap=dtot, tab="7F7F7F", protect=True)
    for k, ar, w, fmt in [("docno", "رقم المستند", 20, None), ("stype", "النوع", 8, None), ("date", "التاريخ", 12, "yyyy-mm-dd"), ("loc", "الموقع", 9, None),
                          ("amount", "المبلغ", 14, "#,##0.00"), ("status", "الحالة", 10, None), ("sheet", "ورقة المصدر", 20, None)]:
        dreg.add(k, ar, "f", w, fmt=fmt, f=(lambda k: lambda x: dr_formula(k, x))(k))
    dreg.add("dup", "تكرار؟", "f", 9, align="center", f=lambda x: f'=IF({x("docno")}="","",IF(COUNTIF({x.col("docno")},{x("docno")})>1,"مكرر",""))')
    dreg.add("flag", "فحص", "f", 9, hide=True, f=lambda x: f'=IF({x("dup")}="مكرر","مكرر","")')

    gl = Sheet("GL_Journal", "دفتر اليومية العام (General Journal) - مولَّد آليًا من المستندات",
               "كل سطر يُولَّد بمعادلة من مستند مصدري. لا يُرحَّل السطر (أثر فعلي = 1) إلا إذا كان المستند (مرحل) وخاليًا من الأخطاء والقيد متوازنًا. للمراجعة: فلتر على عمود (أثر فعلي) أو (سبب عدم الترحيل).",
               cap=tot, tab="C00000", freeze_col=2)
    cols = [
        ("act", "سطر نشط", 6, None), ("id", "رقم السطر", 20, None), ("je", "رقم القيد", 20, None), ("date", "تاريخ القيد", 12, "yyyy-mm-dd"),
        ("doc", "رقم المستند الأصلي", 20, None), ("stype", "مصدر القيد", 8, None), ("acct", "رقم الحساب", 9, None), ("aname", "اسم الحساب", 30, None),
        ("loc", "الفرع", 8, None), ("cc", "مركز التكلفة", 8, None), ("party", "الجهة", 9, None), ("desc", "البيان", 34, None),
        ("dr", "مدين", 13, "#,##0.00"), ("cr", "دائن", 13, "#,##0.00"), ("status", "حالة القيد", 9, None), ("srcchk", "فحص المصدر", 30, None),
        ("ok", "جاهز للترحيل", 7, None), ("eff", "أثر فعلي", 7, None), ("effdr", "مدين مرحّل", 13, "#,##0.00"), ("effcr", "دائن مرحّل", 13, "#,##0.00"),
        ("per", "الفترة", 8, None), ("cash", "نقدي؟", 6, None), ("cfc", "تصنيف التدفق", 10, None), ("issue", "سبب عدم الترحيل", 36, None),
        ("key", "مفتاح الأستاذ", 12, None), ("prep", "مُعِد القيد", 18, None), ("appr", "معتمد القيد", 18, None), ("poster", "المرحِّل", 18, None), ("ref", "مرجع/مرفق", 14, None), ("revref", "مرجع التصحيح", 16, None), ("clos", "قيد إقفال؟", 6, None),
    ]
    for k, ar, w, fmt in cols:
        gl.add(k, ar, "f", w, fmt=fmt, f=(lambda k: lambda x: gl_formula(k, x))(k), hide=(k in ("act", "ok", "per", "cash", "key", "clos")))
    gl.add("flag", "فحص", "f", 6, hide=True, f=lambda x: "")
    n1, n2 = build_rowmaps(gl, dreg)
    assert n1 == tot and n2 == dtot, (n1, tot, n2, dtot)
    return gl, dreg


def dr_formula(key, x):
    b, sr = DR_MAP[x.r]
    sh = b["sheet"]
    st = b["stype"]
    if st == "MJE":
        docexpr = f'IF(AND({src(sh,"doc",sr)}<>"",{src(sh,"first",sr)}=1),{src(sh,"doc",sr)},"")'
    else:
        docexpr = f'IF({src(sh,"doc",sr)}="","",{src(sh,"doc",sr)})'
    if key == "docno":
        return "=" + docexpr
    if key == "stype":
        return f'=IF({x("docno")}="","","{st}")'
    if key == "sheet":
        return f'=IF({x("docno")}="","","{sh}")'
    if key == "date":
        return f'=IF({x("docno")}="","",{src(sh,"date",sr)})'
    if key == "status":
        return f'=IF({x("docno")}="","",{src(sh,"status",sr)})'
    if key == "loc":
        lk = {"SAL": "loc", "PLT": "loc", "CV": "tloc", "EXP": "loc", "SINV": "loc", "SPAY": "tloc", "MJE": "loc"}[st]
        return f'=IF({x("docno")}="","",{src(sh,lk,sr)})'
    if key == "amount":
        ak = {"SAL": "net", "PLT": "actual", "CV": "amount", "EXP": "total", "SINV": "total", "SPAY": "amount", "MJE": "jedr"}[st]
        return f'=IF({x("docno")}="","",{src(sh,ak,sr)})'
    raise KeyError(key)


import re
_SIMPLE = re.compile(r"^[A-Za-z_0-9]+!\$[A-Z]+\d+$")


def bs(e):
    return f'IF({e}="","",{e})' if _SIMPLE.match(e) else e


def gl_formula(key, x):
    b, j, sr = ROWMAP[x.r]
    sl = b["slots"][j]
    A = x("act")
    g = lambda body: f'=IF({A}=1,{body},"")'
    if key == "act":
        return f'=IF({sl["act"](sr)},1,0)'
    if key == "id":
        return g(f'{x("doc")}&"/{j + 1}"')
    if key == "je":
        return g(bs(b["je"](sr)))
    if key == "date":
        return g(b["date"](sr))
    if key == "doc":
        return g(bs(b["doc"](sr)))
    if key == "stype":
        return g(f'"{b["stype"]}"')
    if key == "acct":
        return g(bs(sl["acct"](sr)))
    if key == "aname":
        return g(f'IF(COUNTIF({COAc("code")},{x("acct")})=0,"حساب غير معرف",INDEX({COAc("ar")},MATCH({x("acct")},{COAc("code")},0)))')
    if key == "loc":
        return g(bs(sl["loc"](sr)))
    if key == "cc":
        return g(bs(sl["cc"](sr)))
    if key == "party":
        return g(bs(sl["party"](sr)))
    if key == "desc":
        return g(bs(sl["desc"](sr)))
    if key == "dr":
        return g(sl["dr"](sr))
    if key == "cr":
        return g(sl["cr"](sr))
    if key == "status":
        return g(b["status"](sr))
    if key == "srcchk":
        return g(f'IF({b["srcchk"](sr)}="","",{b["srcchk"](sr)})')
    if key == "ok":
        return g(b["ok"](sr))
    if key == "eff":
        return g(f'IF(AND({x("ok")}=1,COUNTIFS({COAc("code")},{x("acct")},{COAc("valid")},1)=1),1,0)')
    if key == "effdr":
        return g(f'IF({x("eff")}=1,N({x("dr")}),0)')
    if key == "effcr":
        return g(f'IF({x("eff")}=1,N({x("cr")}),0)')
    if key == "per":
        return g(f'YEAR({x("date")})*100+MONTH({x("date")})')
    if key == "cash":
        return g(f'IF(COUNTIF({COAc("code")},{x("acct")})=0,0,INDEX({COAc("cash")},MATCH({x("acct")},{COAc("code")},0)))')
    if key == "cfc":
        return g(bs(sl["cfc"](sr)))
    if key == "issue":
        return g(f'IF({x("eff")}=1,"",IF({x("status")}=ST_POST,IF({x("srcchk")}<>"",{x("srcchk")},"حساب غير صالح أو غير نشط"),IF({x("status")}=ST_REV,"معكوس: لا أثر على الدفاتر",IF({x("status")}=ST_CANC,"ملغى",IF({x("status")}=ST_APPR,"معتمد ولم يُرحَّل","مسودة")))))')
    if key == "key":
        return g(f'IF(AND({x("eff")}=1,OR(LG_Acct="",{x("acct")}=LG_Acct),OR(LG_Party="",{x("party")}=LG_Party),OR(LG_Loc="ALL",{x("loc")}=LG_Loc),OR(LG_CC="ALL",{x("cc")}=LG_CC),{x("date")}>=LG_From,{x("date")}<=LG_To),{x("date")}+ROW()/1000000,"")')
    if key == "prep":
        return g(f'IF({b["prep"](sr)}="","",{b["prep"](sr)})')
    if key == "appr":
        return g(f'IF({b["appr"](sr)}="","",{b["appr"](sr)})')
    if key == "poster":
        return g(f'IF({b["poster"](sr)}="","",{b["poster"](sr)})')
    if key == "ref":
        return g(f'IF({b["ref"](sr)}="","",{b["ref"](sr)})')
    if key == "clos":
        return g(b.get("closing", lambda sr: "0")(sr))
    if key == "revref":
        return g(f'IF({b["rev"](sr)}="","",{b["rev"](sr)})')
    if key == "flag":
        return ""
    raise KeyError(key)


# =================================================================== helpers for reports
def sumifs(sumkey, crits):
    parts = [G(sumkey)]
    for k, c in crits:
        parts.append(G(k))
        parts.append(c)
    return "SUMIFS(" + ",".join(parts) + ")"


def netdr(crits):
    return f'({sumifs("effdr", crits)}-{sumifs("effcr", crits)})'


def put(ws, addr, value, fmt=None, bold=False, fill=None, name=None, wb=None, color=None, size=10, align=None, unlock=False, wrap=False):
    c = ws[addr]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if unlock:
        c.protection = Protection(locked=False)
    c.border = BORDER
    if align or wrap:
        c.alignment = Alignment(horizontal=align, wrap_text=wrap, vertical="center")
    if name and wb is not None:
        name_cell(wb, name, f"'{ws.title}'!${c.column_letter}${c.row}".replace("'", ""))
    return c


def selector(ws, wb, addr, label_addr, label, value, name, fmt=None):
    put(ws, label_addr, label, bold=True, fill=C_HDR_BAND)
    put(ws, addr, value, fmt=fmt, fill=C_IN, name=name, wb=wb, unlock=True)


# =================================================================== GL_Ledger
def define_gl_ledger(wb):
    lg = Sheet("GL_Ledger", "الأستاذ العام وكشف حركة الحساب (General Ledger / Account Activity)",
               "اختر حسابًا (واختياريًا جهة: مورد/عميل/موظف، فرع، مركز تكلفة، فترة) لتظهر حركته مرتبة بالتاريخ مع الرصيد الجاري. أستاذ المورد = حساب 2101 + كود المورد. أستاذ العميل = 1301 + كود العميل. الرصيد: موجب = مدين، سالب = دائن.",
               hdr=12, cap=500, tab="C00000")

    def hdrfill(ws):
        selector(ws, wb, "B4", "A4", "رقم الحساب (فارغ = كل الحسابات)", None, "LG_Acct")
        selector(ws, wb, "B5", "A5", "كود الجهة (مورد/عميل/موظف) - اختياري", None, "LG_Party")
        selector(ws, wb, "B6", "A6", "الفرع/الموقع (ALL = الكل)", "=S_Branch", "LG_Loc")
        selector(ws, wb, "B7", "A7", "مركز التكلفة (ALL = الكل)", "=S_CC", "LG_CC")
        selector(ws, wb, "B8", "A8", "من تاريخ", "=S_From", "LG_From", "yyyy-mm-dd")
        selector(ws, wb, "B9", "A9", "إلى تاريخ", "=S_To", "LG_To", "yyyy-mm-dd")
        put(ws, "D4", "اسم الحساب", bold=True, fill=C_HDR_BAND)
        put(ws, "E4", '=IF(LG_Acct="","(كل الحسابات)",IFERROR(INDEX(' + COAc("ar") + ',MATCH(LG_Acct,' + COAc("code") + ',0)),"حساب غير معرف"))', fill=C_F)
        put(ws, "D5", "الرصيد الافتتاحي (قبل الفترة)", bold=True, fill=C_HDR_BAND)
        put(ws, "E5", '=IF(LG_Acct="",0,' + netdr([("acct", "LG_Acct"), ("date", '"<"&LG_From'), ("loc", 'IF(LG_Loc="ALL","<>",LG_Loc)'), ("cc", 'IF(LG_CC="ALL","<>",LG_CC)'), ("party", 'IF(LG_Party="","<>",LG_Party)')]) + ')', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="LG_Open", wb=wb)
        put(ws, "D6", "عدد الحركات المطابقة", bold=True, fill=C_HDR_BAND)
        put(ws, "E6", f'=COUNT({G("key")})', fill=C_F)
        put(ws, "D7", "تنبيه", bold=True, fill=C_HDR_BAND)
        put(ws, "E7", f'=IF(E6>{lg.cap},"عدد الحركات أكبر من سعة العرض ({lg.cap}) - ضيّق الفترة أو الفلتر",IF(LG_Acct="","اختر حسابًا لعرض الرصيد الجاري",""))', fill=C_F)
        put(ws, "D8", "إجمالي المدين / الدائن", bold=True, fill=C_HDR_BAND)
        put(ws, "E8", f'=SUM({lg.rng("dr")})', fmt="#,##0.00", fill=C_F)
        put(ws, "F8", f'=SUM({lg.rng("cr")})', fmt="#,##0.00", fill=C_F)
        put(ws, "D9", "الرصيد الختامي", bold=True, fill=C_HDR_BAND)
        put(ws, "E9", '=IF(LG_Acct="","",E5+E8-F8)', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F)
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["D"].width = 30
    lg.extra.append(hdrfill)
    K = lambda x: f'INDEX({G("key")},{x("row")})'
    lg.add("n", "م", "f", 6, f=lambda x: f'=ROW()-{lg.hdr}')
    for k, ar, w, fmt in [("date", "التاريخ", 12, "yyyy-mm-dd"), ("je", "رقم القيد", 20, None), ("doc", "المستند", 20, None), ("stype", "المصدر", 7, None),
                          ("acct", "الحساب", 8, None), ("aname", "اسم الحساب", 28, None), ("loc", "الفرع", 8, None), ("cc", "م.ت", 7, None), ("party", "الجهة", 9, None), ("desc", "البيان", 34, None),
                          ("dr", "مدين", 13, "#,##0.00"), ("cr", "دائن", 13, "#,##0.00")]:
        lg.add(k, ar, "f", w, fmt=fmt, f=(lambda k: lambda x: f'=IF({x("row")}="","",INDEX({G(k)},{x("row")}))')(k))
    lg.add("bal", "الرصيد الجاري", "f", 14, fmt="#,##0.00;[Red]-#,##0.00",
           f=lambda x: f'=IF({x("row")}="","",IF({x("n")}=1,LG_Open,${lg.L("bal")}{x.r - 1})+{x("dr")}-{x("cr")})')
    lg.add("key", "مفتاح", "f", 11, hide=True, f=lambda x: f'=IF({x("n")}>COUNT({G("key")}),"",SMALL({G("key")},{x("n")}))')
    lg.add("row", "صف", "f", 6, hide=True, f=lambda x: f'=IF({x("key")}="","",MATCH({x("key")},{G("key")},0))')
    lg.add("flag", "فحص", "f", 6, hide=True, f=lambda x: "")
    return lg


# =================================================================== Trial balance
def define_tb(wb):
    tb = Sheet("Trial_Balance", "ميزان المراجعة (Trial Balance)",
               "يقرأ من دفتر اليومية العام (الأسطر المرحّلة فقط). الحسابات الرئيسية تعرض مجاميع النطاق. الإجماليات تُحسب من الحسابات القابلة للترحيل فقط. أي فرق بين المدين والدائن = خطأ يمنع الإقفال.",
               hdr=14, cap=CAP_COA, tab="C00000")

    def top(ws):
        selector(ws, wb, "C4", "B4", "من تاريخ", "=S_From", "TB_From", "yyyy-mm-dd")
        selector(ws, wb, "C5", "B5", "إلى تاريخ", "=S_To", "TB_To", "yyyy-mm-dd")
        selector(ws, wb, "C6", "B6", "الفرع (ALL = الكل)", "=S_Branch", "TB_Branch")
        selector(ws, wb, "C7", "B7", "مركز التكلفة (ALL = الكل)", "=S_CC", "TB_CC")
        put(ws, "D4", 'معيار الفرع', fill=C_F)
        put(ws, "E4", '=IF(TB_Branch="ALL","<>",TB_Branch)', fill=C_F, name="TB_LOCC", wb=wb)
        put(ws, "D5", 'معيار م.ت', fill=C_F)
        put(ws, "E5", '=IF(TB_CC="ALL","<>",TB_CC)', fill=C_F, name="TB_CCC", wb=wb)
        fR, lR = tb.first, tb.last
        rng = lambda k: f'${tb.L(k)}${fR}:${tb.L(k)}${lR}'
        put(ws, "B9", "الإجماليات (الحسابات القابلة للترحيل فقط)", bold=True, fill=C_HDR_BAND)
        for col, t in [("C", "رصيد افتتاحي صافي"), ("D", "حركة مدين"), ("E", "حركة دائن"), ("F", "ختامي مدين"), ("G", "ختامي دائن"), ("H", "الفرق (مدين - دائن)")]:
            put(ws, f"{col}9", t, bold=True, fill=C_HDR_BAND, align="center")
        put(ws, "B10", "الإجمالي", bold=True)
        put(ws, "C10", f'=SUMIFS({rng("open")},{rng("post")},"نعم")', fmt="#,##0.00;[Red]-#,##0.00")
        put(ws, "D10", f'=SUMIFS({rng("pdr")},{rng("post")},"نعم")', fmt="#,##0.00")
        put(ws, "E10", f'=SUMIFS({rng("pcr")},{rng("post")},"نعم")', fmt="#,##0.00")
        put(ws, "F10", f'=SUMIFS({rng("cdr")},{rng("post")},"نعم")', fmt="#,##0.00")
        put(ws, "G10", f'=SUMIFS({rng("ccr")},{rng("post")},"نعم")', fmt="#,##0.00")
        put(ws, "H10", "=ROUND(F10-G10,2)", fmt="#,##0.00;[Red]-#,##0.00", bold=True, name="TB_Diff", wb=wb)
        put(ws, "B11", "حالة الميزان", bold=True)
        put(ws, "C11", '=IF(AND(ABS(H10)<=S_Tol,ABS(D10-E10)<=S_Tol,ABS(C10)<=S_Tol),"✔ متوازن","✖ غير متوازن - يمنع الإقفال")', bold=True, name="TB_Status", wb=wb)
        ws.merge_cells("C11:E11")
        put(ws, "B12", "مطابقة الأستاذ: (المدين المرحّل - الدائن المرحّل) لكل الدفتر", bold=True)
        put(ws, "C12", f'=ROUND(SUM({G("effdr")})-SUM({G("effcr")}),2)', fmt="#,##0.00;[Red]-#,##0.00", name="GL_TotalDiff", wb=wb)
        put(ws, "D12", '=IF(ABS(C12)<=S_Tol,"✔ متطابق","✖ فرق في الأستاذ")', bold=True)
        ws.column_dimensions["A"].width = 10
        ws.conditional_formatting.add("C11", FormulaRule(formula=['LEFT(C11,1)="✖"'], fill=PatternFill("solid", bgColor=C_RED, fgColor=C_RED)))
        ws.conditional_formatting.add("C11", FormulaRule(formula=['LEFT(C11,1)="✔"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    tb.extra.append(top)
    off = tb.first - 5
    CI = lambda col: (lambda x: f'IF(INDEX(COA!$A:$A,ROW()-{off})="","",INDEX({"COA!$" + col + ":$" + col},ROW()-{off}))')
    cL = lambda k: REG["COA"].L(k)
    tb.add("code", "كود", "f", 10, f=lambda x: "=" + CI("A")(x), align="center")
    tb.add("name", "اسم الحساب", "f", 40, f=lambda x: f'=IF({x("code")}="","",INDEX(COA!${cL("ar")}:${cL("ar")},ROW()-{off}))')
    rngc = lambda x: [("acct", f'">="&{x("from")}'), ("acct", f'"<="&{x("to")}'), ("loc", "TB_LOCC"), ("cc", "TB_CCC")]
    tb.add("open", "رصيد افتتاحي (مدين+/دائن-)", "f", 17, fmt="#,##0.00;[Red]-#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",{netdr(rngc(x) + [("date", chr(34)+"<"+chr(34)+"&TB_From")])})')
    tb.add("pdr", "حركة مدين", "f", 16, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",{sumifs("effdr", rngc(x) + [("date", chr(34)+">="+chr(34)+"&TB_From"), ("date", chr(34)+"<="+chr(34)+"&TB_To")])})')
    tb.add("pcr", "حركة دائن", "f", 16, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",{sumifs("effcr", rngc(x) + [("date", chr(34)+">="+chr(34)+"&TB_From"), ("date", chr(34)+"<="+chr(34)+"&TB_To")])})')
    tb.add("cdr", "ختامي مدين", "f", 16, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MAX({x("close")},0))')
    tb.add("ccr", "ختامي دائن", "f", 16, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MAX(-{x("close")},0))')
    tb.add("close", "رصيد ختامي صافي", "f", 16, fmt="#,##0.00;[Red]-#,##0.00", hide=True, f=lambda x: f'=IF({x("code")}="","",{x("open")}+{x("pdr")}-{x("pcr")})')
    tb.add("lvl", "مستوى", "f", 7, f=lambda x: f'=IF({x("code")}="","",INDEX(COA!${cL("lvl")}:${cL("lvl")},ROW()-{off}))', align="center")
    tb.add("post", "ترحيل", "f", 8, f=lambda x: f'=IF({x("code")}="","",INDEX(COA!${cL("post")}:${cL("post")},ROW()-{off}))', align="center")
    tb.add("from", "من", "f", 8, hide=True, f=lambda x: f'=IF({x("code")}="","",INDEX(COA!${cL("rng_from")}:${cL("rng_from")},ROW()-{off}))')
    tb.add("to", "إلى", "f", 8, hide=True, f=lambda x: f'=IF({x("code")}="","",INDEX(COA!${cL("rng_to")}:${cL("rng_to")},ROW()-{off}))')
    tb.add("flag", "فحص", "f", 6, hide=True, f=lambda x: "")
    return tb


# =================================================================== Statements (custom layout)
def build_income_statement(wb):
    ws = wb.create_sheet("Income_Statement")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    ws.cell(1, 1, "قائمة الدخل (Income Statement) - تقرير الأرباح والخسائر المجمع ومقارنة الفروع").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 17):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "تُقرأ من دفتر اليومية العام فقط (لا توجد خلايا إدخال للأرقام). أعمدة الفروع تتجاهل فلتر الفرع. التحويلات الداخلية (نقدية/مخزون) لا تمر بحسابات الإيراد أو المصروف فلا تتكرر في المجمّع.").font = Font(name=FONT, size=9, italic=True)
    selector(ws, wb, "B4", "A4", "من تاريخ", "=S_From", "IS_From", "yyyy-mm-dd")
    selector(ws, wb, "B5", "A5", "إلى تاريخ", "=S_To", "IS_To", "yyyy-mm-dd")
    selector(ws, wb, "B6", "A6", "الفرع (ALL = الكل)", "=S_Branch", "IS_Branch")
    selector(ws, wb, "B7", "A7", "مركز التكلفة (ALL = الكل)", "=S_CC", "IS_CC")
    put(ws, "D4", "معيار الفرع", fill=C_F)
    put(ws, "E4", '=IF(IS_Branch="ALL","<>",IS_Branch)', fill=C_F, name="IS_LOCC", wb=wb)
    put(ws, "D5", "معيار مركز التكلفة", fill=C_F)
    put(ws, "E5", '=IF(IS_CC="ALL","<>",IS_CC)', fill=C_F, name="IS_CCC", wb=wb)
    put(ws, "D6", "فترة المقارنة", fill=C_F)
    put(ws, "E6", '=YEAR(S_PFrom)&"-"&TEXT(MONTH(S_PFrom),"00")&"-"&TEXT(DAY(S_PFrom),"00")&" → "&YEAR(S_PTo)&"-"&TEXT(MONTH(S_PTo),"00")&"-"&TEXT(DAY(S_PTo),"00")', fill=C_F)
    put(ws, "A8", "حالة النتائج", bold=True, fill=C_HDR_BAND)
    put(ws, "B8", '=IF(CK_Critical>0,"⚠ نتائج أولية: توجد "&CK_Critical&" بنود إقفال حرجة غير مكتملة (انظر Close_Checklist)","✔ لا توجد بنود حرجة مفتوحة في فترة قائمة الإقفال")', bold=True, name="IS_Banner", wb=wb)
    ws.conditional_formatting.add("B8", FormulaRule(formula=['LEFT(B8,1)="⚠"'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    H = 11
    heads = ["البند", "من حساب", "إلى حساب", "الطبيعة", "الفترة الحالية", "% من صافي المبيعات", "الفترة السابقة", "التغيّر"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(H, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_F_H)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    # branch headers
    bcols = [9, 10, 11, 12, 13, 14]  # I..N
    for k, ci in enumerate(bcols, start=1):
        c = ws.cell(H, ci, f'=IF(COUNTIF({REG["Locations"].ws_col("seq")},{k})=0,"",INDEX({REG["Locations"].ws_col("code")},MATCH({k},{REG["Locations"].ws_col("seq")},0)))')
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_IN_H)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for ci, t in [(15, "المخازن/المطبخ/الإدارة (غير الفروع)"), (16, "الإجمالي المجمع (كل المواقع)")]:
        c = ws.cell(H, ci, t)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_TITLE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[H].height = 36
    ws.column_dimensions["A"].width = 44
    for col, w in zip("BCDEFGH", [9, 9, 8, 16, 11, 16, 14]):
        ws.column_dimensions[col].width = w
    for ci in range(9, 17):
        ws.column_dimensions[GL(ci)].width = 15

    lines = [
        ("hdr", None, "الإيرادات", None, None),
        ("rev", "r1", "مبيعات الفطير", 4101, 4101), ("rev", "r2", "مبيعات البيتزا", 4102, 4102),
        ("rev", "r3", "مبيعات المكرونات", 4103, 4103), ("rev", "r4", "مبيعات أخرى", 4104, 4104),
        ("sub", "GS", "إجمالي المبيعات (Gross Sales)", ["r1", "r2", "r3", "r4"], None),
        ("rev", "r6", "(-) خصومات المبيعات والعروض", 4901, 4901), ("rev", "r7", "(-) مرتجعات وإلغاءات المبيعات", 4902, 4902),
        ("sub", "NS", "صافي المبيعات (Net Sales)", ["GS", "r6", "r7"], None),
        ("hdr", None, "تكلفة المبيعات", None, None),
        ("exp", "c1", "تكلفة المواد والعبوات", 5100, 5199), ("exp", "c2", "الهالك", 5200, 5299),
        ("exp", "c3", "الضيافة والمجانيات", 5300, 5399), ("exp", "c4", "فروق الجرد والتكلفة", 5400, 5499),
        ("sub", "COGS", "إجمالي تكلفة المبيعات", ["c1", "c2", "c3", "c4"], None),
        ("tot", "GP", "مجمل الربح (Gross Profit)", ("NS", "COGS"), None),
        ("hdr", None, "المصروفات التشغيلية والإدارية", None, None),
        ("exp", "o1", "الرواتب والأجور", 6100, 6199), ("exp", "o2", "الإيجارات", 6200, 6299), ("exp", "o3", "المرافق والاتصالات", 6300, 6399),
        ("exp", "o4", "التسويق", 6400, 6499), ("exp", "o5", "الصيانة", 6500, 6599), ("exp", "o6", "مصروفات التشغيل (نقل/نظافة/مستلزمات)", 6600, 6699),
        ("exp", "o7", "العمولات والرسوم (منصات/بنوك)", 6700, 6799), ("exp", "o8", "الإهلاك", 6800, 6899),
        ("exp", "o9", "مصروفات إدارية وأخرى (تشمل عجز الخزينة)", 6900, 6999),
        ("sub", "OPEX", "إجمالي المصروفات التشغيلية", ["o1", "o2", "o3", "o4", "o5", "o6", "o7", "o8", "o9"], None),
        ("tot", "OP", "الربح التشغيلي (Operating Profit)", ("GP", "OPEX"), None),
        ("rev", "oi", "إيرادات أخرى (تشمل زيادة الخزينة)", 4200, 4299),
        ("net", "NP", "صافي الربح (Net Profit)", ("OP", "oi"), None),
    ]
    r = H + 1
    plan = []
    keymap = {}
    for kind, key, label, a, b in lines:
        plan.append((r, kind, key, label, a, b))
        if key:
            keymap[key] = r
        r += 1
    allcols = ["E", "G"] + [GL(c) for c in bcols] + ["O", "P"]
    NSr = keymap["NS"]

    def amount(kind, a, b, col):
        if col == "G":
            dcr = [("date", '">="&S_PFrom'), ("date", '"<="&S_PTo')]
        else:
            dcr = [("date", '">="&IS_From'), ("date", '"<="&IS_To')]
        if col in ("E", "G"):
            loccr = [("loc", "IS_LOCC")]
        elif col == "P":
            loccr = [("loc", '"<>"')]
        else:
            loccr = [("loc", f"{col}${H}")]
        crits = [("acct", f'">={a}"'), ("acct", f'"<={b}"')] + dcr + loccr + [("cc", "IS_CCC"), ("clos", "0")]
        if kind == "rev":
            return f'{sumifs("effcr", crits)}-{sumifs("effdr", crits)}'
        return f'{sumifs("effdr", crits)}-{sumifs("effcr", crits)}'

    for (rr, kind, key, label, a, b) in plan:
        c = ws.cell(rr, 1, label)
        c.font = Font(name=FONT, bold=(kind not in ("rev", "exp")), size=10)
        c.border = BORDER
        if kind == "hdr":
            for ci in range(1, 17):
                ws.cell(rr, ci).fill = PatternFill("solid", fgColor=C_HDR_BAND)
            continue
        if kind in ("rev", "exp"):
            ws.cell(rr, 2, a).border = BORDER
            ws.cell(rr, 3, b).border = BORDER
            ws.cell(rr, 4, "دائن" if kind == "rev" else "مدين").border = BORDER
            for col in allcols:
                if col == "O":
                    ws[f"{col}{rr}"].value = f"=P{rr}-SUM(I{rr}:N{rr})"
                else:
                    ws[f"{col}{rr}"].value = "=" + amount(kind, a, b, col)
        else:
            for col in allcols:
                if kind == "sub":
                    f = "=" + "+".join(f"{col}{keymap[k]}" for k in a)
                elif key in ("GP", "OP"):
                    f = f"={col}{keymap[a[0]]}-{col}{keymap[a[1]]}"
                else:
                    f = f"={col}{keymap[a[0]]}+{col}{keymap[a[1]]}"
                ws[f"{col}{rr}"].value = f
            for ci in range(1, 17):
                ws.cell(rr, ci).fill = PatternFill("solid", fgColor="E7F3F5" if kind == "sub" else "C5E0E6")
                ws.cell(rr, ci).font = Font(name=FONT, bold=True, size=10)
        for col in allcols:
            cell = ws[f"{col}{rr}"]
            cell.number_format = "#,##0.00;[Red]-#,##0.00"
            cell.border = BORDER
        pc = ws[f"F{rr}"]
        pc.value = f'=IF(E${NSr}=0,"",E{rr}/E${NSr})'
        pc.number_format = "0.0%"
        pc.border = BORDER
        ch = ws[f"H{rr}"]
        ch.value = f"=E{rr}-G{rr}"
        ch.number_format = "#,##0.00;[Red]-#,##0.00"
        ch.border = BORDER
    last = plan[-1][0]
    other_row = keymap['oi']
    # margins
    m = last + 2
    ws.cell(m, 1, "هوامش الربحية (على صافي المبيعات)").font = Font(name=FONT, bold=True)
    for i, (lab, key) in enumerate([("نسبة تكلفة المبيعات (Food Cost %)", "COGS"), ("هامش مجمل الربح", "GP"), ("هامش الربح التشغيلي", "OP"), ("هامش صافي الربح", "NP")], start=1):
        ws.cell(m + i, 1, lab).font = Font(name=FONT, size=10)
        for col in allcols:
            c = ws[f"{col}{m + i}"]
            c.value = f'=IF({col}{NSr}=0,"",{col}{keymap[key]}/{col}{NSr})'
            c.number_format = "0.0%"
            c.border = BORDER
    # checks
    ck = m + 6
    ws.cell(ck, 1, "فحص: مجموع الفروع + غير الفروع = الإجمالي المجمع (صافي الربح)").font = Font(name=FONT, size=9, italic=True)
    ws.cell(ck, 5, f'=IF(ABS(SUM(I{keymap["NP"]}:O{keymap["NP"]})-P{keymap["NP"]})<=S_Tol,"✔","✖")')
    ws.cell(ck + 1, 1, "فحص: صافي ربح قائمة الدخل = (إيرادات - تكاليف - مصروفات) من الأستاذ مباشرة").font = Font(name=FONT, size=9, italic=True)
    ws.cell(ck + 1, 5, '=IF(ABS(P' + str(keymap["NP"]) + '-(' + sumifs("effcr", [("acct", '">=4000"'), ("acct", '"<=4999"'), ("date", '">="&IS_From'), ("date", '"<="&IS_To'), ("cc", "IS_CCC"), ("clos", "0")]) + '-' + sumifs("effdr", [("acct", '">=4000"'), ("acct", '"<=4999"'), ("date", '">="&IS_From'), ("date", '"<="&IS_To'), ("cc", "IS_CCC"), ("clos", "0")]) + '-(' + sumifs("effdr", [("acct", '">=5000"'), ("acct", '"<=6999"'), ("date", '">="&IS_From'), ("date", '"<="&IS_To'), ("cc", "IS_CCC"), ("clos", "0")]) + '-' + sumifs("effcr", [("acct", '">=5000"'), ("acct", '"<=6999"'), ("date", '">="&IS_From'), ("date", '"<="&IS_To'), ("cc", "IS_CCC"), ("clos", "0")]) + ')))<=S_Tol,"✔","✖")')
    name_cell(wb, "IS_NetSales", f"Income_Statement!$E${keymap['NS']}")
    name_cell(wb, "IS_GrossSales", f"Income_Statement!$E${keymap['GS']}")
    name_cell(wb, "IS_COGS", f"Income_Statement!$E${keymap['COGS']}")
    name_cell(wb, "IS_GrossProfit", f"Income_Statement!$E${keymap['GP']}")
    name_cell(wb, "IS_OpEx", f"Income_Statement!$E${keymap['OPEX']}")
    name_cell(wb, "IS_OpProfit", f"Income_Statement!$E${keymap['OP']}")
    name_cell(wb, "IS_NetProfit", f"Income_Statement!$E${keymap['NP']}")
    name_cell(wb, "IS_NetSales_P", f"Income_Statement!$G${keymap['NS']}")
    name_cell(wb, "IS_GrossProfit_P", f"Income_Statement!$G${keymap['GP']}")
    name_cell(wb, "IS_NetProfit_P", f"Income_Statement!$G${keymap['NP']}")
    name_cell(wb, "IS_OpEx_P", f"Income_Statement!$G${keymap['OPEX']}")
    ws.freeze_panes = ws.cell(H + 1, 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    protect_ws(ws)
    ws.protection.sheet = True
    # allow selector edits
    return ws, keymap, H


def build_balance_sheet(wb):
    ws = wb.create_sheet("Balance_Sheet")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    ws.cell(1, 1, "الميزانية العمومية (Balance Sheet)").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 7):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "تُسحب أرصدتها من الأستاذ العام حتى تاريخ الأرصدة. ربح الفترة الجارية = كل الإيرادات - التكاليف - المصروفات المتراكمة غير المقفلة سنويًا. الميزانية شاملة كل المواقع.").font = Font(name=FONT, size=9, italic=True)
    selector(ws, wb, "B4", "A4", "تاريخ الأرصدة", "=S_To", "BS_AsOf", "yyyy-mm-dd")
    put(ws, "A5", "حالة النتائج", bold=True, fill=C_HDR_BAND)
    put(ws, "B5", "=IS_Banner", bold=True)
    ws.column_dimensions["A"].width = 50
    for c in "BCDE":
        ws.column_dimensions[c].width = 16
    for i, h in enumerate(["البند", "من حساب", "إلى حساب", "المبلغ"], start=1):
        c = ws.cell(7, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_F_H)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    A = '"<="&BS_AsOf'
    def bal(a, b, sign):
        crits = [("acct", f'">={a}"'), ("acct", f'"<={b}"'), ("date", A)]
        if sign == "dr":
            return f'{sumifs("effdr", crits)}-{sumifs("effcr", crits)}'
        return f'{sumifs("effcr", crits)}-{sumifs("effdr", crits)}'
    rows = [
        ("hdr", "الأصول (Assets)"),
        ("l", "النقدية والخزائن", 1100, 1199, "dr"), ("l", "البنوك", 1200, 1299, "dr"), ("l", "العملاء ومستحقات المنصات والبطاقات", 1300, 1399, "dr"),
        ("l", "المخزون", 1400, 1499, "dr"), ("l", "السلف والعهد", 1500, 1599, "dr"), ("l", "المصروفات المقدمة والضرائب المدينة", 1600, 1699, "dr"),
        ("l", "الأصول الثابتة (بالتكلفة)", 1700, 1789, "dr"), ("l", "(-) مجمع الإهلاك", 1790, 1799, "dr"),
        ("t:TA", "إجمالي الأصول", None, None, None),
        ("hdr", "الالتزامات (Liabilities)"),
        ("l", "الموردون", 2100, 2199, "cr"), ("l", "المصروفات والرواتب المستحقة", 2200, 2299, "cr"), ("l", "الضرائب المستحقة", 2300, 2399, "cr"),
        ("l", "التزامات أخرى", 2400, 2499, "cr"), ("l", "الحسابات المعلقة", 2900, 2999, "cr"),
        ("t:TL", "إجمالي الالتزامات", None, None, None),
        ("hdr", "حقوق الملكية (Equity)"),
        ("l", "رأس المال", 3100, 3199, "cr"), ("l", "جاري المالك", 3200, 3299, "cr"), ("l", "أرباح محتجزة", 3300, 3399, "cr"),
        ("l", "(-) المسحوبات والتوزيعات", 3400, 3499, "cr"), ("l", "ملخص الدخل (إقفال)", 3500, 3599, "cr"),
        ("pl", "ربح (خسارة) الفترة الجارية غير المقفلة", 4000, 6999, None),
        ("t:TE", "إجمالي حقوق الملكية", None, None, None),
        ("t:TLE", "إجمالي الالتزامات وحقوق الملكية", None, None, None),
        ("chk", "الفرق (الأصول - الالتزامات - حقوق الملكية)", None, None, None),
    ]
    r = 8
    pos = {}
    start = {}
    for row in rows:
        kind = row[0]
        label = row[1]
        ws.cell(r, 1, label).font = Font(name=FONT, size=10, bold=not kind in ("l", "pl"))
        ws.cell(r, 1).border = BORDER
        if kind == "hdr":
            for ci in range(1, 5):
                ws.cell(r, ci).fill = PatternFill("solid", fgColor=C_HDR_BAND)
            cur = label
            start[label] = r + 1
        elif kind == "l":
            _, _, a, b, sg = row
            ws.cell(r, 2, a).border = BORDER
            ws.cell(r, 3, b).border = BORDER
            c = ws.cell(r, 4, "=" + bal(a, b, sg))
            c.number_format = "#,##0.00;[Red]-#,##0.00"
            c.border = BORDER
        elif kind == "pl":
            crits = [("acct", '">=4000"'), ("acct", '"<=4999"'), ("date", A)]
            crits2 = [("acct", '">=5000"'), ("acct", '"<=6999"'), ("date", A)]
            f = f'={sumifs("effcr", crits)}-{sumifs("effdr", crits)}-({sumifs("effdr", crits2)}-{sumifs("effcr", crits2)})'
            c = ws.cell(r, 4, f)
            c.number_format = "#,##0.00;[Red]-#,##0.00"
            c.border = BORDER
        else:
            tag = kind.split(":")[1] if ":" in kind else kind
            pos[tag] = r
            for ci in range(1, 5):
                ws.cell(r, ci).fill = PatternFill("solid", fgColor="E7F3F5")
            ws.cell(r, 1).font = Font(name=FONT, size=10, bold=True)
        r += 1
    # totals
    def rng(label_hdr, tot_tag):
        return f"D{start[label_hdr]}:D{pos[tot_tag] - 1}"
    ws.cell(pos["TA"], 4, f"=SUM({rng('الأصول (Assets)', 'TA')})")
    ws.cell(pos["TL"], 4, f"=SUM({rng('الالتزامات (Liabilities)', 'TL')})")
    ws.cell(pos["TE"], 4, f"=SUM({rng('حقوق الملكية (Equity)', 'TE')})")
    ws.cell(pos["TLE"], 4, f"=D{pos['TL']}+D{pos['TE']}")
    ws.cell(pos["chk"], 4, f"=ROUND(D{pos['TA']}-D{pos['TLE']},2)")
    for t in ("TA", "TL", "TE", "TLE", "chk"):
        ws.cell(pos[t], 4).number_format = "#,##0.00;[Red]-#,##0.00"
        ws.cell(pos[t], 4).font = Font(name=FONT, bold=True)
        ws.cell(pos[t], 4).border = BORDER
    ws.cell(pos["chk"], 5, f'=IF(ABS(D{pos["chk"]})<=S_Tol,"✔ الميزانية متوازنة","✖ الميزانية غير متوازنة - يمنع الإقفال")')
    name_cell(wb, "BS_Diff", f"Balance_Sheet!$D${pos['chk']}")
    name_cell(wb, "BS_TotalAssets", f"Balance_Sheet!$D${pos['TA']}")
    name_cell(wb, "BS_TotalLiab", f"Balance_Sheet!$D${pos['TL']}")
    name_cell(wb, "BS_TotalEquity", f"Balance_Sheet!$D${pos['TE']}")
    ws.freeze_panes = "A8"
    protect_ws(ws)
    return ws


def build_cash_flow(wb):
    ws = wb.create_sheet("Cash_Flow")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "BF8F00"
    ws.cell(1, 1, "قائمة التدفقات النقدية (الطريقة المباشرة المبسطة)").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 9):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "كل سطر نقدي (خزائن وبنوك) يحمل تصنيف تدفق من مستنده. عمود (التحويلات الداخلية) يجب أن يساوي صفرًا؛ أي رقم فيه يعني تحويلًا غير مكتمل.").font = Font(name=FONT, size=9, italic=True)
    selector(ws, wb, "B4", "A4", "من تاريخ", "=S_From", "CF_From", "yyyy-mm-dd")
    selector(ws, wb, "B5", "A5", "إلى تاريخ", "=S_To", "CF_To", "yyyy-mm-dd")
    ws.column_dimensions["A"].width = 44
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 16
    heads = ["مصدر الحركة", "تشغيلية", "استثمارية", "تمويلية", "تحويلات داخلية (=0)", "رصيد افتتاحي مرحّل", "الإجمالي"]
    classes = ["OPERATING", "INVESTING", "FINANCING", "INTERNAL", "OPENING"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(7, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_F_H)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    srcs = [("SAL", "مقبوضات المبيعات النقدية والتحويلات المباشرة"), ("PLT", "تحصيل تسويات المنصات والبطاقات"), ("SINV", "مشتريات مدفوعة فورًا"),
            ("SPAY", "دفعات الموردين"), ("EXP", "مصروفات مدفوعة فورًا"), ("CV", "سندات الخزينة (قبض/صرف/تحويل/عجز)"), ("MJE", "قيود يدوية تمس النقدية (افتتاحي/تسويات)")]
    r = 8
    for code, lab in srcs:
        ws.cell(r, 1, lab).border = BORDER
        for j, cl in enumerate(classes, start=2):
            crits = [("cash", "1"), ("stype", f'"{code}"'), ("cfc", f'"{cl}"'), ("date", '">="&CF_From'), ("date", '"<="&CF_To')]
            c = ws.cell(r, j, "=" + netdr(crits))
            c.number_format = "#,##0.00;[Red]-#,##0.00"
            c.border = BORDER
        c = ws.cell(r, 7, f"=SUM(B{r}:F{r})")
        c.number_format = "#,##0.00;[Red]-#,##0.00"
        c.border = BORDER
        r += 1
    ws.cell(r, 1, "صافي التدفق").font = Font(name=FONT, bold=True)
    for j in range(2, 8):
        c = ws.cell(r, j, f"=SUM({GL(j)}8:{GL(j)}{r - 1})")
        c.number_format = "#,##0.00;[Red]-#,##0.00"
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="E7F3F5")
        c.border = BORDER
    tot = r
    ws.cell(r + 2, 1, "رصيد النقدية والبنوك أول الفترة").font = Font(name=FONT, bold=True)
    c = ws.cell(r + 2, 2, "=" + netdr([("cash", "1"), ("date", '"<"&CF_From')]))
    ws.cell(r + 3, 1, "(+) صافي التدفق خلال الفترة").font = Font(name=FONT, bold=True)
    ws.cell(r + 3, 2, f"=G{tot}")
    ws.cell(r + 4, 1, "رصيد النقدية والبنوك آخر الفترة (محسوب)").font = Font(name=FONT, bold=True)
    ws.cell(r + 4, 2, f"=B{r + 2}+B{r + 3}")
    ws.cell(r + 5, 1, "رصيد النقدية والبنوك من الأستاذ العام").font = Font(name=FONT, bold=True)
    ws.cell(r + 5, 2, "=" + netdr([("cash", "1"), ("date", '"<="&CF_To')]))
    ws.cell(r + 6, 1, "الفرق").font = Font(name=FONT, bold=True)
    ws.cell(r + 6, 2, f"=ROUND(B{r + 4}-B{r + 5},2)")
    ws.cell(r + 6, 3, f'=IF(ABS(B{r + 6})<=S_Tol,"✔ متطابق","✖ فرق")')
    ws.cell(r + 7, 1, "فحص: التحويلات الداخلية صافيها صفر").font = Font(name=FONT, bold=True)
    ws.cell(r + 7, 2, f"=E{tot}")
    ws.cell(r + 7, 3, f'=IF(ABS(B{r + 7})<=S_Tol,"✔","✖ تحويل غير مكتمل")')
    for k in range(2, 8):
        ws.cell(r + k, 2).number_format = "#,##0.00;[Red]-#,##0.00"
        ws.cell(r + k, 2).border = BORDER
    name_cell(wb, "CF_Diff", f"Cash_Flow!$B${r + 6}")
    name_cell(wb, "CF_InternalNet", f"Cash_Flow!$B${r + 7}")
    name_cell(wb, "CF_Net", f"Cash_Flow!$G${tot}")
    name_cell(wb, "CF_Operating", f"Cash_Flow!$B${tot}")
    ws.freeze_panes = "B8"
    protect_ws(ws)
    return ws


# =================================================================== Treasury / Suppliers / Customers
def define_treasury_balances(wb):
    tr = Sheet("Treasury_Balances", "أرصدة الخزائن والبنوك وحركة النقدية (Cash & Bank Balances)",
               "رصيد كل خزينة = الرصيد الافتتاحي + المقبوضات - المدفوعات (من الأستاذ العام). العمود الأخير يعرض حركة يوم محدد.",
               hdr=8, cap=CAP_TRE, tab="BF8F00")

    def top(ws):
        selector(ws, wb, "B4", "A4", "من تاريخ", "=S_From", "TR_From", "yyyy-mm-dd")
        selector(ws, wb, "B5", "A5", "إلى تاريخ", "=S_To", "TR_To", "yyyy-mm-dd")
        selector(ws, wb, "B6", "A6", "يوم الحركة اليومية", "=S_To", "TR_Day", "yyyy-mm-dd")
        put(ws, "D4", "إجمالي النقدية والبنوك", bold=True, fill=C_HDR_BAND)
        put(ws, "E4", f'=SUM({tr.rng("close")})', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="TR_TotalCash", wb=wb)
        put(ws, "D5", "عدد خزائن برصيد سالب", bold=True, fill=C_HDR_BAND)
        put(ws, "E5", f'=COUNTIFS({tr.rng("close")},"<0")', fill=C_F, name="TR_Negative", wb=wb)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["D"].width = 26
        ws.column_dimensions["E"].width = 17
    tr.extra.append(top)
    off = tr.first - 5
    tL = lambda k: REG["Treasuries"].L(k)
    tr.add("code", "الخزينة/البنك", "f", 14, align="center", f=lambda x: f'=IF(INDEX(Treasuries!$A:$A,ROW()-{off})="","",INDEX(Treasuries!$A:$A,ROW()-{off}))')
    tr.add("name", "الاسم", "f", 32, f=lambda x: f'=IF({x("code")}="","",INDEX(Treasuries!${tL("name")}:${tL("name")},ROW()-{off}))')
    tr.add("type", "النوع", "f", 11, f=lambda x: f'=IF({x("code")}="","",INDEX(Treasuries!${tL("type")}:${tL("type")},ROW()-{off}))')
    tr.add("loc", "الموقع", "f", 8, f=lambda x: f'=IF({x("code")}="","",INDEX(Treasuries!${tL("loc")}:${tL("loc")},ROW()-{off}))', align="center")
    tr.add("gl", "الحساب", "f", 8, f=lambda x: f'=IF({x("code")}="","",INDEX(Treasuries!${tL("gl")}:${tL("gl")},ROW()-{off}))', align="center")
    mk = lambda extra: (lambda x: f'=IF({x("code")}="","",' + netdr([("acct", x("gl"))] + extra) + ')')
    tr.add("open", "الرصيد الافتتاحي", "f", 15, fmt="#,##0.00;[Red]-#,##0.00", f=mk([("date", chr(34)+"<"+chr(34)+"&TR_From")]))
    tr.add("rec", "المقبوضات", "f", 15, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effdr", [("acct", x("gl")), ("date", chr(34)+">="+chr(34)+"&TR_From"), ("date", chr(34)+"<="+chr(34)+"&TR_To")]) + ')')
    tr.add("pay", "المدفوعات", "f", 15, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effcr", [("acct", x("gl")), ("date", chr(34)+">="+chr(34)+"&TR_From"), ("date", chr(34)+"<="+chr(34)+"&TR_To")]) + ')')
    tr.add("close", "الرصيد الختامي", "f", 15, fmt="#,##0.00;[Red]-#,##0.00", f=lambda x: f'=IF({x("code")}="","",{x("open")}+{x("rec")}-{x("pay")})')
    tr.add("dayin", "مقبوضات اليوم المحدد", "f", 15, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effdr", [("acct", x("gl")), ("date", "TR_Day")]) + ')')
    tr.add("dayout", "مدفوعات اليوم المحدد", "f", 15, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effcr", [("acct", x("gl")), ("date", "TR_Day")]) + ')')
    tr.add("flag", "تنبيه", "f", 24, f=lambda x: f'=IF({x("code")}="","",IF({x("close")}<-S_Tol,"رصيد سالب - تحقق من المستندات",""))')
    return tr


def define_supplier_balances(wb):
    sb = Sheet("Supplier_Balances", "أرصدة الموردين وأعمار الديون (Supplier Balances & Aging)",
               "رصيد المورد = الالتزامات المعتمدة - المرتجعات والتسويات والدفعات (حساب 2101 مع كود المورد). أعمار الديون بافتراض سداد الأقدم أولًا (FIFO). لتفاصيل مورد: افتح GL_Ledger واختر 2101 وكود المورد.",
               hdr=10, cap=CAP_SUP, tab="7030A0")

    def top(ws):
        selector(ws, wb, "B4", "A4", "من تاريخ", "=S_From", "SB_From", "yyyy-mm-dd")
        selector(ws, wb, "B5", "A5", "تاريخ الأرصدة والأعمار", "=S_To", "SB_To", "yyyy-mm-dd")
        put(ws, "D4", "إجمالي أرصدة الموردين (دفتر فرعي)", bold=True, fill=C_HDR_BAND)
        put(ws, "E4", f'=SUM({sb.rng("close")})', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="SB_Total", wb=wb)
        put(ws, "D5", "رصيد حساب الموردين 2101 في الأستاذ", bold=True, fill=C_HDR_BAND)
        put(ws, "E5", "=" + f'({sumifs("effcr", [("acct", "ACC_SUPPLIERS"), ("date", chr(34)+"<="+chr(34)+"&SB_To")])}-{sumifs("effdr", [("acct", "ACC_SUPPLIERS"), ("date", chr(34)+"<="+chr(34)+"&SB_To")])})', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="SB_GL", wb=wb)
        put(ws, "D6", "الفرق (الدفتر الفرعي - الأستاذ)", bold=True, fill=C_HDR_BAND)
        put(ws, "E6", "=ROUND(E4-E5,2)", fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="SB_Diff", wb=wb)
        put(ws, "F6", '=IF(ABS(E6)<=S_Tol,"✔ متطابق","✖ حركات على 2101 بدون كود مورد صحيح")', bold=True)
        put(ws, "D7", "أرصدة موردين مدينة (مراجعة)", bold=True, fill=C_HDR_BAND)
        put(ws, "E7", f'=COUNTIFS({sb.rng("close")},"<0")', fill=C_F, name="SB_Negative", wb=wb)
        put(ws, "D8", "مبالغ متأخرة (أكثر من 0 يوم)", bold=True, fill=C_HDR_BAND)
        put(ws, "E8", f'=SUM({sb.rng("b30")})+SUM({sb.rng("b60")})+SUM({sb.rng("b90")})+SUM({sb.rng("b90p")})', fmt="#,##0.00", fill=C_F, name="SB_Overdue", wb=wb)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["E"].width = 18
    sb.extra.append(top)
    off = sb.first - 5
    sL = lambda k: REG["Suppliers"].L(k)
    sb.add("code", "كود المورد", "f", 10, align="center", f=lambda x: f'=IF(INDEX(Suppliers!$A:$A,ROW()-{off})="","",INDEX(Suppliers!$A:$A,ROW()-{off}))')
    sb.add("name", "اسم المورد", "f", 30, f=lambda x: f'=IF({x("code")}="","",INDEX(Suppliers!${sL("name")}:${sL("name")},ROW()-{off}))')
    cond = lambda x, extra: [("acct", "ACC_SUPPLIERS"), ("party", x("code"))] + extra
    sb.add("open", "رصيد أول الفترة", "f", 14, fmt="#,##0.00;[Red]-#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + f'{sumifs("effcr", cond(x, [("date", chr(34)+"<"+chr(34)+"&SB_From")]))}-{sumifs("effdr", cond(x, [("date", chr(34)+"<"+chr(34)+"&SB_From")]))}' + ')')
    sb.add("cr", "فواتير والتزامات (دائن)", "f", 14, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effcr", cond(x, [("date", chr(34)+">="+chr(34)+"&SB_From"), ("date", chr(34)+"<="+chr(34)+"&SB_To")])) + ')')
    sb.add("dr", "دفعات ومرتجعات وتسويات (مدين)", "f", 14, fmt="#,##0.00",
           f=lambda x: f'=IF({x("code")}="","",' + sumifs("effdr", cond(x, [("date", chr(34)+">="+chr(34)+"&SB_From"), ("date", chr(34)+"<="+chr(34)+"&SB_To")])) + ')')
    sb.add("close", "رصيد المورد (دائن+)", "f", 14, fmt="#,##0.00;[Red]-#,##0.00", f=lambda x: f'=IF({x("code")}="","",{x("open")}+{x("cr")}-{x("dr")})')
    SI = lambda k: REG["Supplier_Invoices"].ws_col(k)

    def cum(x, days):
        return f'SUMIFS({SI("due_amt")},{SI("sup")},{x("code")},{SI("ok")},1,{SI("due")},">="&(SB_To-{days}),{SI("date")},"<="&SB_To)'
    B = lambda x: f'MAX({x("close")},0)'
    sb.add("b0", "غير مستحق", "f", 13, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MIN({B(x)},{cum(x, 0)}))')
    sb.add("b30", "متأخر 1-30", "f", 13, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MIN({B(x)},{cum(x, 30)})-MIN({B(x)},{cum(x, 0)}))')
    sb.add("b60", "متأخر 31-60", "f", 13, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MIN({B(x)},{cum(x, 60)})-MIN({B(x)},{cum(x, 30)}))')
    sb.add("b90", "متأخر 61-90", "f", 13, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",MIN({B(x)},{cum(x, 90)})-MIN({B(x)},{cum(x, 60)}))')
    sb.add("b90p", "متأخر أكثر من 90 (وأرصدة افتتاحية)", "f", 14, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",{B(x)}-MIN({B(x)},{cum(x, 90)}))')
    SP = lambda k: REG["Supplier_Payments"].ws_col(k)
    sb.add("lastpay", "آخر دفعة", "f", 12, fmt="yyyy-mm-dd", f=lambda x: f'=IF({x("code")}="","",IF(COUNTIFS({SP("sup")},{x("code")},{SP("ok")},1)=0,"",_xlfn.MAXIFS({SP("date")},{SP("sup")},{x("code")},{SP("ok")},1)))')
    sb.add("flag", "تنبيه", "f", 30, f=lambda x: f'=IF({x("code")}="","",IF({x("close")}<-S_Tol,"رصيد مدين للمورد - مراجعة",IF({x("b30")}+{x("b60")}+{x("b90")}+{x("b90p")}>S_Tol,"يوجد مبالغ متأخرة","")))')
    return sb


def define_customer_balances(wb):
    cb = Sheet("Customer_Balances", "أرصدة العملاء (Customer Balances)", "رصيد العميل = حساب 1301 مع كود العميل (فواتير العملاء الآجلين تسجل بقيد يدوي في المرحلة 1). مبيعات الفروع والمنصات لا تظهر هنا.",
               hdr=8, cap=CAP_CUS, tab="548235")

    def top(ws):
        selector(ws, wb, "B4", "A4", "تاريخ الأرصدة", "=S_To", "CB_To", "yyyy-mm-dd")
        put(ws, "D4", "إجمالي أرصدة العملاء (دفتر فرعي)", bold=True, fill=C_HDR_BAND)
        put(ws, "E4", f'=SUM({cb.rng("bal")})', fmt="#,##0.00;[Red]-#,##0.00", fill=C_F)
        put(ws, "D5", "رصيد حساب العملاء 1301 في الأستاذ", bold=True, fill=C_HDR_BAND)
        put(ws, "E5", "=" + netdr([("acct", "ACC_CUSTOMERS"), ("date", chr(34)+"<="+chr(34)+"&CB_To")]), fmt="#,##0.00;[Red]-#,##0.00", fill=C_F)
        put(ws, "D6", "الفرق", bold=True, fill=C_HDR_BAND)
        put(ws, "E6", "=ROUND(E4-E5,2)", fmt="#,##0.00;[Red]-#,##0.00", fill=C_F, name="CB_Diff", wb=wb)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["E"].width = 18
    cb.extra.append(top)
    off = cb.first - 5
    cL = lambda k: REG["Customers"].L(k)
    cb.add("code", "كود العميل", "f", 10, align="center", f=lambda x: f'=IF(INDEX(Customers!$A:$A,ROW()-{off})="","",INDEX(Customers!$A:$A,ROW()-{off}))')
    cb.add("name", "اسم العميل", "f", 30, f=lambda x: f'=IF({x("code")}="","",INDEX(Customers!${cL("name")}:${cL("name")},ROW()-{off}))')
    cb.add("limit", "حد الائتمان", "f", 13, fmt="#,##0.00", f=lambda x: f'=IF({x("code")}="","",INDEX(Customers!${cL("limit")}:${cL("limit")},ROW()-{off}))')
    cb.add("bal", "الرصيد (مدين+)", "f", 14, fmt="#,##0.00;[Red]-#,##0.00", f=lambda x: f'=IF({x("code")}="","",' + netdr([("acct", "ACC_CUSTOMERS"), ("party", x("code")), ("date", chr(34)+"<="+chr(34)+"&CB_To")]) + ')')
    cb.add("flag", "تنبيه", "f", 28, f=lambda x: f'=IF({x("code")}="","",IF(AND(ISNUMBER({x("limit")}),{x("limit")}>0,{x("bal")}>{x("limit")}),"تجاوز حد الائتمان",IF({x("bal")}<-S_Tol,"رصيد دائن للعميل","")))')
    return cb
