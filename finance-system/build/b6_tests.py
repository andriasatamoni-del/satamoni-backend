# -*- coding: utf-8 -*-
"""ورقة Test_Cases: اختبارات حية بمعادلات تقارن النتيجة الفعلية بنتيجة متوقعة محسوبة مستقلًا في Python."""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from fs_common import *
from b1_masters import G
from b3_ledger import put, sumifs, netdr
import datetime as dt


def n(x):
    x = round(float(x), 2)
    return str(int(x)) if x == int(x) else (f"{x:.2f}".rstrip("0").rstrip("."))


def R(sheet, key):
    return REG[sheet].ws_col(key)


def eff(side, doc=None, acct=None, party=None, ge=None, le=None, date_le=None, loc=None, stype=None):
    cr = []
    if doc is not None:
        cr.append(("doc", f'"{doc}"'))
    if acct is not None:
        cr.append(("acct", str(acct)))
    if ge is not None:
        cr.append(("acct", f'">={ge}"'))
    if le is not None:
        cr.append(("acct", f'"<={le}"'))
    if party is not None:
        cr.append(("party", f'"{party}"'))
    if loc is not None:
        cr.append(("loc", f'"{loc}"'))
    if stype is not None:
        cr.append(("stype", f'"{stype}"'))
    if date_le is not None:
        cr.append(("date", f'"<="&DATE({date_le.year},{date_le.month},{date_le.day})'))
    if not cr:
        cr.append(("act", "1"))
    return sumifs("effdr" if side == "D" else "effcr", cr)


def rnd(x):
    return f"ROUND({x},2)"


def build_tests(wb, E, with_br04, keymap_is, H_is):
    ws = wb.create_sheet("Test_Cases")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    ws.cell(1, 1, "اختبارات الجودة (Test_Cases) - معادلات حية تقارن الفعلي بالمتوقع").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 9):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=C_TITLE)
    ws.cell(2, 1, "عمود (النتيجة الفعلية) معادلة تقرأ من الدفاتر الآن؛ و(المتوقعة) محسوبة مستقلًا خارج الجدول من مدخلات المستندات. تفترض الاختبارات بيانات الملف التجريبي كما هي وفترة Settings = أكتوبر 2026. حالات المخزون والإنتاج مؤجلة للمرحلة 2.").font = Font(name=FONT, size=9, italic=True)
    for col, w in zip("ABCDEFGH", [7, 30, 46, 40, 40, 14, 18, 40]):
        ws.column_dimensions[col].width = w
    heads = ["رقم", "الاختبار", "السيناريو والمستندات المستخدمة", "النتيجة المتوقعة", "النتيجة الفعلية (معادلة حية)", "الحالة", "طريقة القياس", "ملاحظات"]
    H = 7
    for i, h in enumerate(heads, start=1):
        c = ws.cell(H, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_IN_H)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    SR = E["sales_rows"]
    r1 = SR["SAL-BR01-0001"]
    br04 = SR.get("SAL-BR04-0001")
    D1 = dt.date(2026, 10, 31)
    pipe = lambda *parts: '&"|"&'.join(parts)
    tests = []

    def add(no, title, scen, expected, actual, how="معادلة حية", note="", deferred=False):
        tests.append(dict(no=no, title=title, scen=scen, expected=expected, actual=actual, how=how, note=note, deferred=deferred))

    add("T01", "تسجيل عملية شراء نقدية", "SINV-HQ01-0003: عبوات 4,200 من SUP003 مدفوعة فورًا من CSH-HQ01",
        "4200|4200|0",
        "=" + pipe(rnd(eff("D", "SINV-HQ01-0003", 1405)), rnd(eff("K", "SINV-HQ01-0003", 1104)), rnd(f'{eff("K", "SINV-HQ01-0003", 2101)}-{eff("D", "SINV-HQ01-0003", 2101)}')),
        note="مدين مخزون العبوات | دائن خزينة الإدارة | أثر المورد صفر")
    add("T02", "تسجيل شراء آجل", "SINV-WH01-0001: أجبان 30,000 من SUP001 (استحقاق 30 يومًا)",
        "30000|30000|48000",
        "=" + pipe(rnd(eff("D", "SINV-WH01-0001", 1401)), rnd(eff("K", "SINV-WH01-0001", 2101, party="SUP001")), rnd(f'{eff("K", acct=2101, party="SUP001", date_le=D1)}-{eff("D", acct=2101, party="SUP001", date_le=D1)}')),
        note="الرصيد النهائي للمورد: 30,000 افتتاحي + 30,000 - 2,000 مرتجع - 10,000 دفعة = 48,000")
    add("T03", "تسجيل بيع نقدي وبيع عبر منصة", f"SAL-BR01-0001 (10 سبتمبر): نقدي {n(r1['m1'])} + بطاقات {n(r1['m2'])} + طلبات {n(r1['m3'])}",
        f"{n(r1['m1'])}|{n(r1['m2'])}|{n(r1['m3'])}|{n(r1['_gross'])}|{n(r1['disc'])}",
        "=" + pipe(rnd(eff("D", "SAL-BR01-0001", 1101)), rnd(eff("D", "SAL-BR01-0001", 1302)), rnd(eff("D", "SAL-BR01-0001", 1303)), rnd(eff("K", "SAL-BR01-0001", ge=4101, le=4104)), rnd(eff("D", "SAL-BR01-0001", 4901))),
        note="مدين خزينة الفرع | مستحقات بطاقات | مستحقات طلبات | دائن إجمالي الإيراد | مدين الخصومات")
    add("T04", "تسوية عمولة المنصة", "PLT-HQ01-0003: مستحقات 15,000 - عمولة 2,700 - خصومات 150 = متوقع 12,150؛ المستلم 11,800",
        "11800|2700|150|350|15000",
        "=" + pipe(rnd(eff("D", "PLT-HQ01-0003", 1201)), rnd(eff("D", "PLT-HQ01-0003", 6701)), rnd(eff("D", "PLT-HQ01-0003", 6705)), rnd(eff("D", "PLT-HQ01-0003", 1309)), rnd(eff("K", "PLT-HQ01-0003", 1303))),
        note="الفرق 350 في حساب مراجعة 1309 وليس مصروفًا؛ العمولة منفصلة عن خصم العميل")
    for no, t, sc in [("T05", "تحويل مخزون بين موقعين", "إذن تحويل بين المخازن (نموذج 5)"),
                      ("T06", "تحويل إنتاج من السنتر كيتشن إلى فرع", "إذن تحويل إنتاج (نموذج 8)"),
                      ("T07", "صرف مواد للإنتاج", "طلب صرف إنتاج (نموذج 6)"),
                      ("T08", "تسجيل هالك", "تقرير هالك (نموذج 9)")]:
        add(no, t, sc, "مؤجل", "مؤجل - المرحلة 2 (Inventory_Moves/Production)", how="مؤجل", note="هيكل الأوراق جاهز؛ القيود النموذجية في JE_Guide (أمثلة 14-18)", deferred=True)
    add("T09", "تسجيل مصروف مستحق ثم سداده", "EXP-BR01-0003 كهرباء 2,400 آجل (سبتمبر) سُدد بـ PAY-HQ01-0005 في أكتوبر؛ EXP-HQ01-0003 إنترنت 1,500 غير مسدد",
        "2400|2400|مدفوع بالكامل|1500",
        "=" + pipe(rnd(eff("D", "EXP-BR01-0003", 6301)), rnd(eff("K", "EXP-BR01-0003", 2201)),
                   f'INDEX({R("Expenses","paystat")},MATCH("EXP-BR01-0003",{R("Expenses","doc")},0))',
                   rnd(f'INDEX({R("Expenses","remain")},MATCH("EXP-HQ01-0003",{R("Expenses","doc")},0))')),
        note="فصل تاريخ الإثبات (سبتمبر) عن تاريخ الدفع (10 أكتوبر)")
    add("T10", "تسجيل دفعة مورد", "SPAY-HQ01-0002: سداد 12,500 للمورد SUP002 من البنك",
        "12500|12500|12000",
        "=" + pipe(rnd(eff("D", "SPAY-HQ01-0002", 2101)), rnd(eff("K", "SPAY-HQ01-0002", 1201)), rnd(f'{eff("K", acct=2101, party="SUP002", date_le=D1)}-{eff("D", acct=2101, party="SUP002", date_le=D1)}')),
        note="رصيد SUP002: 12,000 افتتاحي + 12,500 فاتورة - 12,500 دفعة = 12,000")
    add("T11", "تسجيل مرتجع مشتريات", "SINV-WH01-0005: إشعار خصم 2,000 من SUP001",
        "2000|2000",
        "=" + pipe(rnd(eff("D", "SINV-WH01-0005", 2101)), rnd(eff("K", "SINV-WH01-0005", 1401))), note="ينقص المورد والمخزون")
    add("T12", "إلغاء عملية وتصحيحها", "EXP-HQ01-0004 (5,000 خطأ) معكوس ← EXP-HQ01-0005 (500 صحيح)",
        "0|500|معكوس",
        "=" + pipe(rnd(f'SUMIFS({G("effdr")},{G("doc")},"EXP-HQ01-0004")'), rnd(eff("D", "EXP-HQ01-0005", 6501)),
                   f'INDEX({R("Expenses","status")},MATCH("EXP-HQ01-0004",{R("Expenses","doc")},0))'),
        note="المستند الأول بلا أثر ويبقى في السجل للمراجعة")
    if with_br04:
        ns = br04["_gross"] - br04["disc"] - br04["ret"]
        add("T13", "إضافة فرع جديد تجريبي (BR04)", f"BR04 في Locations مع خزينة CSH-BR04 وحساب 1107 ومبيعات SAL-BR04-0001 (صافي {n(ns)})",
            f"4|1|{n(ns)}",
            "=" + pipe(f'INDEX({R("Locations","seq")},MATCH("BR04",{R("Locations","code")},0))',
                       f'COUNTIF(Income_Statement!$I${H_is}:$N${H_is},"BR04")',
                       rnd(f'{eff("K", ge=4101, le=4104, loc="BR04")}-{eff("D", ge=4901, le=4902, loc="BR04")}')),
            note="يظهر الفرع آليًا في أعمدة قائمة الدخل ولوحة التحكم دون تعديل معادلات")
    else:
        add("T13", "إضافة فرع جديد تجريبي", "غير مفعّل في النسخة النظيفة", "مؤجل", "—", how="مؤجل", deferred=True)
    add("T14", "اكتشاف مستند مكرر", "SINV-WH01-0007 نسخة من فاتورة V-0777 للمورد SUP005 (نفس الرقم)",
        "فاتورة مكررة (نفس المورد ونفس الرقم)|0",
        "=" + pipe(f'INDEX({R("Supplier_Invoices","chk")},MATCH("SINV-WH01-0007",{R("Supplier_Invoices","doc")},0))', rnd(f'SUMIFS({G("effdr")},{G("doc")},"SINV-WH01-0007")')),
        note="المستند الأصلي SINV-WH01-0006 يبقى سليمًا")
    add("T15", "محاولة ترحيل قيد غير متوازن", "MJE-HQ01-0008: مدين 500 / دائن 300 بحالة (مرحل)",
        "القيد غير متوازن (مدين ≠ دائن) - لن يُرحَّل|0",
        "=" + pipe(f'INDEX({R("Manual_JE","chk")},MATCH("MJE-HQ01-0008",{R("Manual_JE","doc")},0))', rnd(f'SUMIFS({G("effdr")},{G("doc")},"MJE-HQ01-0008")')),
        note="الأسطر تظهر في GL_Journal بأثر فعلي = 0 ويظهر البند في Close_Checklist")
    add("T16", "اكتشاف فرق خزينة وتسويته", "Cash_Close: عجز 150 في خزينة BR02 بتاريخ 12/10 مسوّى بسند ADJ-BR02-0002",
        "-150|0|150",
        "=" + pipe(rnd(f'INDEX({R("Cash_Close","diff")},MATCH("ADJ-BR02-0002",{R("Cash_Close","adj")},0))'), rnd(f'INDEX({R("Cash_Close","resid")},MATCH("ADJ-BR02-0002",{R("Cash_Close","adj")},0))'), rnd(eff("D", "ADJ-BR02-0002", 6910))),
        note="الفرق مكتشف ثم متبقٍ صفر بعد سند العجز")
    add("T17", "اكتشاف فرق جرد مخزون", "محضر جرد ← تسوية مخزون", "مؤجل", "مؤجل - المرحلة 2", how="مؤجل", note="المثال 19 في JE_Guide", deferred=True)
    add("T18", "إقفال شهر تجريبي", "سبتمبر 2026 مقفل بلقطة قيم؛ أكتوبر غير قابل للإقفال لوجود أخطاء مزروعة",
        "مقفلة||blocked",
        "=" + pipe(f'INDEX({R("Period_Close","status")},MATCH(202609,{R("Period_Close","period")},0))',
                   f'INDEX({R("Period_Close","chg")},MATCH(202609,{R("Period_Close","period")},0))&""', 'IF(CK_Critical>0,"blocked","open")'),
        note="تشغيل آلي مسجل: عند Settings=سبتمبر 2026 وتأكيد بنود M01-M14 كانت البنود الحرجة = 0 (انظر docs/09_Test_Run_Log.md)")

    ints = [
        ("V01", "ميزان المراجعة: إجمالي المدين = إجمالي الدائن", "Trial_Balance", "0", "=" + rnd("TB_Diff")),
        ("V02", "الأستاذ العام: مجموع المدين المرحّل = الدائن المرحّل", "GL_Journal", "0", "=" + rnd("GL_TotalDiff")),
        ("V03", "الميزانية متوازنة (أصول = التزامات + ملكية)", "Balance_Sheet", "0", "=" + rnd("BS_Diff")),
        ("V04", "دفتر الموردين الفرعي = حساب الموردين 2101", "Supplier_Balances", "0", "=" + rnd("SB_Diff")),
        ("V05", "التدفقات النقدية تطابق أرصدة النقدية، والتحويلات الداخلية صافيها صفر", "Cash_Flow", "0|0", "=" + pipe(rnd("CF_Diff"), rnd("CF_InternalNet"))),
        ("V06", "صافي ربح قائمة الدخل = (إيرادات - تكاليف - مصروفات) من الأستاذ", "Income_Statement",
         "0", "=" + rnd(f'IS_NetProfit-(({sumifs("effcr", [("acct", chr(34)+">=4000"+chr(34)), ("acct", chr(34)+"<=4999"+chr(34)), ("date", chr(34)+">="+chr(34)+"&S_From"), ("date", chr(34)+"<="+chr(34)+"&S_To"), ("clos", "0")])})-({sumifs("effdr", [("acct", chr(34)+">=4000"+chr(34)), ("acct", chr(34)+"<=4999"+chr(34)), ("date", chr(34)+">="+chr(34)+"&S_From"), ("date", chr(34)+"<="+chr(34)+"&S_To")])})-(({sumifs("effdr", [("acct", chr(34)+">=5000"+chr(34)), ("acct", chr(34)+"<=6999"+chr(34)), ("date", chr(34)+">="+chr(34)+"&S_From"), ("date", chr(34)+"<="+chr(34)+"&S_To")])})-({sumifs("effcr", [("acct", chr(34)+">=5000"+chr(34)), ("acct", chr(34)+"<=6999"+chr(34)), ("date", chr(34)+">="+chr(34)+"&S_From"), ("date", chr(34)+"<="+chr(34)+"&S_To")])})))')),
        ("V07", "لا إيراد من مصدر غير المبيعات (التحويلات ليست إيرادًا)", "GL_Journal", "0",
         "=" + "+".join([f'COUNTIFS({G("eff")},1,{G("acct")},">=4101",{G("acct")},"<=4104",{G("stype")},"{t}")' for t in ["CV", "EXP", "SINV", "SPAY", "PLT"]])),
        ("V08", "رصيد الخزائن والبنوك في التقرير = رصيدها في الأستاذ", "Treasury_Balances", "0", "=" + rnd(f'TR_TotalCash-({sumifs("effdr", [("cash", "1"), ("date", chr(34)+"<="+chr(34)+"&S_To")])}-{sumifs("effcr", [("cash", "1"), ("date", chr(34)+"<="+chr(34)+"&S_To")])})')),
        ("V09", "إجمالي أصول الميزانية = مجموع أرصدة ميزان المراجعة للأصول", "Balance_Sheet/Trial_Balance", "0",
         "=" + rnd(f'BS_TotalAssets-(SUMIFS({R("Trial_Balance","close")},{R("Trial_Balance","post")},"نعم",{R("Trial_Balance","code")},">=1000",{R("Trial_Balance","code")},"<=1999"))')),
        ("V10", "عدم ازدواج المبيعات: إيراد الأستاذ = إجمالي تقارير المبيعات المرحلة", "GL vs Sales_Daily", "0",
         "=" + rnd(f'{eff("K", ge=4101, le=4104)}-SUMIFS({R("Sales_Daily","gross")},{R("Sales_Daily","ok")},1)')),
        ("V11", "عدم ازدواج تكلفة المبيعات: لا أسطر 5101 من غير القيود اليدوية (طريقة دورية)", "GL_Journal", "0",
         "=" + f'COUNTIFS({G("eff")},1,{G("acct")},5101,{G("stype")},"<>MJE")'),
        ("V12", "رصيد الصندوق الافتتاحي + الحركات = الختامي (جميع الخزائن)", "Treasury_Balances", "0", "=" + rnd(f'SUM({R("Treasury_Balances","open")})+SUM({R("Treasury_Balances","rec")})-SUM({R("Treasury_Balances","pay")})-SUM({R("Treasury_Balances","close")})')),
        ("V13", "أرصدة المخزون تطابق سجل الحركات", "Inventory_Moves", "مؤجل", "مؤجل - المرحلة 2"),
        ("V14", "وضوح التنبيهات: وجود بنود حرجة ظاهرة في الملف التجريبي (أخطاء الاختبار 14 و15)", "Close_Checklist/Dashboard", "نعم", '=IF(CK_Critical>0,"نعم","لا")'),
    ]
    def txt(a):
        return a if (a.startswith('=IF(') or '&"|"&' in a or a.startswith('="') or a.startswith("مؤجل")) else a + '&""'
    for no, t, sh, exp, act in ints:
        act = txt(act) if exp != "مؤجل" else act
        add(no, t, f"فحص تكامل ({sh})", exp, act, how="معادلة حية" if exp != "مؤجل" else "مؤجل", deferred=(exp == "مؤجل"))

    r = H + 1
    for t in tests:
        vals = [t["no"], t["title"], t["scen"], t["expected"], t["actual"], None, t["how"], t["note"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if t["deferred"]:
            ws.cell(r, 6, "مؤجل (المرحلة 2)")
        else:
            ws.cell(r, 6, f'=IF(E{r}=D{r},"نجح","فشل")')
        ws.cell(r, 6).font = Font(name=FONT, bold=True)
        r += 1
    last = r - 1
    put(ws, "A4", "عدد الاختبارات الحية", bold=True, fill=C_HDR_BAND)
    put(ws, "B4", f'=COUNTIF(F{H + 1}:F{last},"نجح")+COUNTIF(F{H + 1}:F{last},"فشل")')
    put(ws, "C4", "نجح", bold=True, fill=C_HDR_BAND)
    put(ws, "D4", f'=COUNTIF(F{H + 1}:F{last},"نجح")', name="TEST_Pass", wb=wb)
    put(ws, "C5", "فشل", bold=True, fill=C_HDR_BAND)
    put(ws, "D5", f'=COUNTIF(F{H + 1}:F{last},"فشل")', name="TEST_Fail", wb=wb)
    put(ws, "C6", "مؤجل (المرحلة 2)", bold=True, fill=C_HDR_BAND)
    put(ws, "D6", f'=COUNTIF(F{H + 1}:F{last},"مؤجل*")')
    ws.conditional_formatting.add(f"F{H + 1}:F{last}", FormulaRule(formula=[f'F{H + 1}="نجح"'], fill=PatternFill("solid", bgColor=C_GREEN, fgColor=C_GREEN)))
    ws.conditional_formatting.add(f"F{H + 1}:F{last}", FormulaRule(formula=[f'F{H + 1}="فشل"'], fill=PatternFill("solid", bgColor=C_RED, fgColor=C_RED)))
    ws.conditional_formatting.add(f"F{H + 1}:F{last}", FormulaRule(formula=[f'LEFT(F{H + 1},4)="مؤجل"'], fill=PatternFill("solid", bgColor=C_AMBER, fgColor=C_AMBER)))
    ws.freeze_panes = ws.cell(H + 1, 3)
    protect_ws(ws)
    return ws, tests
